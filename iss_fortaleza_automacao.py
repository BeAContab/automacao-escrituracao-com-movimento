"""Automação visível do portal da ISS de Fortaleza.

Este módulo concentra o fluxo do navegador usando Selenium WebDriver para manter
`extrair_nf_pdfs.py` responsável apenas pela extração dos PDFs e pela orquestração
da CLI.
"""

from __future__ import annotations

import os
import re
import subprocess
import sys
import threading
import time
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any, Callable
from urllib.error import URLError
from urllib.request import urlopen

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from selenium import webdriver
from selenium.common.exceptions import (
    ElementNotInteractableException,
    JavascriptException,
    NoSuchElementException,
    StaleElementReferenceException,
    TimeoutException,
    WebDriverException,
)
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.action_chains import ActionChains
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait
from webdriver_manager.chrome import ChromeDriverManager
from tratamento_erros import registrar_evento_execucao, configurar_pasta_logs

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

URL_ISS_FORTALEZA = "https://iss.fortaleza.ce.gov.br/grpfor/home.seam"
PORTA_DEBUG_CHROME_PADRAO = 9222
PERFIL_CHROME_FUNCAO2_PADRAO = Path("brain/navegador_funcao2_profile")

# Mapeamentos de mês por nome (normalizado sem acento)
MESES_POR_NOME = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}

# Abreviações dos meses usadas no calendário do portal
MESES_ABREVIADOS_PORTAL = {
    1: "jan",
    2: "fev",
    3: "mar",
    4: "abr",
    5: "mai",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "set",
    10: "out",
    11: "nov",
    12: "dez",
}


# ---------------------------------------------------------------------------
# Estruturas de dados
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class CompetenciaTrabalho:
    """Representa a competência trabalhada no portal."""

    mes: int
    ano: int
    rotulo: str

    @property
    def nome_mes(self) -> str:
        nomes = [
            "",
            "Janeiro",
            "Fevereiro",
            "Março",
            "Abril",
            "Maio",
            "Junho",
            "Julho",
            "Agosto",
            "Setembro",
            "Outubro",
            "Novembro",
            "Dezembro",
        ]
        return nomes[self.mes]


@dataclass(frozen=True)
class DocumentoPortalISS:
    """Linha da planilha que será transportada para o formulário do portal."""

    arquivo_pdf: str
    cnpj_prestador: str
    numero_nf: str
    id_cnae_final: str
    data_emissao: str
    descricao_servico: str
    uf_local_prestacao: str
    cidade_local_prestacao: str
    natureza_operacao: str
    iss_retido: str
    valor_servico: str
    # Campos de identificação e endereço do prestador (opcionais — usados no preenchimento manual)
    prefeitura: str = ""
    nome_prestador: str = ""
    uf_prestador: str = ""
    cidade_prestador: str = ""
    cep_prestador: str = ""
    logradouro_prestador: str = ""
    numero_prestador: str = ""
    bairro_prestador: str = ""
    email_prestador: str = ""
    valor_deducoes: str = ""
    descontos_incondicionados: str = ""
    descontos_condicionados: str = ""
    outras_retencoes: str = ""
    ir: str = ""
    pis_nao_retido: str = ""
    cofins_nao_retido: str = ""
    csrf: str = ""
    inss: str = ""
    aliquota: str = ""
    regime_tributario: str = "OUTROS"
    tipo_cliente_prestador: str = ""


# ---------------------------------------------------------------------------
# Exceções customizadas
# ---------------------------------------------------------------------------


class PrestadorNaoEncontradoError(RuntimeError):
    """Sinaliza que o portal não encontrou uma razão social para o CNPJ informado."""

    def __init__(self, cnpj: str, detalhe_tela: str = "") -> None:
        self.cnpj = cnpj
        self.detalhe_tela = detalhe_tela.strip()
        mensagem = f"Nenhuma Razão Social Encontrada para o CNPJ {cnpj}."
        if self.detalhe_tela:
            mensagem = f"{mensagem} Detalhe da tela: {self.detalhe_tela}"
        super().__init__(mensagem)


class CompetenciaSemEscriturarDisponivelError(RuntimeError):
    """Sinaliza que a competência consultada não liberou o botão de escriturar."""


class AutomacaoCanceladaError(RuntimeError):
    """Sinaliza que o operador cancelou a automação em execução pela GUI."""


# ---------------------------------------------------------------------------
# Verificação de pausa global do robô
# ---------------------------------------------------------------------------

# Protegido por Lock para evitar race conditions entre threads de pausa e execução (FALHA-04)
_LOCK_CALLBACK_PAUSA = threading.Lock()
_CALLBACK_PAUSA = None

def _verificar_pausa() -> None:
    """Bloqueia a execução do robô se o operador tiver solicitado a pausa na GUI."""
    with _LOCK_CALLBACK_PAUSA:
        cb = _CALLBACK_PAUSA
    if cb:
        try:
            cb()
        except Exception:
            pass


# ---------------------------------------------------------------------------
# Verificação de cancelamento global do robô
# ---------------------------------------------------------------------------

# Mesmo padrão de Lock+global usado para a pausa (FALHA-04), mas SEM engolir
# exceção: é este utilitário quem levanta AutomacaoCanceladaError, e ela precisa
# se propagar livremente até o finally de executar_fluxo_iss() e o except de app.py.
_LOCK_CALLBACK_CANCELAMENTO = threading.Lock()
_CALLBACK_CANCELAMENTO = None

def _verificar_cancelamento() -> None:
    """Levanta AutomacaoCanceladaError se o operador tiver solicitado o cancelamento
    da automação pela GUI. O callback registrado deve apenas retornar um bool
    (nunca lançar por si só) — é aqui, fora de qualquer try/except supressor, que
    a exceção de cancelamento é de fato levantada."""
    with _LOCK_CALLBACK_CANCELAMENTO:
        cb = _CALLBACK_CANCELAMENTO
    if cb and cb():
        raise AutomacaoCanceladaError("Automação cancelada pelo operador.")


# ---------------------------------------------------------------------------
# Utilitários de texto
# ---------------------------------------------------------------------------


def normalizar_texto(texto: str) -> str:
    """Remove acentos e padroniza o texto para facilitar o parse."""

    normalizado = unicodedata.normalize("NFD", texto)
    return "".join(
        caractere for caractere in normalizado if unicodedata.category(caractere) != "Mn"
    ).lower()


def _interpretar_iss_retido(valor: str) -> tuple[bool, bool]:
    """Interpreta o valor da coluna ISS_RETIDO.

    Retorna (deve_marcar, reconhecido). Quando `reconhecido` é False, o valor não
    corresponde a um formato claro de "Sim"/"Não" — é tratado como "Não retido"
    por segurança (mesmo comportamento de sempre), mas o chamador deve registrar
    um aviso em vez de deixar isso passar silenciosamente.
    """

    normalizado = normalizar_texto(valor or "").strip()
    if normalizado.startswith("sim"):
        return True, True
    if normalizado.startswith("nao") or normalizado == "n":
        return False, True
    return False, False


def limpar_cnpj_para_digitacao(cnpj: str) -> str:
    """Retorna o CNPJ pronto para digitação no portal.

    Remove apenas a máscara de pontuação (".", "-", "/", espaços) e qualquer
    outro caractere que não seja letra ou dígito, preservando as letras do
    novo formato alfanumérico de CNPJ (Receita Federal). O resultado é
    normalizado para maiúsculas, conforme o layout oficial (12 caracteres
    alfanuméricos + 2 dígitos verificadores numéricos).

    Não usar re.sub(r"\\D+", ...) aqui: isso destruiria as letras do CNPJ
    alfanumérico.
    """

    texto = (cnpj or "").strip().upper()
    return re.sub(r"[^0-9A-Z]", "", texto)


def _eh_cpf_prestador(documento: DocumentoPortalISS) -> bool:
    """True se o prestador for pessoa física.

    Fonte primária: a coluna TIPO_CLIENTE da planilha (explícita, auditável,
    gravada por processamento_xml.py para os dois layouts de XML suportados).
    Fallback: comprimento do CPF/CNPJ (11 dígitos = CPF, 14 = CNPJ, mesmo
    alfanumérico), usado apenas quando a coluna estiver ausente ou vazia —
    planilhas geradas antes dessa coluna existir, por exemplo.
    """

    tipo_norm = normalizar_texto(documento.tipo_cliente_prestador or "").strip()
    if tipo_norm.startswith("pessoa fisica"):
        return True
    if tipo_norm.startswith("pessoa juridica"):
        return False
    return len(limpar_cnpj_para_digitacao(documento.cnpj_prestador or "")) == 11


def validar_campos_obrigatorios(documento: DocumentoPortalISS) -> list[str]:
    """Valida se todos os campos obrigatórios necessários para a escrituração e cadastro do prestador
    estão presentes na planilha."""
    ausentes = []
    
    campos_validar = {
        "cnpj_prestador": "CNPJ do Prestador",
        "numero_nf": "Número da Nota Fiscal",
        "data_emissao": "Data de Emissão",
        "id_cnae_final": "Código CNAE (Atividade)",
        "descricao_servico": "Descrição do Serviço",
        "uf_local_prestacao": "UF do Local de Prestação",
        "cidade_local_prestacao": "Cidade do Local de Prestação",
        "natureza_operacao": "Natureza da Operação",
        "iss_retido": "ISS Retido (Sim/Não)",
        "nome_prestador": "Razão Social / Nome do Prestador",
        "uf_prestador": "UF do Prestador",
        "cidade_prestador": "Cidade do Prestador",
        "cep_prestador": "CEP do Prestador",
        "logradouro_prestador": "Logradouro (Endereço) do Prestador",
        "bairro_prestador": "Bairro do Prestador"
    }
    
    for attr, label in campos_validar.items():
        valor = getattr(documento, attr, None)
        if valor is None:
            ausentes.append(label)
        elif isinstance(valor, str) and not valor.strip():
            ausentes.append(label)
            
    # Validação adicional: garante que ISS_RETIDO tenha formato reconhecível (Sim/Não).
    # Sem isso, um valor ambíguo (ex.: "talvez", "S", célula com espaço/typo) seria
    # tratado como "Não retido" no portal sem qualquer aviso ao operador.
    if documento.iss_retido and documento.iss_retido.strip():
        _, iss_retido_reconhecido = _interpretar_iss_retido(documento.iss_retido)
        if not iss_retido_reconhecido:
            ausentes.append(
                f"ISS Retido (valor '{documento.iss_retido}' não reconhecido; use apenas 'Sim' ou 'Não')"
            )

    # Validação do valor do serviço
    val_serv = (documento.valor_servico or "").strip().replace(" ", "").replace(".", "").replace(",", ".")
    try:
        if not val_serv or float(val_serv) <= 0.0:
            ausentes.append("Valor do Serviço (deve ser maior que zero)")
    except ValueError:
        ausentes.append("Valor do Serviço (formato numérico inválido)")
        
    return ausentes


def _formatar_celula_para_string_de_valor(valor: Any) -> str:
    """Normaliza valores numéricos lidos do Excel para formato de string brasileiro."""
    if valor is None:
        return "0,00"
    if isinstance(valor, (int, float)):
        return f"{valor:.2f}".replace(".", ",")
    
    texto = str(valor).strip()
    if not texto:
        return "0,00"
    # Se for string formatada em padrão americano com ponto decimal, ex: "1500.50"
    if "." in texto and "," not in texto:
        try:
            val_f = float(texto)
            return f"{val_f:.2f}".replace(".", ",")
        except ValueError:
            pass
    return texto


def _normalizar_celula_cnpj(valor: Any) -> str:
    """Normaliza a célula CNPJ_PRESTADOR lida do Excel para uma string consistente.

    openpyxl entrega o tipo que o Excel detectou para a célula: int/float quando
    a célula "parece" um número (típico de CNPJ numérico legado, o que pode ter
    descartado um zero à esquerda), ou str em qualquer outro caso — incluindo
    sempre o novo CNPJ alfanumérico, já que o Excel não converte letras em número.

    Valores impossíveis para um CNPJ (célula vazia, booleano, zero ou negativo)
    retornam string vazia de propósito: assim `validar_campos_obrigatorios()`
    continua acusando "CNPJ do Prestador" ausente e a nota é registrada em
    `log_notas_incompletas.txt` em vez de ser escriturada com um CNPJ inventado.
    """

    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, (int, float)):
        if valor <= 0:
            return ""
        # zfill repõe zeros à esquerda que o Excel descartou ao tratar a célula
        # como número (ex.: 191 -> "00000000000191", CNPJ legado válido).
        return str(int(round(valor))).zfill(14)
    return str(valor).strip()


def limpar_valor_para_digitacao(valor: str) -> str:
    """Retorna o valor em formato digitável para o campo mascarado do portal."""

    texto = (valor or "").strip()
    texto = texto.replace("R$", "").strip()
    texto = texto.replace(".", "")
    return texto


def limpar_numero_para_digitacao(numero: str) -> str:
    """Mantém apenas dígitos para o campo de número do documento."""

    return re.sub(r"\D+", "", numero or "")


def _somente_digitos(texto: str) -> str:
    """Normaliza qualquer valor para comparação por dígitos."""

    return re.sub(r"\D+", "", texto or "")


def _texto_indica_prestador_nao_encontrado(texto: str) -> bool:
    """Reconhece mensagens visíveis que indicam ausência de razão social no portal."""

    texto_normalizado = normalizar_texto(texto or "")
    marcadores = (
        "nenhuma razao social encontrada",
        "nenhuma razao social",
        "nenhum resultado encontrado",
        "nao foram encontrados",
    )
    return any(marcador in texto_normalizado for marcador in marcadores)


# ---------------------------------------------------------------------------
# Logging da função 2
# ---------------------------------------------------------------------------


def registrar_log_funcao2(caminho_log: Path, mensagem: str) -> None:
    """Registra eventos da função 2 com data e hora para rastrear as linhas ignoradas."""

    caminho_log.parent.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with caminho_log.open("a", encoding="utf-8") as arquivo:
        arquivo.write(f"[{timestamp}] {mensagem}\n")


# ---------------------------------------------------------------------------
# Competência
# ---------------------------------------------------------------------------


def construir_competencia(mes: int, ano: int) -> CompetenciaTrabalho:
    """Valida mês e ano e devolve a competência estruturada."""

    if not 1 <= mes <= 12:
        raise ValueError("O mês deve estar entre 1 e 12.")
    if ano < 1000 or ano > 9999:
        raise ValueError("O ano deve estar no formato AAAA.")
    return CompetenciaTrabalho(mes=mes, ano=ano, rotulo=f"{mes:02d}/{ano}")


def interpretar_competencia(texto: str) -> CompetenciaTrabalho:
    """Converte uma entrada do usuário em mês e ano.

    Aceita entradas como `maio 2026`, `maio/2026` ou `05/2026`.
    """

    texto_limpo = texto.strip()
    if not texto_limpo:
        raise ValueError("Competência vazia.")

    padrao_numerico = re.fullmatch(r"(\d{1,2})\s*(?:/|\s)\s*(\d{4})", texto_limpo)
    if padrao_numerico:
        mes = int(padrao_numerico.group(1))
        ano = int(padrao_numerico.group(2))
        return construir_competencia(mes, ano)

    texto_norm = normalizar_texto(texto_limpo)
    padrao_nome = re.fullmatch(r"([a-z]+)\s+de\s+(\d{4})", texto_norm) or re.fullmatch(
        r"([a-z]+)\s+(\d{4})", texto_norm
    )
    if not padrao_nome:
        raise ValueError(
            "Informe a competência no formato 'maio 2026' ou '05/2026'."
        )

    nome_mes = padrao_nome.group(1)
    ano = int(padrao_nome.group(2))
    mes = MESES_POR_NOME.get(nome_mes)
    if mes is None:
        raise ValueError(
            "Não foi possível identificar o mês informado. Use, por exemplo, 'maio 2026'."
        )
    return construir_competencia(mes, ano)


def solicitar_competencia() -> CompetenciaTrabalho:
    """Pergunta ao usuário a competência antes de abrir o navegador."""

    while True:
        try:
            mes = int(input("Informe o mês de trabalho (1 a 12): ").strip())
            ano_texto = input("Informe o ano de trabalho (AAAA): ").strip()
            if not re.fullmatch(r"\d{4}", ano_texto):
                raise ValueError("O ano deve conter exatamente 4 dígitos.")
            return construir_competencia(mes, int(ano_texto))
        except ValueError as exc:
            print(f"Entrada inválida: {exc}")
            print()


def solicitar_nova_competencia_para_repetir(
    competencia_atual: CompetenciaTrabalho,
) -> CompetenciaTrabalho:
    """Pergunta ao usuário outra competência para reiniciar a seleção De/Até."""

    while True:
        texto = input(
            f"A competência {competencia_atual.rotulo} está sem o botão Escriturar habilitado. "
            "Informe outra competência para voltar aos campos De e Até (ou pressione Enter para cancelar): "
        ).strip()
        if not texto:
            raise SystemExit(
                "A automação foi interrompida porque não foi informada uma nova competência."
            )
        try:
            return interpretar_competencia(texto)
        except ValueError as exc:
            print(f"Competência inválida: {exc}")
            print()


def solicitar_planilha_automacao() -> Path:
    """Pergunta ao usuário qual planilha XLSX será usada na automação."""

    texto = input(
        "Informe o caminho da planilha XLSX de automação (Enter para usar nf_compilado.xlsx): "
    ).strip().strip('"')
    return Path(texto) if texto else Path("nf_compilado.xlsx")


# ---------------------------------------------------------------------------
# Carregamento da planilha
# ---------------------------------------------------------------------------


def carregar_documentos_xlsx(caminho_xlsx: Path) -> list[DocumentoPortalISS]:
    """Lê a planilha de notas e devolve as linhas prontas para o portal."""

    workbook = load_workbook(caminho_xlsx, data_only=True)
    planilha = workbook.active

    cabecalhos = [celula.value for celula in next(planilha.iter_rows(min_row=1, max_row=1))]
    colunas = {nome: indice for indice, nome in enumerate(cabecalhos)}

    # Suporte flexível para arquivo_xml ou arquivo_pdf
    col_arquivo = "ARQUIVO_XML" if "ARQUIVO_XML" in colunas else ("ARQUIVO_PDF" if "ARQUIVO_PDF" in colunas else None)
    if not col_arquivo:
        raise ValueError("A planilha informada deve conter a coluna 'ARQUIVO_XML' ou 'ARQUIVO_PDF'.")

    campos_obrigatorios = [
        "CNPJ_PRESTADOR",
        "NUMERO_NF",
        "ID_CNAE_FINAL",
        "DATA_EMISSAO",
        "DESCRICAO_SERVICO",
        "UF_LOCAL_PRESTACAO",
        "CIDADE_LOCAL_PRESTACAO",
        "NATUREZA_OPERACAO",
        "ISS_RETIDO",
        "VALOR_SERVICO",
    ]
    faltantes = [campo for campo in campos_obrigatorios if campo not in colunas]
    if faltantes:
        raise ValueError(
            f"A planilha informada não contém as colunas obrigatórias: {', '.join(faltantes)}"
        )

    documentos: list[DocumentoPortalISS] = []
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if not any(valor is not None and str(valor).strip() for valor in linha):
            continue

        valor_servico = _formatar_celula_para_string_de_valor(linha[colunas["VALOR_SERVICO"]]) if "VALOR_SERVICO" in colunas else "0,00"
        valor_deducoes = _formatar_celula_para_string_de_valor(linha[colunas["VALOR_DEDUCOES"]]) if "VALOR_DEDUCOES" in colunas else "0,00"
        descontos_incondicionados = _formatar_celula_para_string_de_valor(linha[colunas["DESCONTOS_INCONDICIONADOS"]]) if "DESCONTOS_INCONDICIONADOS" in colunas else "0,00"
        descontos_condicionados = _formatar_celula_para_string_de_valor(linha[colunas["DESCONTOS_CONDICIONADOS"]]) if "DESCONTOS_CONDICIONADOS" in colunas else "0,00"
        outras_retencoes = _formatar_celula_para_string_de_valor(linha[colunas["OUTRAS_RETENCOES"]]) if "OUTRAS_RETENCOES" in colunas else "0,00"
        ir = _formatar_celula_para_string_de_valor(linha[colunas["IR"]]) if "IR" in colunas else "0,00"
        pis_nao_retido = _formatar_celula_para_string_de_valor(linha[colunas["PIS_NAO_RETIDO"]]) if "PIS_NAO_RETIDO" in colunas else "0,00"
        cofins_nao_retido = _formatar_celula_para_string_de_valor(linha[colunas["COFINS_NAO_RETIDO"]]) if "COFINS_NAO_RETIDO" in colunas else "0,00"
        csrf = _formatar_celula_para_string_de_valor(linha[colunas["CSRF (CSLL + PIS + Cofins Retidos)"]]) if "CSRF (CSLL + PIS + Cofins Retidos)" in colunas else "0,00"
        inss = _formatar_celula_para_string_de_valor(linha[colunas["INSS"]]) if "INSS" in colunas else "0,00"
        aliquota = _formatar_celula_para_string_de_valor(linha[colunas["ALIQUOTA"]]) if "ALIQUOTA" in colunas else ""

        # Leitura flexível da prefeitura
        prefeitura_val = str(linha[colunas["PREFEITURA"]] or "") if "PREFEITURA" in colunas else ""

        documentos.append(
            DocumentoPortalISS(
                arquivo_pdf=str(linha[colunas[col_arquivo]] or ""),
                prefeitura=prefeitura_val,
                cnpj_prestador=_normalizar_celula_cnpj(linha[colunas["CNPJ_PRESTADOR"]]),
                numero_nf=str(linha[colunas["NUMERO_NF"]] or ""),
                id_cnae_final=str(linha[colunas["ID_CNAE_FINAL"]] or ""),
                data_emissao=str(linha[colunas["DATA_EMISSAO"]] or ""),
                descricao_servico=str(linha[colunas["DESCRICAO_SERVICO"]] or ""),
                uf_local_prestacao=str(linha[colunas["UF_LOCAL_PRESTACAO"]] or ""),
                cidade_local_prestacao=str(linha[colunas["CIDADE_LOCAL_PRESTACAO"]] or ""),
                natureza_operacao=str(linha[colunas["NATUREZA_OPERACAO"]] or ""),
                iss_retido=str(linha[colunas["ISS_RETIDO"]] or ""),
                valor_servico=valor_servico,
                # Campos de endereço do prestador (lidos quando presentes na planilha)
                nome_prestador=str(linha[colunas["NOME_PRESTADOR"]] or "") if "NOME_PRESTADOR" in colunas else "",
                uf_prestador=str(linha[colunas["UF_PRESTADOR"]] or "") if "UF_PRESTADOR" in colunas else "",
                cidade_prestador=str(linha[colunas["CIDADE_PRESTADOR"]] or "") if "CIDADE_PRESTADOR" in colunas else "",
                cep_prestador=str(linha[colunas["CEP_PRESTADOR"]] or "") if "CEP_PRESTADOR" in colunas else "",
                logradouro_prestador=str(linha[colunas["LOGRADOURO_PRESTADOR"]] or "") if "LOGRADOURO_PRESTADOR" in colunas else "",
                numero_prestador=str(linha[colunas["NUMERO_PRESTADOR"]] or "") if "NUMERO_PRESTADOR" in colunas else "",
                bairro_prestador=str(linha[colunas["BAIRRO_PRESTADOR"]] or "") if "BAIRRO_PRESTADOR" in colunas else "",
                email_prestador=str(linha[colunas["EMAIL_PRESTADOR"]] or "") if "EMAIL_PRESTADOR" in colunas else "",
                valor_deducoes=valor_deducoes,
                descontos_incondicionados=descontos_incondicionados,
                descontos_condicionados=descontos_condicionados,
                outras_retencoes=outras_retencoes,
                ir=ir,
                pis_nao_retido=pis_nao_retido,
                cofins_nao_retido=cofins_nao_retido,
                csrf=csrf,
                inss=inss,
                aliquota=aliquota,
                regime_tributario=str(linha[colunas["REGIME_TRIBUTARIO"]] or "OUTROS") if "REGIME_TRIBUTARIO" in colunas else "OUTROS",
                tipo_cliente_prestador=str(linha[colunas["TIPO_CLIENTE"]] or "") if "TIPO_CLIENTE" in colunas else "",
            )
        )
    return documentos


# ---------------------------------------------------------------------------
# Inicialização do navegador Chrome com Selenium
# ---------------------------------------------------------------------------


def _criar_opcoes_chrome(
    depuracao: bool = False,
    perfil_persistente: Path | None = None,
    pasta_downloads: Path | None = None,
) -> Options:
    """Monta as opções do Chrome para o Selenium, incluindo proteções anti-detecção."""

    opcoes = Options()
    opcoes.add_argument("--start-maximized")
    opcoes.add_argument("--no-first-run")
    opcoes.add_argument("--no-default-browser-check")
    # Oculta a flag de automação no navegador para evitar bloqueios do portal
    opcoes.add_argument("--disable-blink-features=AutomationControlled")
    opcoes.add_experimental_option("excludeSwitches", ["enable-automation"])
    opcoes.add_experimental_option("useAutomationExtension", False)
    if perfil_persistente:
        perfil_persistente.mkdir(parents=True, exist_ok=True)
        opcoes.add_argument(f"--user-data-dir={perfil_persistente.resolve()}")
    # Configura pasta de download automático sem exibir diálogos de salvamento
    if pasta_downloads:
        pasta_downloads.mkdir(parents=True, exist_ok=True)
        prefs = {
            "download.default_directory": str(pasta_downloads.resolve()),
            "download.prompt_for_download": False,
            "download.directory_upgrade": True,
            "safebrowsing.enabled": True,
        }
        opcoes.add_experimental_option("prefs", prefs)
    return opcoes


def _inicializar_driver(opcoes: Options) -> WebDriver:
    """Inicializa o ChromeDriver com estratégia híbrida: Selenium Manager nativo → webdriver-manager.
    
    Centraliza a lógica de fallback que antes era duplicada em abrir_navegador_visivel()
    e abrir_navegador_com_perfil_persistente() (MELHORIA-07).
    """
    try:
        # Tenta inicializar nativamente usando o Selenium Manager (embutido no Selenium 4.6+)
        # Isso evita problemas com downloads de drivers bloqueados por proxy/firewall ou falhas de SSL.
        driver = webdriver.Chrome(options=opcoes)
    except Exception as e_native:
        print(f"Aviso: não foi possível iniciar o Chrome de forma nativa ({e_native}). Tentando com webdriver-manager...")
        try:
            # Fallback para o webdriver-manager convencional
            servico = Service(ChromeDriverManager().install())
            driver = webdriver.Chrome(service=servico, options=opcoes)
        except Exception as e_fallback:
            msg_erro = (
                f"Erro crítico: Não foi possível inicializar o navegador Google Chrome.\n"
                f"Detalhes do erro nativo: {e_native}\n"
                f"Detalhes do erro de fallback: {e_fallback}\n"
                f"Por favor, verifique se o Google Chrome está instalado corretamente nesta máquina."
            )
            registrar_evento_execucao(msg_erro, "ISS Fortaleza")
            raise RuntimeError(msg_erro) from e_fallback
    return driver


def abrir_navegador_visivel(
    depuracao: bool = False,
    pasta_downloads: Path | None = None,
) -> WebDriver:
    """Abre um Chrome visível e maximizado, com inicialização híbrida resiliente."""

    opcoes = _criar_opcoes_chrome(depuracao=depuracao, pasta_downloads=pasta_downloads)
    driver = _inicializar_driver(opcoes)

    # Remove o atributo webdriver para evitar a detecção pelo portal
    driver.execute_script(
        "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
    )
    driver.get(URL_ISS_FORTALEZA)
    return driver


def abrir_navegador_com_perfil_persistente(
    perfil_persistente: Path,
    depuracao: bool = False,
    pasta_downloads: Path | None = None,
) -> WebDriver:
    """Abre o Chrome com perfil de usuário persistente com inicialização híbrida resiliente."""

    opcoes = _criar_opcoes_chrome(depuracao=depuracao, perfil_persistente=perfil_persistente, pasta_downloads=pasta_downloads)
    # Reutiliza _inicializar_driver() para evitar duplicação da lógica de fallback (MELHORIA-07)
    driver = _inicializar_driver(opcoes)

    driver.execute_script(
        "Object.defineProperty(navigator, 'webdriver', {get: () => undefined})"
    )
    driver.get(URL_ISS_FORTALEZA)
    registrar_evento_execucao(
        f"Navegador com perfil persistente aberto: {perfil_persistente.resolve()}",
        "ISS Fortaleza",
    )
    return driver


# ---------------------------------------------------------------------------
# Esperas e cliques com Selenium
# ---------------------------------------------------------------------------


def _aguardar_elemento(
    driver: WebDriver,
    by: str,
    seletor: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda um elemento ficar visível e retorna o WebElement."""

    return WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((by, seletor))
    )


def _aguardar_elemento_clicavel(
    driver: WebDriver,
    by: str,
    seletor: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda um elemento ficar clicável (visível + habilitado) e o retorna."""

    return WebDriverWait(driver, timeout).until(
        EC.element_to_be_clickable((by, seletor))
    )


def _clicar_com_espera(
    driver: WebDriver,
    by: str,
    seletor: str,
    descricao: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda o elemento e efetua o clique, com fallback via JavaScript e retry para elementos obsoletos."""

    _verificar_pausa()
    _verificar_cancelamento()
    for tentativa in range(3):
        try:
            elemento = _aguardar_elemento_clicavel(driver, by, seletor, timeout=timeout)
            try:
                elemento.click()
            except (ElementNotInteractableException, WebDriverException) as exc:
                if isinstance(exc, StaleElementReferenceException):
                    raise
                # Fallback via JavaScript para elementos que bloqueiam o clique nativo
                driver.execute_script("arguments[0].click();", elemento)
            return elemento
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)
    # Sentinela: garante que falhas de timeout ou inesperadas sejam exibidas ao chamador
    raise RuntimeError(f"Não foi possível clicar no elemento '{descricao}' após 3 tentativas.")


def _clicar_por_id(driver: WebDriver, elemento_id: str, timeout: int = 10) -> None:
    """Clica em um elemento pelo ID com fallback via JavaScript."""

    _clicar_com_espera(driver, By.ID, elemento_id, f"elemento ID={elemento_id}", timeout=timeout)


# O portal usa o MESMO componente de modal (RichFaces) para vários avisos de
# confirmação ao clicar em Gravar — só o texto interno muda. Exemplos já
# observados: "Já existe documento fiscal escriturado com o CNPJ do
# prestador, com o mesmo número de nota, na competência ..." (possível
# duplicata) e "Prestador não inscrito no CPOM. ..." (aviso informativo).
_XPATH_MODAL_CONFIRMACAO_CONTAINER = (
    '//*[@id="digitarDocumentoForm:confirmacao_customizadaContainer"]'
)
_XPATH_MODAL_CONFIRMACAO_H3 = _XPATH_MODAL_CONFIRMACAO_CONTAINER + '//h3'
_XPATH_HEADING_SUCESSO_GRAVACAO = (
    '//*[@id="content"]/legend/h2[contains(., "digitado com") '
    'and (contains(., "Sucesso") or contains(., "sucesso"))]'
)

# Marcador de texto (normalizado, sem acento/caixa) que identifica especificamente
# o aviso de possível nota duplicada — o único caso em que a decisão é "Não".
# Qualquer outro texto nesse mesmo componente é tratado como aviso informativo
# (decisão "Sim").
_MARCADOR_TEXTO_NOTA_DUPLICADA = "documento fiscal escriturado"

# Mensagem de validação inline (NÃO é o modal de confirmação acima — é o
# componente <rich:messages> do JSF, sem botão, que bloqueia a gravação até o
# formulário ser corrigido): "CNAE não incide Imposto Sobre Serviço. Solução:
# Selecione a Natureza de Operação 'Não Incidência'." A correção é trocar a
# Natureza da Operação e tentar gravar de novo — não há botão para clicar aqui.
_XPATH_MENSAGEM_CNAE_NAO_INCIDE = (
    '//span[contains(@class, "rich-messages-label") and contains(., "CNAE não incide")]'
)


def _clicar_botao_modal_confirmacao(driver: WebDriver, valor_botao: str) -> None:
    """Clica no botão ("Sim" ou "Não") do modal de confirmação genérico do
    portal e aguarda o modal desaparecer da tela antes de prosseguir."""

    xpath_botao = f'{_XPATH_MODAL_CONFIRMACAO_CONTAINER}//input[@value="{valor_botao}"]'
    try:
        driver.find_element(By.XPATH, xpath_botao).click()
    except Exception:
        # Fallback pelos IDs auto-gerados pelo JSF observados no momento da implementação
        # (frágil a mudanças de versão do portal, por isso o XPath acima é a via principal).
        id_fallback = "digitarDocumentoForm:j_id489" if valor_botao == "Sim" else "digitarDocumentoForm:j_id493"
        driver.find_element(By.ID, id_fallback).click()

    WebDriverWait(driver, 10).until(
        EC.invisibility_of_element_located((By.XPATH, _XPATH_MODAL_CONFIRMACAO_H3))
    )


def _aguardar_resultado_da_gravacao(driver: WebDriver, timeout: int = 15) -> str:
    """Aguarda o desfecho do clique em "Gravar": sucesso normal, ou algum dos
    modais de confirmação que o portal pode exibir antes de gravar de fato.

    - Se o texto do modal indicar possível nota duplicada (mesmo CNPJ do
      prestador + mesmo número de nota já escriturados em outra competência),
      clica em "Não" (decisão do operador: nunca confirmar a geração de uma
      possível duplicata automaticamente) e retorna "duplicata".
    - Qualquer outro texto de confirmação (ex.: "Prestador não inscrito no
      CPOM...") é tratado como aviso informativo: clica em "Sim" e retorna
      "confirmacao_generica" — cabe ao chamador clicar em Gravar de novo.
    - Se a mensagem de validação "CNAE não incide Imposto Sobre Serviço..."
      aparecer (sem botão, bloqueia a gravação), retorna "cnae_nao_incide" —
      cabe ao chamador trocar a Natureza da Operação e tentar gravar de novo.
    - Quando o cabeçalho de sucesso aparece, retorna "sucesso".
    - Se nada aparecer dentro do prazo, levanta TimeoutException.
    """

    fim = time.monotonic() + timeout
    while time.monotonic() < fim:
        _verificar_cancelamento()

        if driver.find_elements(By.XPATH, _XPATH_HEADING_SUCESSO_GRAVACAO):
            return "sucesso"

        elementos_cnae = driver.find_elements(By.XPATH, _XPATH_MENSAGEM_CNAE_NAO_INCIDE)
        if any(el.is_displayed() for el in elementos_cnae):
            return "cnae_nao_incide"

        elementos_modal = driver.find_elements(By.XPATH, _XPATH_MODAL_CONFIRMACAO_H3)
        elementos_modal = [el for el in elementos_modal if el.is_displayed()]
        if elementos_modal:
            texto_modal = normalizar_texto(elementos_modal[0].text or "")
            if _MARCADOR_TEXTO_NOTA_DUPLICADA in texto_modal:
                registrar_evento_execucao(
                    "Portal indicou que já existe documento fiscal escriturado com o mesmo "
                    "CNPJ do prestador e o mesmo número de nota (em outra competência). "
                    "Clicando em 'Não' para não gravar uma possível duplicata.",
                    "ISS Fortaleza",
                )
                _clicar_botao_modal_confirmacao(driver, "Não")
                return "duplicata"

            registrar_evento_execucao(
                f"Portal exibiu confirmação: '{elementos_modal[0].text.strip()}'. "
                "Clicando em 'Sim'.",
                "ISS Fortaleza",
            )
            _clicar_botao_modal_confirmacao(driver, "Sim")
            return "confirmacao_generica"

        time.sleep(0.3)

    raise TimeoutException("Timeout aguardando confirmação de gravação do documento.")


def _clicar_gravar_documento(driver: WebDriver) -> None:
    """Clica no botão "Gravar Documento", com fallback via XPath por valor."""

    try:
        _clicar_por_id(driver, "digitarDocumentoForm:j_id477", timeout=15)
    except Exception:
        el = driver.find_element(
            By.XPATH,
            "//*[@id='digitarDocumentoForm:j_id477'] | //input[@value='Gravar' or @value='Gravar Documento']",
        )
        driver.execute_script("arguments[0].click();", el)


def _gravar_documento_com_confirmacoes(
    driver: WebDriver, numero_nf: str, timeout: int = 15, max_confirmacoes: int = 4
) -> str:
    """Clica em Gravar Documento e resolve os bloqueios que o portal pode
    encadear antes de gravar de fato: nota duplicada (clica "Não", nunca
    grava), avisos informativos como "Prestador não inscrito no CPOM" (clica
    "Sim" e tenta Gravar de novo), ou o erro de validação "CNAE não incide
    Imposto Sobre Serviço" (troca a Natureza da Operação para "Não
    Incidência" e tenta Gravar de novo) — repete até `max_confirmacoes`
    vezes, já que o portal pode encadear mais de um bloqueio para a mesma
    nota. Retorna "sucesso" ou "duplicata".
    """

    _clicar_gravar_documento(driver)
    for _ in range(max_confirmacoes):
        resultado = _aguardar_resultado_da_gravacao(driver, timeout=timeout)
        if resultado in ("sucesso", "duplicata"):
            return resultado

        if resultado == "cnae_nao_incide":
            # O portal não deixa gravar esse CNAE com a Natureza atual; a correção
            # é forçar "Não Incidência" (sobrescrevendo o que a planilha continha
            # para esta nota) e tentar gravar de novo. O checkbox de ISS Retido
            # não precisa de tratamento aqui: o próprio portal o remove da tela ao
            # selecionar "Não Incidência".
            registrar_evento_execucao(
                f"Portal recusou a gravação da NF {numero_nf}: CNAE não incide ISS. "
                "Selecionando Natureza da Operação = 'Não Incidência' (sobrescrevendo "
                "o valor da planilha para esta nota) e tentando gravar novamente.",
                "ISS Fortaleza",
            )
            _selecionar_opcao_por_texto(
                driver,
                By.ID,
                "digitarDocumentoForm:comboEscolherLocalPrestacao",
                ["Não Incidência", "Nao Incidencia"],
            )
            # Mesma pausa de estabilização de AJAX já usada para esse campo em
            # preencher_documento_servico().
            time.sleep(1.5)
        else:
            # resultado == "confirmacao_generica": _aguardar_resultado_da_gravacao já
            # clicou "Sim" no modal; tenta gravar de novo para ver o próximo desfecho.
            registrar_evento_execucao(
                f"Confirmação genérica do portal resolvida (Sim) para a NF {numero_nf}; "
                "tentando Gravar Documento novamente.",
                "ISS Fortaleza",
            )

        _clicar_gravar_documento(driver)

    raise TimeoutException(
        f"Excedeu o limite de {max_confirmacoes} confirmações encadeadas ao gravar a NF {numero_nf}."
    )


def _pausar_para_depuracao(
    driver: WebDriver,
    habilitado: bool,
    mensagem: str,
) -> None:
    """Interrompe o fluxo para inspeção manual quando o modo de depuração estiver ativo."""

    if not habilitado:
        return
    print(mensagem)
    if sys.stdin.isatty():
        input(f"{mensagem}\nPressione Enter para continuar com o próximo clique...")


# ---------------------------------------------------------------------------
# Digitação e seleção de campos
# ---------------------------------------------------------------------------


def _digitar_campo(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
    delay_ms: int = 80,
    normalizador_fallback: Callable[[str], str] | None = None,
) -> None:
    """Localiza o campo, limpa e digita o texto tecla por tecla com pequena pausa entre elas.

    Possui lógica de retry para se recuperar caso o elemento mude sob AJAX durante a digitação.

    O parâmetro opcional `normalizador_fallback` substitui a normalização padrão
    (somente dígitos) usada na verificação de estabilidade do campo. É necessário
    para campos que podem conter letras (ex.: CNPJ alfanumérico): usar apenas
    dígitos como fallback nesse caso mascararia silenciosamente uma eventual perda
    de letras durante a digitação. Quando omitido, o comportamento é idêntico ao
    anterior para todos os outros campos.
    """

    _verificar_pausa()
    _verificar_cancelamento()
    normalizar = normalizador_fallback or _somente_digitos
    for tentativa in range(3):
        try:
            campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
            campo.click()
            campo.send_keys(Keys.CONTROL + "a")
            campo.send_keys(Keys.BACKSPACE)
            time.sleep(0.12)

            # Digita caractere por caractere para disparar os eventos key* do portal
            for caractere in texto:
                campo.send_keys(caractere)
                time.sleep(delay_ms / 1000)

            # Verifica se o valor ficou estável antes de seguir
            esperado_normalizado = normalizar(texto)
            for _ in range(20):
                try:
                    valor_atual = campo.get_attribute("value") or ""
                    if valor_atual.strip() == texto.strip():
                        return
                    if esperado_normalizado and normalizar(valor_atual) == esperado_normalizado:
                        return
                except StaleElementReferenceException:
                    raise
                time.sleep(0.15)

            raise RuntimeError(f"O campo {seletor} não permaneceu com o valor esperado: {texto}")
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)


def _definir_valor_instantaneo(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
) -> None:
    """Define o valor de um campo de texto instantaneamente executando JavaScript no navegador.

    Isso é utilizado para campos de textos muito longos, como a descrição dos serviços,
    evitando o gargalo de tempo da digitação caractere por caractere via Selenium.
    Dispara os eventos 'input' e 'change' para garantir que os ouvintes de eventos da página
    e o estado do JSF reconheçam a alteração do valor.
    """
    _verificar_pausa()
    _verificar_cancelamento()
    for tentativa in range(3):
        try:
            campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
            driver.execute_script(
                "arguments[0].value = arguments[1];"
                "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
                "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
                campo,
                texto,
            )
            return
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)


def _definir_campo_rapido(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
    delay_ms: int = 80,
    normalizador_fallback: Callable[[str], str] | None = None,
) -> None:
    """Tenta preencher o campo instantaneamente via JavaScript (bem mais rápido
    que digitar caractere por caractere) e confere se o valor realmente ficou
    correto. Se não ficou — vazio, ou diferente do esperado (ex.: alguma
    validação/máscara do site descartando o valor) — cai para `_digitar_campo`
    (digitação caractere por caractere, comportamento de sempre) como fallback
    garantido.
    """

    _verificar_pausa()
    _verificar_cancelamento()
    normalizar = normalizador_fallback or _somente_digitos

    try:
        campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
        driver.execute_script(
            "arguments[0].value = arguments[1];"
            "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
            "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
            campo,
            texto,
        )
        time.sleep(0.3)  # pequena espera para eventual validação JS do portal reagir

        valor_atual = campo.get_attribute("value") or ""
        if valor_atual.strip() == texto.strip() or normalizar(valor_atual) == normalizar(texto):
            return
    except (StaleElementReferenceException, NoSuchElementException):
        pass  # cai para o fallback abaixo

    _digitar_campo(driver, by, seletor, texto, delay_ms=delay_ms, normalizador_fallback=normalizador_fallback)


def _campo_aliquota_esta_editavel(driver: WebDriver) -> WebElement | None:
    """Retorna o elemento do campo Alíquota se ele estiver editável (um <input>
    sem disabled/readonly), ou None se estiver bloqueado pelo portal.

    O campo é calculado automaticamente para alguns CNAEs — nesse caso o portal
    renderiza um <span> (não um <input>), sem nenhum campo para digitar. Só
    tentamos preencher quando existe de fato um <input> habilitado.
    """

    try:
        elemento = driver.find_element(By.ID, "digitarDocumentoForm:idAliquota")
    except NoSuchElementException:
        return None
    if elemento.tag_name.lower() != "input":
        return None
    if elemento.get_attribute("readonly") or elemento.get_attribute("disabled"):
        return None
    if "disabled" in (elemento.get_attribute("class") or ""):
        return None
    return elemento


def _ler_estado_checkbox_apos_espera(
    driver: WebDriver,
    by: str,
    seletor: str,
    espera: float,
):
    """Aguarda `espera` segundos e retorna checkbox.is_selected(), ou None se o
    elemento não for encontrado.

    Usado para dar tempo a um onchange AJAX assíncrono do portal (ex.: a
    mudança de Natureza da Operação, que pode marcar/desmarcar sozinho o
    checkbox de ISS Retido) terminar de processar antes de ler o estado real.
    Uma leitura por "debounce" (esperar até duas leituras seguidas baterem)
    foi cogitada e descartada: como o polling costuma acontecer mais rápido
    que a resposta do AJAX, duas leituras consecutivas podem coincidir
    ANTES da mudança tardia chegar, declarando "estável" cedo demais — uma
    espera de duração fixa (mesmo padrão já usado nesta função para Status
    NFSE e UF do Prestador) é mais confiável aqui.
    """

    time.sleep(espera)
    try:
        return driver.find_element(by, seletor).is_selected()
    except NoSuchElementException:
        return None


def _selecionar_opcao_por_texto(
    driver: WebDriver,
    by: str,
    seletor: str,
    textos: list[str],
    timeout: int = 10,
) -> None:
    """Seleciona a primeira opção disponível em um <select> a partir de uma lista de rótulos."""

    _verificar_pausa()
    _verificar_cancelamento()
    # Aumentado para 8 tentativas para maior tolerância ao AJAX do portal
    for tentativa in range(8):
        try:
            elemento = _aguardar_elemento_clicavel(driver, by, seletor, timeout=timeout)
            select = Select(elemento)

            # Se a lista de opções estiver vazia ou contiver apenas o placeholder inicial, forçamos o retry
            # pois o portal (RichFaces) provavelmente ainda está popularizando o <select> via AJAX assíncrono.
            if len(select.options) <= 1:
                raise NoSuchElementException("Select ainda não populado com opções (AJAX pendente).")

            # 1. Tenta correspondência exata normalizada (ignora caixa e acentos)
            for texto in textos:
                texto_norm = normalizar_texto(texto)
                for opcao in select.options:
                    opcao_texto_norm = normalizar_texto(opcao.text)
                    opcao_valor_norm = normalizar_texto(opcao.get_attribute("value") or "")
                    if (texto_norm == opcao_texto_norm) or (texto_norm == opcao_valor_norm):
                        select.select_by_value(opcao.get_attribute("value"))
                        return

            # 2. Tenta correspondência parcial normalizada (ex: "Fortaleza" contido em "Fortaleza - CE")
            for texto in textos:
                texto_norm = normalizar_texto(texto)
                for opcao in select.options:
                    opcao_texto_norm = normalizar_texto(opcao.text)
                    if texto_norm in opcao_texto_norm:
                        # Ignora placeholders de instruções como "selecione" ou "escolha"
                        if "selecione" in opcao_texto_norm or "escolha" in opcao_texto_norm:
                            continue
                        select.select_by_value(opcao.get_attribute("value"))
                        return

            raise NoSuchElementException(f"Nenhuma das opções {textos} foi localizada no select.")
            
        except (StaleElementReferenceException, NoSuchElementException) as exc:
            if tentativa == 7:
                raise RuntimeError(
                    f"Não foi possível selecionar nenhuma opção entre {textos} após várias tentativas."
                ) from exc
            time.sleep(0.8)


# ---------------------------------------------------------------------------
# Aguardando login manual
# ---------------------------------------------------------------------------


def aguardar_login_manual(driver: WebDriver) -> None:
    """Mostra a instrução e aguarda o usuário concluir o login manual no terminal."""

    print("O navegador foi aberto no portal da ISS de Fortaleza.")
    print("Faça o login manualmente e, quando terminar, volte aqui e pressione Enter.")
    input("Pressione Enter somente após o login estar concluído...")
    time.sleep(1)


def aguardar_login_manual_por_arquivo(driver: WebDriver, caminho_confirmacao: Path) -> None:
    """Espera até que o arquivo de confirmação seja criado pelo operador."""

    caminho_confirmacao.parent.mkdir(parents=True, exist_ok=True)
    if caminho_confirmacao.exists():
        caminho_confirmacao.unlink()

    print("O navegador foi aberto no portal da ISS de Fortaleza.")
    print("Faça o login manualmente e, ao terminar, crie o arquivo de confirmação para continuar.")
    print(f"Arquivo de confirmação esperado: {caminho_confirmacao.resolve()}")

    while not caminho_confirmacao.exists():
        _verificar_cancelamento()
        time.sleep(1)

    # Evita que o próximo teste reutilize a mesma confirmação antiga
    try:
        caminho_confirmacao.unlink()
    except Exception:
        pass

    time.sleep(1)


# ---------------------------------------------------------------------------
# Seleção de competência via calendário RichFaces
# ---------------------------------------------------------------------------


def _valor_campo_competencia(driver: WebDriver, base_id: str) -> str:
    """Lê o valor do campo de data do calendário RichFaces."""

    try:
        campo = driver.find_element(By.ID, f"{base_id}InputDate")
        return campo.get_attribute("value") or ""
    except NoSuchElementException:
        return ""


def _abrir_editor_calendario(driver: WebDriver, base_id: str, rotulo: str) -> None:
    """Clica no botão de edição do calendário RichFaces para abrir o editor de mês/ano.

    Garante que o calendário popup seja aberto antes de tentar acessar seu cabeçalho.
    """

    # Clica no botão popup do calendário (ícone do calendário) para exibir o calendário na tela.
    # Sem isso, o cabeçalho '#{base_id}Header' não é renderizado ou fica invisível no DOM.
    popup_btn_id = f"{base_id}PopupButton"
    try:
        popup_btn = WebDriverWait(driver, 8).until(
            EC.element_to_be_clickable((By.ID, popup_btn_id))
        )
        popup_btn.click()
        # Pequena pausa para garantir a renderização visual do calendário popup
        time.sleep(0.5)
    except Exception as e:
        print(f"Aviso ao abrir popup do calendário para {rotulo}: {e}")

    # O botão de edição fica no cabeçalho do calendário.
    # Atenção: o base_id contém ':' (ex: 'manterEscrituracaoForm:dataInicial'),
    # portanto não pode ser usado diretamente como seletor CSS '#id'. Usamos o seletor
    # de atributo [id='...'] para contornar isso sem necessidade de escape manual.
    botao = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable(
            (By.CSS_SELECTOR, f"[id='{base_id}Header'] .rich-calendar-tool-btn")
        )
    )
    botao.click()
    # Aguarda o botão OK do editor ficar visível como sinal de abertura do editor
    WebDriverWait(driver, 10).until(
        EC.visibility_of_element_located((By.ID, f"{base_id}DateEditorButtonOk"))
    )


def _selecionar_mes_no_editor(driver: WebDriver, base_id: str, competencia: CompetenciaTrabalho) -> None:
    """Clica no mês desejado na grade do editor do calendário RichFaces."""

    indice_mes = competencia.mes - 1
    id_mes = f"{base_id}DateEditorLayoutM{indice_mes}"
    _clicar_por_id(driver, id_mes)


def _selecionar_ano_no_editor(driver: WebDriver, base_id: str, competencia: CompetenciaTrabalho) -> None:
    """Navega pelos anos do editor do calendário RichFaces e clica no ano desejado."""

    for _ in range(20):
        anos_visiveis: list[tuple[int, WebElement]] = []
        for indice in range(10):
            id_ano = f"{base_id}DateEditorLayoutY{indice}"
            try:
                el = driver.find_element(By.ID, id_ano)
                if not el.is_displayed():
                    continue
                texto = el.text.strip()
                if texto.isdigit():
                    anos_visiveis.append((int(texto), el))
            except NoSuchElementException:
                continue

        for ano_visivel, el_ano in anos_visiveis:
            if ano_visivel == competencia.ano:
                el_ano.click()
                return

        if not anos_visiveis:
            raise RuntimeError(
                f"Não consegui ler os anos visíveis no calendário de competência ({base_id})."
            )

        menor_ano = min(a for a, _ in anos_visiveis)
        maior_ano = max(a for a, _ in anos_visiveis)

        # Navega para esquerda ou direita conforme o ano-alvo
        if competencia.ano < menor_ano:
            sinal = "<"
        else:
            sinal = ">"

        try:
            # Seletor por atributo para evitar o erro de CSS com ':' no ID
            botoes = driver.find_elements(By.CSS_SELECTOR, f"[id='{base_id}'] .rich-calendar-editor-btn")
            clicou = False
            for btn in botoes:
                if btn.text.strip() == sinal and btn.is_displayed():
                    btn.click()
                    clicou = True
                    break
            if not clicou:
                raise RuntimeError("Botão de navegação do calendário não encontrado.")
        except Exception as exc:
            raise RuntimeError(
                f"Não consegui navegar no calendário para chegar ao ano {competencia.ano}."
            ) from exc

        time.sleep(0.25)

    raise RuntimeError(
        f"Não consegui localizar o ano {competencia.ano} no calendário de competência ({base_id})."
    )


def _confirmar_editor_calendario(driver: WebDriver, base_id: str, esperado: str) -> None:
    """Clica em OK no editor do calendário e confirma que o campo assumiu o valor esperado."""

    _clicar_por_id(driver, f"{base_id}DateEditorButtonOk")
    # Aguarda o campo de data refletir o valor confirmado
    WebDriverWait(driver, 10).until(
        lambda d: (
            lambda el: el is not None and el.get_attribute("value") == esperado
        )(d.find_element(By.ID, f"{base_id}InputDate"))
    )


def _selecionar_campo_competencia(
    driver: WebDriver,
    base_id: str,
    rotulo: str,
    competencia: CompetenciaTrabalho,
) -> None:
    """Seleciona o mês e o ano no calendário RichFaces de um campo de competência."""

    esperado = f"{competencia.mes:02d}/{competencia.ano}"
    valor_atual = _valor_campo_competencia(driver, base_id)

    if valor_atual == esperado:
        registrar_evento_execucao(
            f"Campo {rotulo} da competência já estava em {esperado}.",
            "ISS Fortaleza",
        )
        return

    _abrir_editor_calendario(driver, base_id, rotulo)
    _selecionar_mes_no_editor(driver, base_id, competencia)
    _selecionar_ano_no_editor(driver, base_id, competencia)
    _confirmar_editor_calendario(driver, base_id, esperado)
    registrar_evento_execucao(
        f"Campo {rotulo} da competência ajustado para {esperado}.",
        "ISS Fortaleza",
    )


def selecionar_competencia_na_tela_richfaces(
    driver: WebDriver,
    competencia: CompetenciaTrabalho,
) -> None:
    """Seleciona a competência nos calendários `De` e `Até` usando o editor do RichFaces."""

    # Garante que a tela carregou antes de tentar acessar os calendários
    WebDriverWait(driver, 120).until(
        EC.element_to_be_clickable((By.ID, "manterEscrituracaoForm:btnConsultar"))
    )
    time.sleep(1.2)

    _selecionar_campo_competencia(driver, "manterEscrituracaoForm:dataInicial", "De", competencia)
    _selecionar_campo_competencia(driver, "manterEscrituracaoForm:dataFinal", "Até", competencia)
    registrar_evento_execucao(
        f"Competência selecionada na tela: {competencia.rotulo}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Botão Escriturar
# ---------------------------------------------------------------------------


def aguardar_botao_escriturar(driver: WebDriver, timeout: int = 120) -> bool:
    """Espera a consulta retornar e verifica se o botão de escriturar está habilitado."""

    id_ativo = "manterEscrituracaoForm:dataTable:0:linkEscriturar"
    id_desabilitado = "manterEscrituracaoForm:dataTable:0:linkEscriturarDesabilitado"

    prazo = time.monotonic() + timeout
    while time.monotonic() < prazo:
        try:
            el = driver.find_element(By.ID, id_ativo)
            if el.is_displayed():
                return True
        except NoSuchElementException:
            pass
        try:
            el = driver.find_element(By.ID, id_desabilitado)
            if el.is_displayed():
                return False
        except NoSuchElementException:
            pass
        time.sleep(0.5)

    raise CompetenciaSemEscriturarDisponivelError(
        "Não consegui identificar se a competência consultada liberou o botão de escriturar."
    )


# ---------------------------------------------------------------------------
# Tela de Escrituração Fiscal
# ---------------------------------------------------------------------------


def aguardar_tela_escrituracao_fiscal(driver: WebDriver, timeout: int = 120) -> None:
    """Espera a tela de Escrituração Fiscal ficar pronta após o clique em Escriturar."""

    prazo = time.monotonic() + timeout
    while time.monotonic() < prazo:
        try:
            el = driver.find_element(By.ID, "tabEncerramentoEscrituracao")
            if el.is_displayed():
                return
        except NoSuchElementException:
            pass
        try:
            el = driver.find_element(By.ID, "aba_tomados_lbl")
            if el.is_displayed():
                return
        except NoSuchElementException:
            pass
        time.sleep(0.25)

    raise RuntimeError(
        "A tela de Escrituração Fiscal não ficou visível após o clique em Escriturar."
    )


# ---------------------------------------------------------------------------
# Aba Serviços Tomados
# ---------------------------------------------------------------------------


def clicar_aba_servicos_tomados(driver: WebDriver) -> None:
    """Clica na aba 'Serviços Tomados' com estratégia específica para o portal RichFaces."""

    # Aguarda a tela de Escrituração Fiscal aparecer primeiro
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "tabEncerramentoEscrituracao"))
    )
    try:
        _clicar_por_id(driver, "aba_tomados_lbl", timeout=120)
        return
    except Exception:
        pass

    # Fallback: busca o elemento pelo onclick
    try:
        el = driver.find_element(By.CSS_SELECTOR, '[onclick*="atualizarDadosDaAbaTomados"]')
        el.click()
    except NoSuchElementException:
        raise RuntimeError("Não foi possível clicar na aba 'Serviços Tomados'.")


# ---------------------------------------------------------------------------
# Tela Digitar Documento
# ---------------------------------------------------------------------------


def aguardar_tela_digitar_documento(driver: WebDriver, timeout: int = 120) -> None:
    """Espera a tela de Digitar Documento carregar por completo."""

    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm"))
    )
    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:divPesquisaTomador"))
    )
    # Aguarda o rádio de CNPJ aparecer como confirmação final de carregamento
    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located(
            (By.ID, "digitarDocumentoForm:tipoPesquisaTomadorRb:1")
        )
    )


# ---------------------------------------------------------------------------
# Preenchimento manual dos dados do prestador
# ---------------------------------------------------------------------------


def preencher_dados_prestador(
    driver: WebDriver,
    documento: "DocumentoPortalISS",
    depuracao: bool = False,
) -> None:
    """Preenche manualmente o formulário de identificação do prestador com os dados da planilha.

    Este fluxo substitui a busca via autocomplete pelo CNPJ, permitindo escriturar
    prestadores que ainda não estão cadastrados no portal da ISS Fortaleza.
    Os dados utilizados são os extraídos pela IA durante o processamento das notas fiscais.
    """

    registrar_evento_execucao(
        f"Preenchendo dados do prestador manualmente para CNPJ {documento.cnpj_prestador}",
        "ISS Fortaleza",
    )
    aguardar_tela_digitar_documento(driver)

    # --- Tipo de Cliente/Fornecedor: "Pessoa Física" ou "Pessoa Jurídica" conforme o prestador ---
    eh_cpf = _eh_cpf_prestador(documento)
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherTipoNaturezaJuridica",
        ["Pessoa Física", "Pessoa Fisica"] if eh_cpf else ["Pessoa Jurídica", "Pessoa Juridica", "PJ"],
        timeout=10,
    )
    registrar_evento_execucao(
        f"Tipo Natureza Jurídica definido como {'Pessoa Física' if eh_cpf else 'Pessoa Jurídica'} "
        f"(prestador identificado por {'CPF' if eh_cpf else 'CNPJ'})",
        "ISS Fortaleza",
    )
    time.sleep(0.5)

    # --- CPF/CNPJ do Prestador ---
    cnpj_limpo = limpar_cnpj_para_digitacao(documento.cnpj_prestador)
    if cnpj_limpo:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idCPFCNPJ",
            cnpj_limpo,
            delay_ms=60,
            normalizador_fallback=limpar_cnpj_para_digitacao,
        )
        registrar_evento_execucao(f"CPF/CNPJ preenchido: {cnpj_limpo}", "ISS Fortaleza")

    # --- Nome/Denominação ---
    nome = documento.nome_prestador.strip()
    if nome:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idNome",
            nome,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Nome/Denominação preenchido: {nome}", "ISS Fortaleza")

    # --- UF do Prestador (dispara AJAX de carregamento de cidades) ---
    uf = documento.uf_prestador.strip().upper()
    if uf:
        _selecionar_opcao_por_texto(
            driver,
            By.ID,
            "digitarDocumentoForm:comboEscolherEstado",
            [uf],
            timeout=10,
        )
        registrar_evento_execucao(f"UF do prestador selecionada: {uf}", "ISS Fortaleza")
        # Aguarda o AJAX recarregar o combo de cidades após a seleção de estado
        time.sleep(1.5)

    # --- Cidade do Prestador ---
    cidade = documento.cidade_prestador.strip()
    if cidade:
        _selecionar_opcao_por_texto(
            driver,
            By.ID,
            "digitarDocumentoForm:comboEscolherCidade",
            [cidade],
            timeout=15,
        )
        registrar_evento_execucao(f"Cidade do prestador selecionada: {cidade}", "ISS Fortaleza")
        # Aguarda o AJAX da Cidade processar antes de digitar o CEP — o combo de
        # cidade também dispara um postback do próprio portal (mesmo padrão de
        # estabilização já usado acima para a UF), e sem essa espera o CEP pode
        # ser digitado com sucesso e depois limpo por uma resposta tardia desse
        # AJAX, sem que nada perceba.
        time.sleep(1.5)

    # --- CEP ---
    cep = re.sub(r"\D+", "", documento.cep_prestador or "")
    if cep:
        _digitar_campo(
            driver,
            By.ID,
            "digitarDocumentoForm:idCEP",
            cep,
            delay_ms=40,
        )

        # Confere de novo após uma pequena espera extra: se uma resposta AJAX
        # ainda mais tardia (da seleção de Cidade) limpar o campo, redigita e
        # registra a correção no log de auditoria.
        time.sleep(0.8)
        try:
            valor_pos_espera = re.sub(r"\D+", "", driver.find_element(By.ID, "digitarDocumentoForm:idCEP").get_attribute("value") or "")
        except NoSuchElementException:
            valor_pos_espera = None
        if valor_pos_espera is not None and valor_pos_espera != cep:
            registrar_evento_execucao(
                f"CEP do prestador foi sobrescrito/limpo pelo portal após a digitação "
                f"(ficou '{valor_pos_espera}'); redigitando '{cep}'.",
                "ISS Fortaleza",
            )
            _digitar_campo(
                driver,
                By.ID,
                "digitarDocumentoForm:idCEP",
                cep,
                delay_ms=40,
            )

        registrar_evento_execucao(f"CEP preenchido: {cep}", "ISS Fortaleza")

    # --- Logradouro ---
    logradouro = documento.logradouro_prestador.strip()
    if logradouro:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idEndereco",
            logradouro,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Logradouro preenchido: {logradouro}", "ISS Fortaleza")

    # --- Número (usa 'S/N' se o campo estiver vazio na planilha) ---
    numero = documento.numero_prestador.strip() or "S/N"
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:idNumero",
        numero,
        delay_ms=40,
    )
    registrar_evento_execucao(f"Número preenchido: {numero}", "ISS Fortaleza")

    # --- Bairro ---
    bairro = documento.bairro_prestador.strip()
    if bairro:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idBairro",
            bairro,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Bairro preenchido: {bairro}", "ISS Fortaleza")

    # --- Email (campo não obrigatório — preenche somente se disponível) ---
    email = documento.email_prestador.strip()
    if email:
        try:
            _definir_campo_rapido(
                driver,
                By.ID,
                "digitarDocumentoForm:inputEmail3",
                email,
                delay_ms=20,
            )
            registrar_evento_execucao(f"E-mail preenchido: {email}", "ISS Fortaleza")
        except Exception as exc_email:
            # E-mail é opcional: registra o aviso mas não interrompe o fluxo
            registrar_evento_execucao(
                f"Aviso: não foi possível preencher o e-mail '{email}': {exc_email}",
                "ISS Fortaleza",
            )

    _pausar_para_depuracao(
        driver,
        depuracao,
        f"Dados do prestador {documento.cnpj_prestador} preenchidos. Confira antes de continuar.",
    )
    registrar_evento_execucao(
        f"Dados do prestador preenchidos com sucesso para CNPJ {documento.cnpj_prestador}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Modal de pesquisa de CNAE
# ---------------------------------------------------------------------------


def _forcar_abertura_modal_cnae(driver: WebDriver) -> None:
    """Chama a API nativa do RichFaces via JavaScript para forçar a abertura do modal."""

    try:
        driver.execute_script("Richfaces.showModalPanel('pesquisarCnaeModal')")
    except JavascriptException:
        # Fallback: chama a função com prefixo alternativo que alguns portais usam
        driver.execute_script("RichFaces.showModalPanel('pesquisarCnaeModal')")


def preencher_modal_pesquisar_cnae(
    driver: WebDriver,
    id_cnae_final: str,
    depuracao: bool = False,
) -> None:
    """Preenche o modal de pesquisa de CNAE com o código final da planilha e dispara a busca.

    Caso o modal não apareça dentro do tempo padrão após o clique, utiliza a API
    nativa do RichFaces via JavaScript para forçá-lo a abrir — estratégia de
    fallback que corrige o timeout observado quando o JSF ainda está reprocessando.
    """

    valor_cnae = limpar_numero_para_digitacao(id_cnae_final)
    registrar_evento_execucao(
        f"Abrindo modal de pesquisa de CNAE para o código {valor_cnae}",
        "ISS Fortaleza",
    )

    # Tenta aguardar o modal ficar visível normalmente
    try:
        WebDriverWait(driver, 8).until(
            EC.visibility_of_element_located(
                (By.ID, "digitarDocumentoForm:pesquisarCnaeModalContainer")
            )
        )
    except TimeoutException:
        # Fallback via JavaScript: força o RichFaces a exibir o modal
        registrar_evento_execucao(
            "Modal de CNAE não abriu sozinho; forçando via JavaScript do RichFaces.",
            "ISS Fortaleza",
        )
        _forcar_abertura_modal_cnae(driver)
        # Aguarda mais tempo após a chamada forçada
        WebDriverWait(driver, 10).until(
            EC.visibility_of_element_located(
                (By.ID, "digitarDocumentoForm:pesquisarCnaeModalContainer")
            )
        )

    # Preenche o campo de pesquisa dentro do modal
    _definir_campo_rapido(
        driver,
        By.ID,
        "digitarDocumentoForm:idFormularioPesquisaCnae:idCnaePesquisa",
        valor_cnae,
        delay_ms=45,
    )
    _pausar_para_depuracao(
        driver,
        depuracao,
        f"O modal de pesquisa de CNAE está preenchido com {valor_cnae}.",
    )

    # Dispara a pesquisa AJAX do CNAE escolhido
    _clicar_por_id(driver, "digitarDocumentoForm:idFormularioPesquisaCnae:idPesquisar")
    time.sleep(1.2)
    registrar_evento_execucao(
        f"Pesquisa de CNAE enviada para {valor_cnae}",
        "ISS Fortaleza",
    )
    _pausar_para_depuracao(
        driver,
        depuracao,
        f"A pesquisa de CNAE para {valor_cnae} foi enviada. Confirme o resultado na lista.",
    )

    # Seleciona a primeira linha de resultado retornada pelo portal
    # O ID do link de seleção usa um índice dinâmico; buscamos com seletor parcial
    seletor_resultado = (
        "[id^='digitarDocumentoForm:idFormularioPesquisaCnae:idDatatableListaCnae:0:']"
    )
    el_resultado = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, seletor_resultado))
    )
    el_resultado.click()
    time.sleep(1.2)
    registrar_evento_execucao(
        f"CNAE selecionado na lista de resultados: {valor_cnae}",
        "ISS Fortaleza",
    )


def _digitar_campo_se_nao_zero(
    driver: WebDriver,
    seletor_id: str,
    valor: str,
) -> None:
    """Preenche o campo apenas se o valor for numérico e diferente de zero."""
    if not valor:
        return

    # Remove tudo que não for dígito para checar se é zero
    digitos = re.sub(r"\D+", "", valor)
    if not digitos or int(digitos) == 0:
        return

    _digitar_campo(
        driver,
        By.ID,
        seletor_id,
        limpar_valor_para_digitacao(valor),
    )


# ---------------------------------------------------------------------------
# Preenchimento da aba Serviço
# ---------------------------------------------------------------------------


def preencher_documento_servico(
    driver: WebDriver,
    documento: DocumentoPortalISS,
    depuracao: bool = False,
) -> None:
    """Preenche a aba de serviço com os dados de uma linha da planilha."""

    # Ativa a aba Serviço real do portal antes de preencher os campos
    _clicar_por_id(driver, "digitarDocumentoForm:abaServico_lbl")
    registrar_evento_execucao("Aba Serviço acionada", "ISS Fortaleza")
    time.sleep(0.5)

    # Aguarda os campos da aba Serviço ficarem disponíveis
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm"))
    )
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:tipoDocumentoDigitado"))
    )
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:statusNfse"))
    )
    registrar_evento_execucao(
        "Formulário da aba Serviço ficou visível e pronto para preenchimento",
        "ISS Fortaleza",
    )

    registrar_evento_execucao(
        f"Iniciando preenchimento da aba Serviço para a NF {documento.numero_nf} ({documento.cnpj_prestador})",
        "ISS Fortaleza",
    )

    # Tipo de documento — MEI estabelecido em Fortaleza/CE (a exceção que faz a nota
    # ser escriturada apesar de ser da própria capital, ver regra em executar_fluxo_iss)
    # usa "NFS-e Nacional" em vez do "NFS-e de Outro Município" padrão.
    cidade_prest_doc = normalizar_texto(documento.cidade_prestador).upper()
    uf_prest_doc = normalizar_texto(documento.uf_prestador).upper()
    is_fortaleza_ce_doc = (uf_prest_doc == "CE" and "FORTALEZA" in cidade_prest_doc)
    is_mei_doc = (documento.regime_tributario.upper().strip() == "MEI")

    if is_mei_doc and is_fortaleza_ce_doc:
        opcoes_tipo_documento = ["NFS-e Nacional", "NFS-e nacional"]
        registrar_evento_execucao(
            f"Prestador MEI estabelecido em Fortaleza/CE (NF {documento.numero_nf}): "
            "Tipo do Documento Digitado definido como 'NFS-e Nacional'.",
            "ISS Fortaleza",
        )
    elif _eh_cpf_prestador(documento):
        opcoes_tipo_documento = ["NFS Avulsa de outro município", "NFS Avulsa de outro municipio"]
        registrar_evento_execucao(
            f"Prestador pessoa física (NF {documento.numero_nf}): "
            "Tipo do Documento Digitado definido como 'NFS Avulsa de outro município'.",
            "ISS Fortaleza",
        )
    else:
        opcoes_tipo_documento = ["NFS-e de Outro Município", "NFS-e de outro município", "707"]

    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:tipoDocumentoDigitado",
        opcoes_tipo_documento,
    )

    # Número da nota fiscal
    # Usa digitação real (não _definir_campo_rapido): este campo é seguido de perto
    # por um select (Status NFSE) que dispara reprocessamento AJAX do JSF — sem
    # foco/clique real no campo, o valor definido via JS pode nunca ser confirmado
    # pelo lado servidor e acaba sendo limpo quando esse AJAX reprocessa o painel.
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:numeroDocumentoDigitado",
        limpar_numero_para_digitacao(documento.numero_nf),
    )

    # Data de emissão
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:dataEmissaoInputDate",
        documento.data_emissao,
    )

    # Status NFSE
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:statusNfse",
        ["NORMAL", "Normal", "472"],
    )
    registrar_evento_execucao(
        f"Campos iniciais da aba Serviço preenchidos para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Pausa de estabilização para o JSF reprocessar após selecionar o Status
    time.sleep(1.2)

    # Confere se o reprocessamento do JSF não limpou o Número da NF (mesma classe
    # de corrida já vista para Natureza/ISS Retido e Cidade/CEP nesta sessão)
    numero_nf_limpo = limpar_numero_para_digitacao(documento.numero_nf)
    try:
        valor_pos_status = driver.find_element(By.ID, "digitarDocumentoForm:numeroDocumentoDigitado").get_attribute("value") or ""
    except NoSuchElementException:
        valor_pos_status = None
    if valor_pos_status is not None and _somente_digitos(valor_pos_status) != numero_nf_limpo:
        registrar_evento_execucao(
            f"Número da NF foi limpo pelo portal após selecionar Status NFSE "
            f"(ficou '{valor_pos_status}'); redigitando '{numero_nf_limpo}'.",
            "ISS Fortaleza",
        )
        _digitar_campo(driver, By.ID, "digitarDocumentoForm:numeroDocumentoDigitado", numero_nf_limpo, delay_ms=40)

    # Abre o modal de pesquisa de CNAE
    _clicar_por_id(driver, "digitarDocumentoForm:idLinkPesquisarCnae")
    registrar_evento_execucao(
        f"Botão Pesquisar CNAE acionado para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Preenche o modal de CNAE (com fallback via JavaScript)
    preencher_modal_pesquisar_cnae(driver, documento.id_cnae_final, depuracao=depuracao)

    # --- Alíquota (bloqueada pelo portal para alguns CNAEs, calculada automaticamente) ---
    aliquota = documento.aliquota.strip()
    if aliquota:
        if _campo_aliquota_esta_editavel(driver) is not None:
            _digitar_campo(driver, By.ID, "digitarDocumentoForm:idAliquota", aliquota, delay_ms=40)
            registrar_evento_execucao(f"Alíquota preenchida: {aliquota}", "ISS Fortaleza")
        else:
            registrar_evento_execucao(
                f"Alíquota não preenchida: campo bloqueado pelo portal para este CNAE "
                f"(calculado automaticamente). Valor da planilha: {aliquota}.",
                "ISS Fortaleza",
            )

    # Descrição do serviço preenchida de forma instantânea via JS
    _definir_valor_instantaneo(
        driver,
        By.ID,
        "digitarDocumentoForm:idDescricaoServico",
        documento.descricao_servico,
    )

    # UF de prestação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherEstadoLocalPrestacao",
        [documento.uf_local_prestacao],
    )

    # Cidade de prestação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherCidadeLocalPrestacao",
        [documento.cidade_local_prestacao],
    )

    # Natureza da operação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherLocalPrestacao",
        [documento.natureza_operacao],
    )
    # Aguarda o AJAX da Natureza da Operação processar antes de mexer no checkbox
    # ISS Retido — o portal marca/desmarca esse checkbox sozinho ao mudar para
    # certas naturezas (ex.: "Tributação Fora do Município"), e sem essa espera
    # o código abaixo lê o estado do checkbox antes dessa reação do portal
    # terminar (mesmo padrão de estabilização já usado para Status NFSE e UF do
    # Prestador nesta função).
    time.sleep(1.5)

    # Checkbox ISS retido
    deve_marcar, iss_retido_reconhecido = _interpretar_iss_retido(documento.iss_retido)
    if not iss_retido_reconhecido:
        # Defesa em profundidade: validar_campos_obrigatorios() já deveria ter
        # barrado esta nota antes de chegar aqui, mas registra o aviso mesmo assim.
        registrar_evento_execucao(
            f"AVISO: valor de ISS_RETIDO '{documento.iss_retido}' não reconhecido para "
            f"a NF {documento.numero_nf}. Tratado como 'Não retido' por padrão.",
            "ISS Fortaleza",
        )
    try:
        seletor_checkbox_iss = (By.NAME, "digitarDocumentoForm:j_id361")

        # A essa altura já esperamos 1.5s pelo AJAX da Natureza da Operação
        # (acima), então esta leitura reflete o estado já assentado do portal.
        marcado_agora = driver.find_element(*seletor_checkbox_iss).is_selected()

        if deve_marcar != marcado_agora:
            driver.find_element(*seletor_checkbox_iss).click()

        # Confere de novo após uma pequena espera extra: se uma resposta AJAX
        # ainda mais tardia sobrescrever o clique, corrige e registra a
        # correção no log de auditoria.
        estado_pos_clique = _ler_estado_checkbox_apos_espera(driver, *seletor_checkbox_iss, espera=1.0)
        if estado_pos_clique is not None and estado_pos_clique != deve_marcar:
            registrar_evento_execucao(
                f"Checkbox ISS Retido da NF {documento.numero_nf} foi sobrescrito pelo "
                f"portal logo após o clique (ficou "
                f"{'marcado' if estado_pos_clique else 'desmarcado'}); corrigindo para "
                f"{'marcado' if deve_marcar else 'desmarcado'}.",
                "ISS Fortaleza",
            )
            driver.find_element(*seletor_checkbox_iss).click()

        estado_final = driver.find_element(*seletor_checkbox_iss).is_selected()
        if estado_final != deve_marcar:
            registrar_evento_execucao(
                f"ALERTA: checkbox ISS Retido da NF {documento.numero_nf} ficou "
                f"{'marcado' if estado_final else 'desmarcado'}, mas o esperado era "
                f"{'marcado' if deve_marcar else 'desmarcado'} "
                f"(valor na planilha: '{documento.iss_retido}').",
                "ISS Fortaleza",
            )
        else:
            registrar_evento_execucao(
                f"Checkbox ISS Retido da NF {documento.numero_nf} definido como "
                f"{'marcado' if estado_final else 'desmarcado'} "
                f"(valor na planilha: '{documento.iss_retido}').",
                "ISS Fortaleza",
            )
    except NoSuchElementException:
        registrar_evento_execucao(
            f"ERRO: elemento do checkbox 'ISS Retido' (digitarDocumentoForm:j_id361) não "
            f"foi encontrado na tela para a NF {documento.numero_nf}. O estado do ISS "
            "Retido NÃO pôde ser verificado/definido nesta nota — revise manualmente no portal.",
            "ISS Fortaleza",
        )
    registrar_evento_execucao(
        f"Local de prestação e ISS tratados para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Valor do serviço e retenções/deduções (somente preenche se não for zero)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idValorServicoPrestado", documento.valor_servico)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idValorDeducoes", documento.valor_deducoes)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idDescontosIncondicionados", documento.descontos_incondicionados)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idDescontosCondicionados", documento.descontos_condicionados)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idOutrasRetencoes", documento.outras_retencoes)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idIR", documento.ir)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idPis", documento.pis_nao_retido)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idConfins", documento.cofins_nao_retido)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idCSLL", documento.csrf)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idINSS", documento.inss)

    registrar_evento_execucao(
        f"Aba Serviço preenchida com sucesso para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Fluxo principal da ISS
# ---------------------------------------------------------------------------


def resetar_tela_para_digitar_documento(driver: WebDriver, competencia: CompetenciaTrabalho) -> None:
    """Tenta navegar de volta para a tela inicial de Digitar Documento caso a gravação de uma nota tenha falhado
    e o formulário tenha ficado travado em estado intermediário (efeito cascata)."""
    registrar_evento_execucao(
        "Iniciando recuperação de estado do portal para limpar formulário travado...",
        "ISS Fortaleza",
    )
    # 1. Clica no menu topo Escrituração
    try:
        el_menu = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable(
                (By.XPATH, "//a[contains(@class, 'dropdown-toggle') and (contains(., 'Escritura') or contains(., 'Escrituracao'))]")
            )
        )
        el_menu.click()
        time.sleep(0.3)
        _clicar_por_id(driver, "formMenuTopo:menuEscrituracao:j_id80")
    except Exception:
        # Fallback de menu por índice
        menus = driver.find_elements(By.CSS_SELECTOR, "a.dropdown-toggle")
        if len(menus) > 4:
            menus[4].click()
            time.sleep(0.3)
            _clicar_por_id(driver, "formMenuTopo:menuEscrituracao:j_id80")
            
    # 2. Aguarda a tela de Manter Escrituração e consulta a competência novamente
    WebDriverWait(driver, 30).until(
        EC.element_to_be_clickable((By.ID, "manterEscrituracaoForm:btnConsultar"))
    )
    selecionar_competencia_na_tela_richfaces(driver, competencia)
    _clicar_por_id(driver, "manterEscrituracaoForm:btnConsultar")
    
    # 3. Aguarda o botão Escriturar e clica
    aguardar_botao_escriturar(driver)
    _clicar_por_id(driver, "manterEscrituracaoForm:dataTable:0:linkEscriturar")
    
    # 4. Vai para a aba Serviços Tomados e clica em Digitar Documento
    clicar_aba_servicos_tomados(driver)
    WebDriverWait(driver, 30).until(
        EC.visibility_of_element_located((By.ID, "servico_tomado_form:seamj_id849"))
    )
    _clicar_por_id(driver, "servico_tomado_form:seamj_id849")
    aguardar_tela_digitar_documento(driver)
    time.sleep(0.5)



def executar_fluxo_iss(
    documentos: list[DocumentoPortalISS],
    competencia: CompetenciaTrabalho,
    caminho_log: Path,
    depuracao: bool = False,
    aguardar_login_por_arquivo: bool = False,
    caminho_confirmacao_login: Path | None = None,
    reutilizar_navegador: bool = False,
    perfil_navegador: Path = PERFIL_CHROME_FUNCAO2_PADRAO,
    callback_progresso: Any = None,
    executando_em_gui: bool = False,
    callback_pausa: Any = None,
    callback_cancelamento: Any = None,
) -> None:
    """Executa o fluxo visual completo do portal da ISS até o preenchimento do formulário."""

    global _CALLBACK_PAUSA, _CALLBACK_CANCELAMENTO

    # Protege a escrita do callback global com o mesmo Lock usado na leitura (FALHA-04)
    # BUG CORRIGIDO: sem o 'global' acima, esta atribuição criava uma variável LOCAL a
    # esta função (sombreando o nome do módulo), e o global de verdade lido por
    # _verificar_pausa()/_verificar_cancelamento() nunca era atualizado — a pausa nos
    # helpers de baixo nível (_clicar_com_espera, _digitar_campo, etc.) nunca chegava
    # a funcionar de fato, apesar do comentário abaixo dizer o contrário.
    with _LOCK_CALLBACK_PAUSA:
        _CALLBACK_PAUSA = callback_pausa
    with _LOCK_CALLBACK_CANCELAMENTO:
        _CALLBACK_CANCELAMENTO = callback_cancelamento

    if not documentos:
        raise RuntimeError("A planilha não contém linhas válidas para a automação.")

    # Resolve o perfil do Chrome para o AppData do usuário se estiver usando o padrão relativo,
    # garantindo que o perfil persistente seja acessível após a instalação do executável.
    if reutilizar_navegador and perfil_navegador == PERFIL_CHROME_FUNCAO2_PADRAO:
        appdata = os.environ.get("APPDATA", str(Path.home()))
        perfil_navegador = Path(appdata) / "BeAContab" / "brain" / "navegador_funcao2_profile"

    # Abre o Chrome com ou sem perfil persistente, conforme solicitado
    if reutilizar_navegador:
        driver = abrir_navegador_com_perfil_persistente(perfil_navegador, depuracao=depuracao)
    else:
        driver = abrir_navegador_visivel(depuracao=depuracao)

    cancelado_pelo_usuario = False
    try:
        registrar_evento_execucao(
            f"Fluxo da ISS iniciado com {len(documentos)} documento(s) e competência {competencia.rotulo}",
            "ISS Fortaleza",
        )

        # Aguarda login manual do operador
        if aguardar_login_por_arquivo:
            # Resolve o caminho de confirmacao usando AppData para evitar falhas no instalador
            if caminho_confirmacao_login:
                caminho = caminho_confirmacao_login
            else:
                appdata = os.environ.get("APPDATA", str(Path.home()))
                caminho = Path(appdata) / "BeAContab" / "brain" / "funcao2_login_ok.flag"
            aguardar_login_manual_por_arquivo(driver, caminho)
        else:
            aguardar_login_manual(driver)
        registrar_evento_execucao("Login manual confirmado pelo usuário", "ISS Fortaleza")

        print("Login confirmado. Aguardando o menu do portal ficar disponível...")
        try:
            WebDriverWait(driver, 120).until(
                EC.presence_of_element_located(
                    (By.XPATH, "//*[contains(text(), 'Escrituração') or contains(text(), 'Escrituracao')]")
                )
            )
        except TimeoutException:
            registrar_evento_execucao(
                "Timeout ao aguardar menu Escrituração visível", "ISS Fortaleza"
            )
            print("AVISO: Timeout ao aguardar menu Escrituração. Prosseguindo mesmo assim...")
        _pausar_para_depuracao(
            driver,
            depuracao,
            "O portal foi carregado e o login já foi confirmado.",
        )

        # Navega para Escrituração > Manter Escrituração
        print("Acessando Escrituração > Manter Escrituração...")
        try:
            # Busca o menu Escrituração pelo texto para ser resiliente a mudanças na ordem dos menus
            el_menu = WebDriverWait(driver, 10).until(
                EC.element_to_be_clickable(
                    (By.XPATH, "//a[contains(@class, 'dropdown-toggle') and (contains(., 'Escritura') or contains(., 'Escrituracao'))]")
                )
            )
            el_menu.click()
        except Exception:
            # Fallback por índice para compatibilidade com variações do portal
            menus = driver.find_elements(By.CSS_SELECTOR, "a.dropdown-toggle")
            if len(menus) > 4:
                menus[4].click()
        time.sleep(0.5)
        _clicar_por_id(driver, "formMenuTopo:menuEscrituracao:j_id80")
        registrar_evento_execucao("Menu Escrituração acionado", "ISS Fortaleza")

        WebDriverWait(driver, 120).until(
            EC.element_to_be_clickable((By.ID, "manterEscrituracaoForm:btnConsultar"))
        )
        registrar_evento_execucao("Tela Manter Escrituração aberta", "ISS Fortaleza")

        # Seleciona a competência nos calendários De e Até
        print(
            f"Selecionando a competência {competencia.nome_mes} {competencia.ano} na tela de manutenção..."
        )
        selecionar_competencia_na_tela_richfaces(driver, competencia)

        # Aciona a consulta da competência
        print("Consultando a competência selecionada...")
        _clicar_por_id(driver, "manterEscrituracaoForm:btnConsultar")
        registrar_evento_execucao("Botão Consultar acionado", "ISS Fortaleza")

        # Loop para lidar com competências sem botão Escriturar habilitado
        while True:
            botao_escriturar_habilitado = aguardar_botao_escriturar(driver)
            if botao_escriturar_habilitado:
                break

            mensagem_sem_escriturar = (
                f"A competência {competencia.rotulo} está desabilitada para escriturar. "
                "Escolha outro mês para continuar a automação."
            )
            print(mensagem_sem_escriturar)
            registrar_evento_execucao(mensagem_sem_escriturar, "ISS Fortaleza")

            if not sys.stdin.isatty():
                raise CompetenciaSemEscriturarDisponivelError(mensagem_sem_escriturar)

            competencia = solicitar_nova_competencia_para_repetir(competencia)
            print(
                f"Reiniciando a seleção da competência {competencia.nome_mes} {competencia.ano} "
                "a partir dos campos De e Até..."
            )
            registrar_evento_execucao(
                f"Nova competência informada pelo usuário para retry: {competencia.rotulo}",
                "ISS Fortaleza",
            )
            selecionar_competencia_na_tela_richfaces(driver, competencia)
            _clicar_por_id(driver, "manterEscrituracaoForm:btnConsultar")
            registrar_evento_execucao("Botão Consultar acionado", "ISS Fortaleza")

        # Abre o formulário de escrituração
        print("Abrindo a rotina de escrituração...")
        _clicar_por_id(driver, "manterEscrituracaoForm:dataTable:0:linkEscriturar")
        registrar_evento_execucao("Botão Escriturar acionado", "ISS Fortaleza")
        aguardar_tela_escrituracao_fiscal(driver)
        registrar_evento_execucao("Tela Escrituração Fiscal aberta", "ISS Fortaleza")

        # Clica na aba Serviços Tomados
        print("Selecionando a aba Serviços Tomados...")
        clicar_aba_servicos_tomados(driver)
        registrar_evento_execucao("Aba Serviços Tomados acionada", "ISS Fortaleza")

        # Aguarda e clica no botão Digitar Documento
        WebDriverWait(driver, 120).until(
            EC.visibility_of_element_located((By.ID, "servico_tomado_form:seamj_id849"))
        )
        print("Abrindo Digitar Documento...")
        _clicar_por_id(driver, "servico_tomado_form:seamj_id849")
        registrar_evento_execucao("Tela Digitar Documento aberta", "ISS Fortaleza")
        aguardar_tela_digitar_documento(driver)
        time.sleep(0.5)
        _pausar_para_depuracao(
            driver,
            depuracao,
            "A tela Digitar Documento está aberta e pronta para o próximo clique.",
        )

        print(
            f"A planilha possui {len(documentos)} linha(s) válida(s). O fluxo iniciará o processamento em lote."
        )

        pasta_log = caminho_log.parent
        if pasta_log.name != "log":
            pasta_log = pasta_log / "log"
        pasta_log.mkdir(parents=True, exist_ok=True)
        
        # Limpa arquivos de logs anteriores para esta execução
        for nome in ["log_escrituradas_sucesso.txt", "log_prefeitura_fortaleza.txt", "log_notas_incompletas.txt"]:
            caminho_antigo = pasta_log / nome
            if caminho_antigo.exists():
                try:
                    caminho_antigo.unlink()
                except Exception:
                    pass

        configurar_pasta_logs(pasta_log)
        notas_sucesso: list[str] = []

        for indice, candidato in enumerate(documentos, start=1):
            _verificar_cancelamento()
            if callback_progresso:
                try:
                    callback_progresso(indice, len(documentos))
                except Exception:
                    pass
            
            # Regra: Se o prestador for de Fortaleza/CE, ignora a nota, EXCETO se for MEI
            cidade_prest = normalizar_texto(candidato.cidade_prestador).upper()
            uf_prest = normalizar_texto(candidato.uf_prestador).upper()
            is_fortaleza_ce = (uf_prest == "CE" and "FORTALEZA" in cidade_prest)
            is_mei = (candidato.regime_tributario.upper().strip() == "MEI")
            
            if is_fortaleza_ce and not is_mei:
                mensagem_log = (
                    f"ARQUIVO_PDF={candidato.arquivo_pdf} | "
                    f"CNPJ_PRESTADOR={candidato.cnpj_prestador} | "
                    f"IGNORADO: Prestador estabelecido em Fortaleza/CE (não MEI)"
                )
                registrar_log_funcao2(caminho_log, mensagem_log)
                print(f"Linha {indice}/{len(documentos)} ignorada: Prestador de Fortaleza/CE e não é MEI.")
                
                # Grava no log de Fortaleza/CE
                caminho_fortaleza = pasta_log / "log_prefeitura_fortaleza.txt"
                try:
                    if not caminho_fortaleza.exists():
                        with open(caminho_fortaleza, "w", encoding="utf-8") as f:
                            f.write("As seguintes notas fiscais foram ignoradas por terem sido emitidas por prestadores estabelecidos em Fortaleza/CE (não MEI):\n\n")
                    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    with open(caminho_fortaleza, "a", encoding="utf-8") as f:
                        f.write(f"[{timestamp}] {candidato.arquivo_pdf} - CNPJ: {candidato.cnpj_prestador} - NF: {candidato.numero_nf} - Cidade: {candidato.cidade_prestador}\n")
                except Exception as exc_fort:
                    print(f"Erro ao registrar log de Fortaleza: {exc_fort}")
                continue

            # Valida se há alguma informação obrigatória vazia/ausente na planilha antes de prosseguir
            campos_ausentes = validar_campos_obrigatorios(candidato)
            if campos_ausentes:
                motivo = f"Campos obrigatórios ausentes na planilha: {', '.join(campos_ausentes)}"
                mensagem_log = (
                    f"ARQUIVO_PDF={candidato.arquivo_pdf} | "
                    f"CNPJ_PRESTADOR={candidato.cnpj_prestador} | "
                    f"IGNORADO: {motivo}"
                )
                registrar_log_funcao2(caminho_log, mensagem_log)
                print(f"Linha {indice}/{len(documentos)} ignorada. Motivo: {motivo}")
                
                # Grava no log exclusivo de notas incompletas
                caminho_incompletas = pasta_log / "log_notas_incompletas.txt"
                try:
                    if not caminho_incompletas.exists():
                        with open(caminho_incompletas, "w", encoding="utf-8") as f:
                            f.write("As seguintes notas fiscais foram ignoradas por estarem com campos obrigatórios ausentes na planilha:\n\n")
                    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                    with open(caminho_incompletas, "a", encoding="utf-8") as f:
                        f.write(f"[{timestamp}] PDF: {candidato.arquivo_pdf} | CNPJ: {candidato.cnpj_prestador} | NF: {candidato.numero_nf}\n")
                        f.write(f"  -> Ausente(s): {', '.join(campos_ausentes)}\n\n")
                except Exception as exc_inc:
                    print(f"Erro ao registrar log de notas incompletas: {exc_inc}")
                continue


            print(
                f"Processando linha {indice}/{len(documentos)}: "
                f"{candidato.cnpj_prestador} - NF {candidato.numero_nf}"
            )
            try:
                # Preenche os dados do prestador manualmente com as informações da planilha,
                # independentemente de ele estar ou não cadastrado no portal da ISS Fortaleza.
                preencher_dados_prestador(driver, candidato, depuracao=depuracao)
            except Exception as exc_prestador:
                if isinstance(exc_prestador, AutomacaoCanceladaError):
                    raise
                mensagem_log = (
                    f"ARQUIVO_PDF={candidato.arquivo_pdf} | "
                    f"CNPJ_PRESTADOR={candidato.cnpj_prestador} | "
                    f"ERRO ao preencher dados do prestador: {exc_prestador}"
                )
                registrar_log_funcao2(caminho_log, mensagem_log)
                print(f"Erro ao preencher dados do prestador {candidato.cnpj_prestador}: {exc_prestador}; avançando.")
                continue

            try:
                preencher_documento_servico(driver, candidato, depuracao=depuracao)
                
                # Gravar documento (resolve sozinho os modais de confirmação que o
                # portal pode encadear: nota duplicada -> "Não" e desiste da nota;
                # avisos informativos como "Prestador não inscrito no CPOM" -> "Sim"
                # e tenta gravar de novo, até um limite de tentativas)
                print("Gravando documento no portal...")
                registrar_evento_execucao(f"Gravando NF {candidato.numero_nf}", "ISS Fortaleza")
                resultado_gravacao = _gravar_documento_com_confirmacoes(driver, candidato.numero_nf, timeout=15)

                if resultado_gravacao == "duplicata":
                    nf_info = f"{candidato.arquivo_pdf} - {candidato.cnpj_prestador} - {candidato.numero_nf}"
                    print(f"Linha {indice}/{len(documentos)} ignorada: portal indicou nota fiscal já escriturada (mesmo CNPJ e número de nota).")

                    mensagem_log = (
                        f"ARQUIVO_PDF={candidato.arquivo_pdf} | "
                        f"CNPJ_PRESTADOR={candidato.cnpj_prestador} | "
                        f"IGNORADO: Portal indicou documento fiscal já escriturado com o mesmo "
                        f"CNPJ e número de nota, em outra competência"
                    )
                    registrar_log_funcao2(caminho_log, mensagem_log)

                    # Grava no log exclusivo de notas duplicadas
                    caminho_duplicadas = pasta_log / "log_notas_duplicadas.txt"
                    try:
                        if not caminho_duplicadas.exists():
                            with open(caminho_duplicadas, "w", encoding="utf-8") as f:
                                f.write("As seguintes notas fiscais foram ignoradas porque o portal indicou que já existe documento fiscal escriturado com o mesmo CNPJ e número de nota, em outra competência:\n\n")
                        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        with open(caminho_duplicadas, "a", encoding="utf-8") as f:
                            f.write(f"[{timestamp}] {nf_info}\n")
                    except Exception as exc_dup:
                        print(f"Erro ao registrar log de notas duplicadas: {exc_dup}")
                else:
                    registrar_evento_execucao(f"Sucesso na gravação da NF {candidato.numero_nf}", "ISS Fortaleza")
                    nf_info = f"{candidato.arquivo_pdf} - {candidato.cnpj_prestador} - {candidato.numero_nf}"
                    notas_sucesso.append(nf_info)

                    # Grava no log de sucesso em tempo real
                    caminho_sucesso = pasta_log / "log_escrituradas_sucesso.txt"
                    try:
                        if not caminho_sucesso.exists():
                            with open(caminho_sucesso, "w", encoding="utf-8") as f:
                                f.write("As seguintes notas fiscais foram escrituradas e gravadas com sucesso no portal:\n\n")
                        with open(caminho_sucesso, "a", encoding="utf-8") as f:
                            f.write(f"- {nf_info}\n")
                    except Exception as exc_suc:
                        print(f"Erro ao salvar log de sucesso em tempo real: {exc_suc}")

            except Exception as exc:
                if isinstance(exc, AutomacaoCanceladaError):
                    raise
                registrar_evento_execucao(f"Falha ao processar NF {candidato.numero_nf}: {exc}", "ISS Fortaleza")
                print(f"Erro ao processar NF {candidato.numero_nf}: {exc}")
                registrar_log_funcao2(caminho_log, f"ERRO ao processar {candidato.arquivo_pdf}: {exc}")

            finally:
                # Clica em Digitar novo documento para limpar o formulário para a próxima linha
                try:
                    _clicar_com_espera(driver, By.XPATH, '//*[@id="j_id165:novo"]', "botão Novo Documento", timeout=15)
                    time.sleep(2)
                except Exception as exc:
                    if isinstance(exc, AutomacaoCanceladaError):
                        raise
                    print(f"Não foi possível clicar em 'Novo documento': {exc}. Realizando reset preventivo de tela.")
                    try:
                        resetar_tela_para_digitar_documento(driver, competencia)
                    except Exception as exc_reset:
                        print(f"Falha crítica no reset de tela: {exc_reset}. O fluxo pode falhar na próxima iteração.")

        print()
        print("Gravação de documentos em lote concluída.")
        if not executando_em_gui:
            print("Revise a tela no navegador e, se quiser encerrar, volte ao terminal.")
            input("Pressione Enter para encerrar esta sessão automatizada e manter a tela aberta...")
        else:
            print("Execução da GUI concluída. O navegador permanecerá aberto.")

    except AutomacaoCanceladaError:
        cancelado_pelo_usuario = True
        registrar_evento_execucao(
            "Automação cancelada pelo operador. Encerrando o navegador para permitir reinício limpo.",
            "ISS Fortaleza",
        )
        raise

    finally:
        # Mantém o navegador aberto em modo de sessão persistente ou quando executado via GUI,
        # facilitando o teste assistido e a análise em caso de problemas — EXCETO quando o
        # operador cancelou a automação, caso em que o navegador é sempre fechado para que
        # a próxima execução comece do zero, com login manual novo.
        if cancelado_pelo_usuario:
            try:
                driver.quit()
            except Exception:
                pass
        elif not reutilizar_navegador and not executando_em_gui:
            try:
                driver.quit()
            except Exception:
                pass
        else:
            registrar_evento_execucao(
                "Sessão persistente da função 2 ou execução pela GUI preservada para análise/reutilização",
                "ISS Fortaleza",
            )


# ---------------------------------------------------------------------------
# Ponto de entrada síncrono exposto para a CLI
# ---------------------------------------------------------------------------


def executar_automacao_iss(
    caminho_xlsx: Path,
    competencia: CompetenciaTrabalho,
    depuracao: bool = False,
    usar_inspector: bool = False,
    aguardar_login_por_arquivo: bool = False,
    caminho_confirmacao_login: Path | None = None,
    reutilizar_navegador: bool = False,
    perfil_navegador: Path = PERFIL_CHROME_FUNCAO2_PADRAO,
    porta_debug_navegador: int = PORTA_DEBUG_CHROME_PADRAO,
    callback_progresso: Any = None,
    executando_em_gui: bool = False,
    callback_pausa: Any = None,
    callback_cancelamento: Any = None,
) -> None:
    """Ponto de entrada síncrono para a opção 2 da CLI."""

    registrar_evento_execucao(
        f"Automação da ISS solicitada com planilha {caminho_xlsx.resolve()} e competência {competencia.rotulo}",
        "ISS Fortaleza",
    )
    documentos = carregar_documentos_xlsx(caminho_xlsx)
    caminho_log = caminho_xlsx.parent / "log" / f"{caminho_xlsx.stem}_log_funcao2.txt"
    executar_fluxo_iss(
        documentos,
        competencia,
        caminho_log,
        depuracao=depuracao,
        aguardar_login_por_arquivo=aguardar_login_por_arquivo,
        caminho_confirmacao_login=caminho_confirmacao_login,
        reutilizar_navegador=reutilizar_navegador,
        perfil_navegador=perfil_navegador,
        callback_progresso=callback_progresso,
        executando_em_gui=executando_em_gui,
        callback_pausa=callback_pausa,
        callback_cancelamento=callback_cancelamento,
    )


# ---------------------------------------------------------------------------
# Extração avulsa do portal (funcionalidade auxiliar mantida)
# ---------------------------------------------------------------------------


def _limpar_texto_exibido(texto: str | None) -> str:
    """Normaliza espaços para deixar a extração mais legível no XLSX."""

    return " ".join((texto or "").split()).strip()


def extrair_dados_da_pagina_iss(driver: WebDriver) -> dict[str, object]:
    """Coleta o conteúdo visível da tela logada para exportação e análise posterior."""

    titulo = driver.title
    url = driver.current_url

    try:
        texto_bruto = driver.find_element(By.TAG_NAME, "body").text
    except Exception:
        texto_bruto = ""

    dados: dict[str, object] = {
        "titulo": titulo,
        "url": url,
        "extraido_em": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "texto_visivel": [linha.strip() for linha in texto_bruto.splitlines() if linha.strip()],
        "elementos": _coletar_elementos_visiveis(driver),
    }
    return dados


def _coletar_elementos_visiveis(driver: WebDriver) -> list[dict[str, str]]:
    """Captura campos e botões visíveis para documentar a tela atual do portal."""

    elementos: list[dict[str, str]] = []
    for tag in ("input", "select", "textarea", "button"):
        for el in driver.find_elements(By.TAG_NAME, tag):
            try:
                if not el.is_displayed():
                    continue
                elementos.append({
                    "tag": tag,
                    "id": el.get_attribute("id") or "",
                    "name": el.get_attribute("name") or "",
                    "type": el.get_attribute("type") or "",
                    "value": el.get_attribute("value") or "",
                    "texto": el.text or "",
                })
            except StaleElementReferenceException:
                continue
    return elementos


def _aplicar_cabecalho(aba) -> None:
    """Padroniza o estilo visual das planilhas geradas pela extração."""

    for celula in aba[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="1F4E78")


def ajustar_largura_colunas(aba) -> None:
    """Ajusta a largura das colunas da aba para facilitar a leitura."""

    for coluna in aba.columns:
        largura_max = max(
            (len(str(celula.value)) if celula.value else 0) for celula in coluna
        )
        aba.column_dimensions[coluna[0].column_letter].width = min(largura_max + 4, 80)


def exportar_extracao_iss(dados: dict[str, object], destino: Path) -> None:
    """Exporta a página extraída para XLSX."""

    wb = Workbook()

    aba_resumo = wb.active
    aba_resumo.title = "Resumo"
    aba_resumo.append(["Campo", "Valor"])
    aba_resumo.append(["Título", dados.get("titulo", "")])
    aba_resumo.append(["URL", dados.get("url", "")])
    aba_resumo.append(["Extraído em", dados.get("extraido_em", "")])
    aba_resumo.append(["Linhas de texto", len(dados.get("texto_visivel", []))])
    aba_resumo.append(["Elementos", len(dados.get("elementos", []))])
    _aplicar_cabecalho(aba_resumo)

    aba_texto = wb.create_sheet("Texto")
    aba_texto.append(["Linha", "Conteúdo"])
    for indice, linha in enumerate(dados.get("texto_visivel", []), start=1):
        aba_texto.append([indice, linha])
    _aplicar_cabecalho(aba_texto)

    aba_elementos = wb.create_sheet("Elementos")
    aba_elementos.append(["Tag", "ID", "Name", "Type", "Value", "Texto"])
    for el in dados.get("elementos", []):
        aba_elementos.append([
            el.get("tag", ""),
            el.get("id", ""),
            el.get("name", ""),
            el.get("type", ""),
            el.get("value", ""),
            el.get("texto", ""),
        ])
    _aplicar_cabecalho(aba_elementos)

    for aba in wb.worksheets:
        ajustar_largura_colunas(aba)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def executar_extracao_portal_iss(destino_xlsx: Path = Path("iss_extracao.xlsx")) -> None:
    """Ponto de entrada síncrono para extrair o conteúdo visível do portal logado."""

    registrar_evento_execucao(
        f"Extração avulsa do portal ISS iniciada com destino {destino_xlsx.resolve()}",
        "ISS Fortaleza",
    )
    driver = abrir_navegador_visivel()
    try:
        aguardar_login_manual(driver)

        print()
        print("Quando estiver na tela exata da qual deseja extrair os dados, volte ao terminal.")
        input("Pressione Enter para capturar a página atual...")

        time.sleep(1)
        dados = extrair_dados_da_pagina_iss(driver)
        exportar_extracao_iss(dados, destino_xlsx)
        registrar_evento_execucao(
            f"Extração avulsa concluída com sucesso em {destino_xlsx.resolve()}",
            "ISS Fortaleza",
        )

        print(f"Extração concluída com sucesso: {destino_xlsx.resolve()}")
        input("Pressione Enter para encerrar esta sessão automatizada e fechar o navegador...")
    finally:
        try:
            driver.quit()
        except Exception:
            pass
