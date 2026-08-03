"""Módulo de processamento de arquivos XML de NFS-e com classificação de CNAE via Tess AI."""

from __future__ import annotations

import json
import os
import re
import shutil
import xml.etree.ElementTree as ET
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
import requests

from cnae_final import (
    _classificar_localmente,
    _normalizar_codigo_cnae,
    _normalizar_texto,
    _selecionar_candidatos,
    carregar_cnaes_oficiais,
)

# Mapeamento do prefixo do código do município IBGE (2 dígitos) para a UF correspondente.
MAPA_IBGE_UF = {
    "11": "RO", "12": "AC", "13": "AM", "14": "RR", "15": "PA", "16": "AP", "17": "TO",
    "21": "MA", "22": "PI", "23": "CE", "24": "RN", "25": "PB", "26": "PE", "27": "AL", "28": "SE", "29": "BA",
    "31": "MG", "32": "ES", "33": "RJ", "35": "SP",
    "41": "PR", "42": "SC", "43": "RS",
    "50": "MS", "51": "MT", "52": "GO", "53": "DF"
}

def formatar_monetario(valor_str: str) -> str:
    """Formata uma string numérica (ex: 28600.00) para o padrão brasileiro (ex: 28.600,00)."""
    if not valor_str:
        return "0,00"
    try:
        val = float(valor_str)
        return f"{val:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    except ValueError:
        return valor_str

def obter_texto_tag(parent: ET.Element, tag_name: str, padrao: str = "") -> str:
    """Busca um elemento de forma recursiva ignorando namespaces e retorna seu valor textual."""
    elem = parent.find(f".//{{*}}{tag_name}")
    if elem is not None and elem.text:
        return elem.text.strip()
    return padrao

def resolver_cnae_tess_ai(
    desc_cnae: str,
    descricao_servico: str = "",
    id_cnae: str = "",
    tess_key: str = "",
    tess_agent_id: str = "",
    caminho_oficial: str = "cnae_oficial.xlsx"
) -> tuple[str, str]:
    """Consulta o Tess AI para selecionar o melhor CNAE oficial da planilha oficial."""
    
    chave = tess_key or os.getenv("TESS_API_KEY", "").strip()
    agente = tess_agent_id or os.getenv("TESS_AGENT_ID", "").strip()
    
    # Validação crítica: caso não existam credenciais do Tess AI, utiliza o fallback local
    if not chave or not agente:
        print("[Tess AI] Chaves ou ID de agente ausentes para classificação. Utilizando fallback local.", flush=True)
        try:
            oficiais = carregar_cnaes_oficiais(caminho_oficial)
            return _classificar_localmente(desc_cnae, descricao_servico, id_cnae, oficiais)
        except Exception:
            return id_cnae.strip(), desc_cnae.strip()
            
    try:
        oficiais = carregar_cnaes_oficiais(caminho_oficial)
    except FileNotFoundError:
        return id_cnae.strip(), desc_cnae.strip()
        
    candidatos = _selecionar_candidatos(desc_cnae, descricao_servico, id_cnae, oficiais)
    if not candidatos:
        return id_cnae.strip(), desc_cnae.strip()
        
    # Construção do prompt alinhado para que a IA escolha exclusivamente uma opção do cnae_oficial.xlsx
    prompt = (
        "Você é um especialista em classificação de atividades econômicas (CNAE).\n"
        "Sua tarefa é selecionar exatamente um único código e descrição da lista de candidatos oficiais de 'cnae_oficial.xlsx' que represente a melhor opção para a NFS-e fornecida.\n\n"
        "Dados extraídos do XML da nota fiscal:\n"
        f"- ID CNAE Extraído: {id_cnae}\n"
        f"- Descrição CNAE Extraída: {desc_cnae}\n"
        f"- Discriminação do Serviço: {descricao_servico}\n\n"
        "Lista de opções válidas em 'cnae_oficial.xlsx' (Candidatos):\n"
        f"{json.dumps([{'codigo': c.codigo, 'descricao': c.descricao} for c in candidatos], ensure_ascii=False, indent=2)}\n\n"
        "Escolha a melhor correspondência oficial da lista de candidatos acima. Não crie novos códigos.\n"
        "Retorne EXCLUSIVAMENTE um objeto JSON estruturado no formato abaixo, sem tags markdown ou textos adicionais:\n"
        "{\n"
        '  "id_cnae_final": "código do item escolhido da lista",\n'
        '  "desc_cnae_final": "descrição oficial exata do item escolhido"\n'
        "}"
    )
    
    url = f"https://api.tess.im/agents/{agente}/execute"
    headers = {
        "Authorization": f"Bearer {chave}",
        "Content-Type": "application/json",
    }
    payload = {
        "temperature": "0",
        "messages": [
            {
                "role": "user",
                "content": prompt,
            },
        ],
        "waitExecution": True,
    }
    
    try:
        response = requests.post(url, json=payload, headers=headers, timeout=60)
        response.raise_for_status()
        dados_resp = response.json()
        texto_resposta = dados_resp["responses"][0]["output"]
        
        if texto_resposta:
            limpo = texto_resposta.strip()
            # Limpa decorações de markdown (```json ... ```) se houver
            if limpo.startswith("```"):
                limpo = re.sub(r"^```[a-zA-Z]*\n", "", limpo)
                limpo = re.sub(r"\n```$", "", limpo)
                limpo = limpo.strip()
                
            resultado = json.loads(limpo)
            id_cnae_final = str(resultado.get("id_cnae_final", "")).strip()
            desc_cnae_final = str(resultado.get("desc_cnae_final", "")).strip()
            if id_cnae_final and desc_cnae_final:
                return id_cnae_final, desc_cnae_final
    except Exception as e:
        print(f"[Tess AI] Erro ao classificar CNAE: {e}. Executando fallback local.", flush=True)
        
    return _classificar_localmente(desc_cnae, descricao_servico, id_cnae, oficiais)

def extrair_dados_xml(caminho_xml: Path) -> dict[str, Any] | None:
    """Realiza o parse de um XML de NFS-e e retorna um dicionário mapeado com os campos fiscais."""
    try:
        tree = ET.parse(caminho_xml)
        root = tree.getroot()
    except Exception as e:
        print(f"Erro ao parsear arquivo XML {caminho_xml.name}: {e}", flush=True)
        return None
        
    # Garante que é um XML de nota fiscal válido (deve conter a tag infNFSe ou similar)
    infNFSe = root.find(".//{*}infNFSe")
    if infNFSe is None:
        return None
        
    # 1. Informações básicas da nota
    numero_nf = obter_texto_tag(root, "nNFSe")
    
    # 2. Data de emissão (limpa timestamp e formata para DD/MM/AAAA)
    data_br = ""
    dhEmi = obter_texto_tag(root, "dhEmi")
    if not dhEmi:
        dhEmi = obter_texto_tag(root, "dhProc")
    if dhEmi:
        try:
            # Pega apenas a data antes do 'T'
            data_iso = dhEmi.split("T")[0]
            dt = datetime.strptime(data_iso, "%Y-%m-%d")
            data_br = dt.strftime("%d/%m/%Y")
        except Exception:
            data_br = dhEmi
            
    # 3. Prefeitura emissora
    xLocEmi = obter_texto_tag(root, "xLocEmi")
    prefeitura = f"PREFEITURA MUNICIPAL DE {xLocEmi.upper()}" if xLocEmi else ""
    
    # 4. Dados do Prestador (emitente)
    cnpj_prestador = ""
    nome_prestador = ""
    logradouro_prestador = ""
    numero_prestador = ""
    bairro_prestador = ""
    uf_prestador = ""
    cep_prestador = ""
    email_prestador = ""
    
    emit = root.find(".//{*}emit")
    if emit is not None:
        cnpj_prestador = obter_texto_tag(emit, "CNPJ")
        nome_prestador = obter_texto_tag(emit, "xNome")
        
        ender = emit.find(".//{*}enderNac")
        if ender is not None:
            logradouro_prestador = obter_texto_tag(ender, "xLgr")
            numero_prestador = obter_texto_tag(ender, "nro")
            bairro_prestador = obter_texto_tag(ender, "xBairro")
            uf_prestador = obter_texto_tag(ender, "UF")
            cep_prestador = obter_texto_tag(ender, "CEP")
            
    prest = root.find(".//{*}prest")
    if prest is not None:
        email_prestador = obter_texto_tag(prest, "email")
        if not cnpj_prestador:
            cnpj_prestador = obter_texto_tag(prest, "CNPJ")
        if not nome_prestador:
            nome_prestador = obter_texto_tag(prest, "xNome")

    # 5. Dados do Serviço e Local de Prestação
    # Busca de forma ampla no XML para maior resiliência em relação à estrutura
    id_cnae = obter_texto_tag(root, "cTribNac")
    if not id_cnae:
        id_cnae = obter_texto_tag(root, "cTribMun")
    if not id_cnae:
        id_cnae = obter_texto_tag(root, "cServ")
        
    desc_cnae = obter_texto_tag(root, "xTribNac")
    if not desc_cnae:
        desc_cnae = obter_texto_tag(root, "xTribMun")
        
    descricao_servico = obter_texto_tag(root, "xDescServ")
    cidade_local_prestacao = obter_texto_tag(root, "xLocPrestacao")
    uf_local_prestacao = ""
    
    serv = root.find(".//{*}serv")
    if serv is not None:
        cLocPrestacao = obter_texto_tag(serv, "cLocPrestacao")
        if cLocPrestacao and len(cLocPrestacao) >= 2:
            uf_local_prestacao = MAPA_IBGE_UF.get(cLocPrestacao[:2], "")

    # Se faltar local de prestacao e o emitente for o mesmo, assume o do emitente
    if not cidade_local_prestacao:
        cidade_local_prestacao = xLocEmi
    if not uf_local_prestacao:
        uf_local_prestacao = uf_prestador

    # 6. Valores da Nota
    # Tenta obter vServ (dentro do DPS/valores/vServPrest/vServ) ou vBC/vLiq da NFS-e superior
    vServ = obter_texto_tag(root, "vServ")
    if not vServ:
        vServ = obter_texto_tag(root, "vBC")
    if not vServ:
        vServ = obter_texto_tag(root, "vLiq")
        
    valor_servico = formatar_monetario(vServ) if vServ else "0,00"
    
    # Alíquota e Retenções
    pAliqAplic = obter_texto_tag(root, "pAliqAplic")
    if not pAliqAplic:
        pAliqAplic = obter_texto_tag(root, "pAliq")
    aliquota = formatar_monetario(pAliqAplic) if pAliqAplic else "0,00"
    
    vTotalRet = obter_texto_tag(root, "vTotalRet")
    vTotalRet_val = 0.0
    if vTotalRet:
        try:
            vTotalRet_val = float(vTotalRet)
        except ValueError:
            pass

    # Determinação da retenção do ISS (tpRetISSQN: 1=Não, 2=Sim)
    tpRetISSQN = obter_texto_tag(root, "tpRetISSQN")
    iss_retido = "NÃO"
    if tpRetISSQN == "2":
        iss_retido = "SIM"
    elif tpRetISSQN == "1":
        iss_retido = "NÃO"
    else:
        # Fallback caso a tag tpRetISSQN não esteja presente
        if vTotalRet_val > 0.0:
            iss_retido = "SIM"

    # Natureza de operação: "Tributação no Município" se cidade for FORTALEZA e UF for CE, senão "Tributação Fora do Município"
    cidade_norm = _normalizar_texto(cidade_local_prestacao or "")
    uf_norm = (uf_local_prestacao or "").strip().upper()
    if "fortaleza" in cidade_norm and uf_norm == "CE":
        natureza_operacao = "Tributação no Município"
    else:
        natureza_operacao = "Tributação Fora do Município"

    # 7. Tributos Federais e Deduções
    ir = "0,00"
    inss = "0,00"
    csrf = "0,00"
    pis_nao_retido = "0,00"
    cofins_nao_retido = "0,00"
    valor_deducoes = "0,00"
    descontos_incondicionados = "0,00"
    descontos_condicionados = "0,00"
    outras_retencoes = "0,00"

    vRetIRRF = obter_texto_tag(root, "vRetIRRF")
    if vRetIRRF:
        ir = formatar_monetario(vRetIRRF)
        
    vRetCP = obter_texto_tag(root, "vRetCP")
    if vRetCP:
        inss = formatar_monetario(vRetCP)
        
    # CSRF é a somatória das retenções de CSLL, PIS e COFINS
    vRetCSLL = obter_texto_tag(root, "vRetCSLL")
    vRetPIS = obter_texto_tag(root, "vRetPIS")
    vRetCofins = obter_texto_tag(root, "vRetCofins")
    
    soma_csrf = 0.0
    for v in [vRetCSLL, vRetPIS, vRetCofins]:
        if v:
            try:
                soma_csrf += float(v)
            except ValueError:
                pass
    if soma_csrf > 0.0:
        csrf = formatar_monetario(soma_csrf)

    vPis = obter_texto_tag(root, "vPis")
    if vPis:
        pis_nao_retido = formatar_monetario(vPis)
        
    vCofins = obter_texto_tag(root, "vCofins")
    if vCofins:
        cofins_nao_retido = formatar_monetario(vCofins)

    vDed = obter_texto_tag(root, "vDed")
    if not vDed:
        vDed = obter_texto_tag(root, "vDedRed")
    if vDed:
        valor_deducoes = formatar_monetario(vDed)

    vDescIncond = obter_texto_tag(root, "vDescIncond")
    if vDescIncond:
        descontos_incondicionados = formatar_monetario(vDescIncond)

    vDescCond = obter_texto_tag(root, "vDescCond")
    if vDescCond:
        descontos_condicionados = formatar_monetario(vDescCond)

    opSimpNac = obter_texto_tag(root, "opSimpNac")
    regime_tributario = "MEI" if opSimpNac == "2" else "OUTROS"

    return {
        "arquivo_xml": caminho_xml.name,
        "prefeitura": prefeitura,
        "cnpj_prestador": cnpj_prestador,
        "nome_prestador": nome_prestador,
        "uf_prestador": uf_prestador,
        "cidade_prestador": xLocEmi or cidade_local_prestacao,
        "cep_prestador": cep_prestador,
        "logradouro_prestador": logradouro_prestador,
        "numero_prestador": numero_prestador,
        "bairro_prestador": bairro_prestador,
        "email_prestador": email_prestador,
        "numero_nf": numero_nf,
        "data_emissao": data_br,
        "id_cnae": id_cnae,
        "desc_cnae": desc_cnae,
        "descricao_servico": descricao_servico,
        "uf_local_prestacao": uf_local_prestacao,
        "cidade_local_prestacao": cidade_local_prestacao,
        "natureza_operacao": natureza_operacao,
        "iss_retido": iss_retido,
        "valor_servico": valor_servico,
        "valor_deducoes": valor_deducoes,
        "descontos_incondicionados": descontos_incondicionados,
        "descontos_condicionados": descontos_condicionados,
        "outras_retencoes": outras_retencoes,
        "ir": ir,
        "pis_nao_retido": pis_nao_retido,
        "cofins_nao_retido": cofins_nao_retido,
        "csrf": csrf,
        "inss": inss,
        "aliquota": aliquota,
        "regime_tributario": regime_tributario,
    }

def processar_pasta_xmls(
    pasta_origem: str,
    tess_key: str = "",
    tess_agent_id: str = "",
    caminho_oficial_cnae: str = "cnae_oficial.xlsx",
    modelo_planilha: str = "EXEMPLOS/planilha exemplo.xlsx",
    callback_log = None,
    callback_progresso = None,
    checkpoint_a_cada: int = 50,
    arquivos_xml_forcados: list = None,
) -> str:
    """Escaneia a pasta recursivamente procurando XMLs, classifica com Tess AI e gera a planilha de saída.

    Parâmetros:
        checkpoint_a_cada: A planilha é salva parcialmente a cada N notas processadas,
                           evitando perda total de dados em caso de falha (MELHORIA-03).
        arquivos_xml_forcados: Quando informado, usa esta lista de arquivos em vez de
                               escanear `pasta_origem` recursivamente. Usado por
                               `processar_pasta_xmls_auto` para gerar uma planilha só
                               com os XMLs diretos de uma pasta, sem descer para dentro
                               de subpastas que serão processadas separadamente.
    """

    def log(msg: str, is_error: bool = False) -> None:
        if callback_log:
            callback_log(msg, is_error)
        else:
            print(f"{'[ERRO] ' if is_error else ''}{msg}", flush=True)

    caminho_origem = Path(pasta_origem)
    if not caminho_origem.exists() or not caminho_origem.is_dir():
        log(f"Pasta de origem inválida: {pasta_origem}", True)
        raise FileNotFoundError(f"Pasta de origem {pasta_origem} não encontrada.")

    # Localiza arquivos XML: usa a lista explícita se fornecida, senão escaneia
    # a pasta recursivamente (comportamento padrão, inalterado).
    arquivos_xml = list(arquivos_xml_forcados) if arquivos_xml_forcados is not None else list(caminho_origem.rglob("*.xml"))
    if not arquivos_xml:
        log("Nenhum arquivo XML encontrado na pasta selecionada.", True)
        # Exibe estado de erro na barra de progresso, não "Concluído!" (FALHA-05)
        if callback_progresso:
            callback_progresso(0, "Erro: nenhum XML encontrado.")
        raise FileNotFoundError("Nenhum arquivo XML encontrado para processamento.")

    log(f"Total de XMLs localizados para análise: {len(arquivos_xml)}")

    # Define o arquivo XLSX resultante no diretório de origem selecionado
    caminho_destino = caminho_origem / "notas_xml_processadas.xlsx"

    # Copia a planilha exemplo modelo para servir de base estrutural exata
    path_modelo = Path(modelo_planilha)
    # Se o modelo relativo não for achado, tenta no diretório de recursos do
    # executável empacotado (PyInstaller) e, por fim, no diretório do script
    if not path_modelo.exists():
        import sys
        if hasattr(sys, "_MEIPASS"):
            path_modelo = Path(sys._MEIPASS) / modelo_planilha
        else:
            path_modelo = Path(__file__).parent / modelo_planilha

    if not path_modelo.exists():
        log(f"Planilha de exemplo modelo não encontrada em {modelo_planilha}.", True)
        raise FileNotFoundError(f"Modelo {modelo_planilha} ausente.")

    log(f"Copiando modelo de estrutura para o destino: {caminho_destino.name}")
    shutil.copy(path_modelo, caminho_destino)

    # Abre a planilha de destino usando openpyxl
    wb = openpyxl.load_workbook(caminho_destino)
    ws = wb.active

    # Se a segunda coluna for "PREFEITURA", apaga-a dinamicamente para alinhar com o novo modelo
    if ws.max_column >= 2:
        col_b_val = str(ws.cell(row=1, column=2).value or "").strip().upper()
        if col_b_val == "PREFEITURA":
            ws.delete_cols(2)

    # Adiciona o cabeçalho REGIME_TRIBUTARIO no final da linha 1, somente se ainda não existir (FALHA-02)
    cabecalhos_existentes = [
        str(ws.cell(row=1, column=c).value or "").strip().upper()
        for c in range(1, ws.max_column + 1)
    ]
    if "REGIME_TRIBUTARIO" not in cabecalhos_existentes:
        ws.cell(row=1, column=ws.max_column + 1).value = "REGIME_TRIBUTARIO"

    # Garante a limpeza de possíveis linhas remanescentes abaixo do cabeçalho
    if ws.max_row > 1:
        for r in range(ws.max_row, 1, -1):
            ws.delete_rows(r)

    processadas = 0
    ignoradas = 0
    total = len(arquivos_xml)
    cache_cnae: dict[tuple[str, str, str], tuple[str, str]] = {}

    for idx, xml_file in enumerate(arquivos_xml, start=1):
        # Emite progresso em tempo real
        pct = int(((idx - 1) / total) * 100)
        if callback_progresso:
            callback_progresso(pct, f"Processando XML {idx} de {total}")

        # Extração de campos estruturados do XML
        dados = extrair_dados_xml(xml_file)
        if dados is None:
            ignoradas += 1
            log(f"[{idx}/{total}] Ignorado: {xml_file.name} (não é uma NFS-e válida)")
            continue

        log(f"[{idx}/{total}] Processando NFS-e {dados['numero_nf']} do prestador {dados['nome_prestador']}")

        # Classificação CNAE final com IA do Tess (com cache local por prestador + serviço)
        id_cnae_final = ""
        desc_cnae_final = ""
        
        cnae_orig = (dados["id_cnae"] or "").strip()
        desc_serv_norm = " ".join((dados["descricao_servico"] or "").split()).lower()
        cnpj_prest = (dados["cnpj_prestador"] or "").strip()
        chave_cache = (cnpj_prest, cnae_orig, desc_serv_norm)
        
        if chave_cache in cache_cnae:
            id_cnae_final, desc_cnae_final = cache_cnae[chave_cache]
            log(f"[{idx}/{total}] Reutilizando classificação de CNAE do cache local para o prestador {dados['nome_prestador']}")
        else:
            try:
                id_cnae_final, desc_cnae_final = resolver_cnae_tess_ai(
                    desc_cnae=dados["desc_cnae"],
                    descricao_servico=dados["descricao_servico"],
                    id_cnae=dados["id_cnae"],
                    tess_key=tess_key,
                    tess_agent_id=tess_agent_id,
                    caminho_oficial=caminho_oficial_cnae
                )
                if id_cnae_final and desc_cnae_final:
                    cache_cnae[chave_cache] = (id_cnae_final, desc_cnae_final)
            except Exception as err:
                log(f"Erro na IA do Tess ao classificar CNAE para a nota {dados['numero_nf']}: {err}", True)

        dados["id_cnae_final"] = id_cnae_final
        dados["desc_cnae_final"] = desc_cnae_final

        # Regra de negócio: Se o CNAE final for um dos IDs de eventos, força ISS retido como SIM
        id_cnae_final_norm = _normalizar_codigo_cnae(id_cnae_final)
        if id_cnae_final_norm in {"932989910", "900190201"}:
            dados["iss_retido"] = "SIM"

        # Organiza a linha no formato sequencial da planilha exemplo, sem a coluna PREFEITURA
        linha = [
            dados["arquivo_xml"],
            dados["cnpj_prestador"],
            dados["nome_prestador"],
            dados["uf_prestador"],
            dados["cidade_prestador"],
            dados["cep_prestador"],
            dados["logradouro_prestador"],
            dados["numero_prestador"],
            dados["bairro_prestador"],
            dados["email_prestador"],
            dados["numero_nf"],
            dados["data_emissao"],
            dados["id_cnae"],
            dados["desc_cnae"],
            dados["descricao_servico"],
            dados["uf_local_prestacao"],
            dados["cidade_local_prestacao"],
            dados["natureza_operacao"],
            dados["iss_retido"],
            dados["valor_servico"],
            dados["valor_deducoes"],
            dados["descontos_incondicionados"],
            dados["descontos_condicionados"],
            dados["outras_retencoes"],
            dados["ir"],
            dados["pis_nao_retido"],
            dados["cofins_nao_retido"],
            dados["csrf"],
            dados["inss"],
            dados["aliquota"],
            dados["id_cnae_final"],
            dados["desc_cnae_final"],
            dados["regime_tributario"],
        ]

        ws.append(linha)
        processadas += 1

        # Checkpoint: salva a planilha parcialmente a cada N notas para não perder progresso (MELHORIA-03)
        if checkpoint_a_cada > 0 and processadas % checkpoint_a_cada == 0:
            try:
                wb.save(caminho_destino)
                log(f"Checkpoint: planilha salva com {processadas} nota(s) até agora.")
            except Exception as e_ckpt:
                log(f"Aviso: falha ao salvar checkpoint ({e_ckpt}).", True)

    wb.save(caminho_destino)
    wb.close()

    if callback_progresso:
        callback_progresso(100, "Concluído!")

    msg_conclusao = f"Concluído! {processadas} nota(s) processada(s) e gravada(s) em {caminho_destino.name}. Ignorados: {ignoradas} arquivo(s)."
    log(msg_conclusao)

    return str(caminho_destino)


def processar_pasta_xmls_auto(
    pasta_origem: str,
    tess_key: str = "",
    tess_agent_id: str = "",
    caminho_oficial_cnae: str = "cnae_oficial.xlsx",
    modelo_planilha: str = "EXEMPLOS/planilha exemplo.xlsx",
    callback_log = None,
    callback_progresso = None,
    checkpoint_a_cada: int = 50,
) -> list[str]:
    """Detecta automaticamente a estrutura de `pasta_origem` e gera uma planilha
    para cada "lote" de XMLs encontrado, sem nunca misturar XMLs de locais diferentes
    num único arquivo:

    - Se houver .xml diretamente em `pasta_origem` (checagem NÃO recursiva): gera uma
      planilha só com esses arquivos, na própria pasta selecionada.
    - Cada subpasta imediata que contenha XML (em qualquer profundidade dentro dela)
      gera sua própria "notas_xml_processadas.xlsx", independente das demais.

    Os dois casos acima não são mutuamente exclusivos: uma pasta com XML solto na
    raiz E subpastas gera 1 planilha para os arquivos soltos + 1 planilha por
    subpasta — nunca uma única planilha misturando tudo via varredura recursiva.

    Retorna a lista de planilhas geradas.
    """

    def log(msg: str, is_error: bool = False) -> None:
        if callback_log:
            callback_log(msg, is_error)
        else:
            print(f"{'[ERRO] ' if is_error else ''}{msg}", flush=True)

    caminho_origem = Path(pasta_origem)
    if not caminho_origem.exists() or not caminho_origem.is_dir():
        log(f"Pasta de origem inválida: {pasta_origem}", True)
        raise FileNotFoundError(f"Pasta de origem {pasta_origem} não encontrada.")

    xmls_diretos = list(caminho_origem.glob("*.xml"))  # não recursivo, só a pasta selecionada
    subpastas = sorted(p for p in caminho_origem.iterdir() if p.is_dir())

    if not xmls_diretos and not subpastas:
        msg = "Nenhum arquivo XML encontrado na pasta selecionada e nenhuma subpasta existe."
        log(msg, True)
        if callback_progresso:
            callback_progresso(0, "Erro: nenhum XML encontrado.")
        raise FileNotFoundError(msg)

    resultados: list[str] = []
    subpastas_vazias: list[str] = []

    # Lote 1: XMLs soltos diretamente na pasta selecionada (se houver), escopado
    # apenas a esses arquivos via arquivos_xml_forcados — NÃO usa rglob aqui, para
    # não puxar também o conteúdo das subpastas (que são tratadas à parte abaixo).
    if xmls_diretos:
        log(f"Processando {len(xmls_diretos)} XML(s) diretamente na pasta selecionada.")
        resultado = processar_pasta_xmls(
            pasta_origem=pasta_origem,
            tess_key=tess_key,
            tess_agent_id=tess_agent_id,
            caminho_oficial_cnae=caminho_oficial_cnae,
            modelo_planilha=modelo_planilha,
            callback_log=callback_log,
            callback_progresso=callback_progresso,
            checkpoint_a_cada=checkpoint_a_cada,
            arquivos_xml_forcados=xmls_diretos,
        )
        resultados.append(resultado)
        log(f"Planilha da pasta selecionada concluída: {resultado}")

    # Lote 2+: cada subpasta imediata com XML é um lote independente
    total_sub = len(subpastas)
    for idx, subpasta in enumerate(subpastas, start=1):
        # Pré-checagem recursiva só para decidir se a subpasta tem conteúdo e poder
        # pular/logar subpastas vazias sem abortar o lote inteiro; o processamento
        # real reaproveita o rglob interno de processar_pasta_xmls.
        if not list(subpasta.rglob("*.xml")):
            subpastas_vazias.append(subpasta.name)
            log(f"Subpasta '{subpasta.name}' ignorada: nenhum XML encontrado.", True)
            continue

        log(f"Processando subpasta {idx} de {total_sub}: {subpasta.name}")

        def callback_progresso_sub(pct, status, _nome=subpasta.name, _idx=idx, _total=total_sub):
            if callback_progresso:
                callback_progresso(pct, f"[{_idx}/{_total}] {_nome}: {status}")

        resultado = processar_pasta_xmls(
            pasta_origem=str(subpasta),
            tess_key=tess_key,
            tess_agent_id=tess_agent_id,
            caminho_oficial_cnae=caminho_oficial_cnae,
            modelo_planilha=modelo_planilha,
            callback_log=callback_log,
            callback_progresso=callback_progresso_sub,
            checkpoint_a_cada=checkpoint_a_cada,
        )
        resultados.append(resultado)
        log(f"Subpasta '{subpasta.name}' concluída. Planilha: {resultado}")

    if not resultados:
        msg = (
            "Nenhum arquivo XML encontrado em nenhuma subpasta. "
            f"Subpastas vazias: {', '.join(subpastas_vazias)}"
        )
        log(msg, True)
        if callback_progresso:
            callback_progresso(0, "Erro: nenhum XML encontrado.")
        raise FileNotFoundError(msg)

    if callback_progresso:
        callback_progresso(100, "Concluído!")

    return resultados
