"""
Módulo de união de planilhas XLSX.

Consolida múltiplos arquivos Excel (gerados pela extração por prefeitura
ou por IA) em um único arquivo de saída, mantendo a formatação contábil.
"""
from __future__ import annotations

from pathlib import Path
from typing import Callable, Optional

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

# Nome da aba esperada nas planilhas de origem
ABA_PADRAO = "NF Extraídas"

# Colunas que devem receber formatação contábil brasileira
COLUNAS_CONTABEIS = {
    "VALOR_SERVICO", "VALOR_DEDUCOES", "DESCONTOS_INCONDICIONADOS",
    "DESCONTOS_CONDICIONADOS", "OUTRAS_RETENCOES", "IR",
    "PIS_NAO_RETIDO", "COFINS_NAO_RETIDO", "CSRF", "INSS",
}
COLUNAS_PERCENTUAL = {"ALIQUOTA"}


def _aplicar_formato_contabil(cell) -> None:
    """Aplica o formato contábil brasileiro à célula."""
    cell.number_format = '#,##0.00'


def _aplicar_formato_percentual(cell) -> None:
    """Aplica o formato percentual à célula (valor já deve estar dividido por 100)."""
    cell.number_format = '0.00%'


def _valor_para_float(valor: str) -> Optional[float]:
    """
    Converte uma string no formato brasileiro ('1.234,56' ou '12,34%')
    para float. Retorna None em caso de falha.
    """
    if not valor:
        return None
    s = str(valor).strip()
    # Remove símbolo de porcentagem para tratar separadamente
    eh_percentual = s.endswith("%")
    s = s.replace("%", "").strip()
    # Remove pontos de milhar e converte vírgula decimal
    s = s.replace(".", "").replace(",", ".")
    try:
        f = float(s)
        if eh_percentual:
            f = f / 100.0
        return f
    except ValueError:
        return None


def detectar_xlsx_modelo(caminho_dir: Path) -> list[Path]:
    """Identifica recursivamente todos os arquivos .xlsx na pasta que possuem a aba modelo."""
    arquivos_modelo = []
    # Procura arquivos .xlsx de forma recursiva nas subpastas
    for caminho in caminho_dir.glob("**/*.xlsx"):
        if caminho.name.startswith("~$") or caminho.name == "planilha_unificada.xlsx":
            continue
        try:
            wb = load_workbook(str(caminho), read_only=True)
            if ABA_PADRAO in wb.sheetnames:
                arquivos_modelo.append(caminho)
            wb.close()
        except Exception:
            continue
    return sorted(arquivos_modelo)


def unir_xlsx(
    lista_caminhos: list[str | Path],
    caminho_destino: str | Path,
    remover_duplicatas: bool = True,
    callback_log: Optional[Callable[[str], None]] = None,
) -> Path:
    """
    Une múltiplas planilhas XLSX em um único arquivo consolidado.

    Parâmetros:
    - lista_caminhos: lista de caminhos para arquivos .xlsx ou diretórios de origem.
    - caminho_destino: pasta ou caminho completo do arquivo de saída.
    - remover_duplicatas: flag para remover linhas duplicadas idênticas.
    - callback_log: função opcional para emitir logs de progresso.

    Retorna o caminho do arquivo XLSX gerado.
    """
    def log(msg: str) -> None:
        print(msg)
        if callback_log:
            callback_log(msg)

    # Normaliza o destino
    destino = Path(caminho_destino)
    if destino.is_dir():
        destino = destino / "planilha_unificada.xlsx"

    # Resolve os caminhos (se houver pastas, extrai as planilhas modelo delas)
    arquivos_finais: list[Path] = []
    for caminho in lista_caminhos:
        p = Path(caminho)
        if p.is_dir():
            log(f"Escaneando pasta por planilhas de modelo: {p.name}")
            encontrados = detectar_xlsx_modelo(p)
            log(f"  ↳ {len(encontrados)} planilha(s) modelo encontrada(s) em {p.name}")
            arquivos_finais.extend(encontrados)
        elif p.is_file() and p.suffix.lower() == ".xlsx" and p.name != "planilha_unificada.xlsx":
            arquivos_finais.append(p)

    if not arquivos_finais:
        log("[AVISO] Nenhuma planilha modelo válida encontrada para unificação.")
        # Retorna o destino mesmo assim para evitar quebrar o fluxo
        return destino

    # Cria workbook de saída
    wb_saida = Workbook()
    ws_saida = wb_saida.active
    ws_saida.title = ABA_PADRAO

    # Estilo do cabeçalho
    fill_cabecalho = PatternFill(start_color="1F3864", end_color="1F3864", fill_type="solid")
    fonte_cabecalho = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    alinhamento_centro = Alignment(horizontal="center", vertical="center")

    cabecalhos_escritos = False
    indices_contabeis: set[int] = set()
    indices_percentuais: set[int] = set()
    total_linhas = 0
    seen_rows = set()

    for caminho in arquivos_finais:
        log(f"Lendo: {caminho.name}")
        try:
            wb_orig = load_workbook(str(caminho), data_only=True)
        except Exception as e:
            log(f"[ERRO] Não foi possível abrir {caminho.name}: {e}")
            continue

        # Localiza a aba correta
        if ABA_PADRAO in wb_orig.sheetnames:
            ws_orig = wb_orig[ABA_PADRAO]
        else:
            ws_orig = wb_orig.active

        linhas = list(ws_orig.iter_rows(values_only=True))
        if not linhas:
            log(f"[AVISO] {caminho.name} está vazio.")
            continue

        if not cabecalhos_escritos:
            # Escreve o cabeçalho a partir do primeiro arquivo
            cabecalhos = linhas[0]
            ws_saida.append(list(cabecalhos))

            # Identifica colunas contábeis e percentuais pelos nomes do cabeçalho
            for idx, nome_col in enumerate(cabecalhos, start=1):
                nome_upper = str(nome_col).upper().replace(" ", "_") if nome_col else ""
                if nome_upper in COLUNAS_CONTABEIS:
                    indices_contabeis.add(idx)
                if nome_upper in COLUNAS_PERCENTUAL:
                    indices_percentuais.add(idx)

            # Aplica estilos no cabeçalho
            for cell in ws_saida[1]:
                cell.fill = fill_cabecalho
                cell.font = fonte_cabecalho
                cell.alignment = alinhamento_centro

            # Congela o painel no cabeçalho
            ws_saida.freeze_panes = "A2"
            cabecalhos_escritos = True
            dados_inicio = 1
        else:
            dados_inicio = 1

        linhas_importadas = 0
        linhas_duplicadas = 0

        # Copia as linhas de dados
        for linha in linhas[dados_inicio:]:
            if all(v is None or str(v).strip() == "" for v in linha):
                continue  # Pula linhas completamente vazias

            # Detecção de duplicatas (converte todos os campos para string limpa de forma consistente)
            linha_tuple = tuple(str(v).strip() if v is not None else "" for v in linha)
            if remover_duplicatas and linha_tuple in seen_rows:
                linhas_duplicadas += 1
                continue
            seen_rows.add(linha_tuple)

            ws_saida.append(list(linha))
            linhas_importadas += 1
            total_linhas += 1
            linha_atual = ws_saida.max_row

            # Aplica formatação nas células da linha adicionada
            for idx_col in indices_contabeis:
                cell = ws_saida.cell(row=linha_atual, column=idx_col)
                valor_raw = cell.value
                if valor_raw is not None:
                    f = _valor_para_float(str(valor_raw))
                    if f is not None:
                        cell.value = f
                _aplicar_formato_contabil(cell)

            for idx_col in indices_percentuais:
                cell = ws_saida.cell(row=linha_atual, column=idx_col)
                valor_raw = cell.value
                if valor_raw is not None:
                    f = _valor_para_float(str(valor_raw))
                    if f is not None:
                        cell.value = f
                _aplicar_formato_percentual(cell)

        log(f"  ↳ {linhas_importadas} linha(s) importada(s).")
        if remover_duplicatas and linhas_duplicadas > 0:
            log(f"  ↳ {linhas_duplicadas} linha(s) duplicada(s) descartada(s).")

    if total_linhas == 0:
        log("[AVISO] Nenhuma linha de dados foi encontrada nos arquivos selecionados.")

    # Ajusta largura automática das colunas
    for col in ws_saida.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                if cell.value:
                    max_len = max(max_len, len(str(cell.value)))
            except Exception:
                pass
        ws_saida.column_dimensions[col_letter].width = min(max(max_len + 2, 12), 60)

    destino.parent.mkdir(parents=True, exist_ok=True)
    wb_saida.save(str(destino))
    log(f"Planilha unificada salva em: {destino} ({total_linhas} linhas no total).")
    return destino
