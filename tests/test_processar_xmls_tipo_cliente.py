import shutil
import tempfile
import unittest
from pathlib import Path

import openpyxl

from processamento_xml import COLUNAS_OBRIGATORIAS_AUTOMACAO, processar_pasta_xmls

PASTA_AMOSTRAS = Path(__file__).resolve().parent.parent / "Layout Portal Da Paraiba"
AMOSTRA_CPF = PASTA_AMOSTRAS / "NFSe_2026000000001_2WRB-S3DR.xml"
AMOSTRA_CNPJ = PASTA_AMOSTRAS / "NFSe_2026000000001_QRJF-LJEP.xml"
# Prestador "51.502.950 YURA PRISCILA BARBOSA RIQUE" — nome no padrão MEI
# (raiz do CNPJ formatada + nome).
AMOSTRA_MEI = PASTA_AMOSTRAS / "NFSe_2026000000019_AIP1-TJ45.xml"


@unittest.skipUnless(
    AMOSTRA_CPF.exists() and AMOSTRA_CNPJ.exists(),
    "Amostras de 'Layout Portal Da Paraiba' não encontradas",
)
class TestColunaTipoClienteEDestaqueDeCabecalho(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp_dir = tempfile.TemporaryDirectory()
        pasta = Path(cls.tmp_dir.name)
        shutil.copy(AMOSTRA_CPF, pasta / AMOSTRA_CPF.name)
        shutil.copy(AMOSTRA_CNPJ, pasta / AMOSTRA_CNPJ.name)

        caminho_planilha = processar_pasta_xmls(pasta_origem=str(pasta))
        cls.wb = openpyxl.load_workbook(caminho_planilha)
        cls.ws = cls.wb.active
        cls.cabecalhos = [str(c.value or "").strip().upper() for c in cls.ws[1]]
        cls.indice_tipo_cliente = cls.cabecalhos.index("TIPO_CLIENTE")

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        cls.tmp_dir.cleanup()

    def test_coluna_tipo_cliente_existe(self):
        self.assertIn("TIPO_CLIENTE", self.cabecalhos)

    def test_valores_da_coluna_tipo_cliente(self):
        valores = {
            row[0]: row[self.indice_tipo_cliente]
            for row in self.ws.iter_rows(min_row=2, values_only=True)
        }
        self.assertEqual(valores[AMOSTRA_CPF.name], "Pessoa Física")
        self.assertEqual(valores[AMOSTRA_CNPJ.name], "Pessoa Jurídica")

    def test_cabecalhos_obrigatorios_tem_destaque_de_cor(self):
        for celula in self.ws[1]:
            texto = str(celula.value or "").strip().upper()
            if texto in COLUNAS_OBRIGATORIAS_AUTOMACAO:
                self.assertIsNotNone(celula.fill.fgColor.rgb)
                self.assertNotEqual(celula.fill.fill_type, None)
                self.assertTrue(
                    str(celula.fill.fgColor.rgb).upper().endswith("FFC000"),
                    f"Cabeçalho obrigatório '{texto}' sem destaque de cor esperado.",
                )

    def test_cabecalho_nao_obrigatorio_sem_destaque(self):
        indice_arquivo = self.cabecalhos.index("ARQUIVO_XML")
        celula = self.ws[1][indice_arquivo]
        self.assertNotEqual(
            str(celula.fill.fgColor.rgb or "").upper()[-6:], "FFC000"
        )


@unittest.skipUnless(AMOSTRA_MEI.exists(), "Amostra MEI de 'Layout Portal Da Paraiba' não encontrada")
class TestColunaTipoTributacao(unittest.TestCase):
    """NFSe_2026000000019_AIP1-TJ45.xml tem prestador com nome no padrão MEI
    ("51.502.950 YURA PRISCILA BARBOSA RIQUE") — confirma a detecção e a
    regra de ISS_RETIDO="NÃO" para Simples Nacional MEI num caso real (não
    sintético). O conflito de prioridade com a regra de ID_CNAE 1207/1213 é
    coberto separadamente em tests/test_iss_retido.py com dados sintéticos,
    já que nenhuma amostra real tem as duas condições ao mesmo tempo (todo
    prestador MEI nas amostras tem a tag CodigoCnae presente, que sempre
    tem prioridade sobre ItemListaServico para compor o ID_CNAE)."""

    @classmethod
    def setUpClass(cls):
        cls.tmp_dir = tempfile.TemporaryDirectory()
        pasta = Path(cls.tmp_dir.name)
        shutil.copy(AMOSTRA_MEI, pasta / AMOSTRA_MEI.name)

        caminho_planilha = processar_pasta_xmls(pasta_origem=str(pasta))
        cls.wb = openpyxl.load_workbook(caminho_planilha)
        cls.ws = cls.wb.active
        cabecalhos = [str(c.value or "").strip().upper() for c in cls.ws[1]]
        linha = next(cls.ws.iter_rows(min_row=2, values_only=True))
        cls.valores = dict(zip(cabecalhos, linha))

    @classmethod
    def tearDownClass(cls):
        cls.wb.close()
        cls.tmp_dir.cleanup()

    def test_tipo_tributacao_e_simples_nacional_mei(self):
        self.assertEqual(self.valores["TIPO_TRIBUTACAO"], "Simples Nacional MEI")

    def test_iss_retido_fica_nao(self):
        self.assertEqual(self.valores["ISS_RETIDO"], "NÃO")


if __name__ == "__main__":
    unittest.main()
