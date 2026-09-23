"""Testes do Encerramento ISS — Múltiplos Meses (encerramento_iss_multi_periodo.py).

REGRA: nenhum teste abre um navegador nem toca o portal ISS. As interações com o portal
(login, busca de inscrição, navegação, processamento de cada competência) são sempre trocadas
por dublês (`unittest.mock`); só a leitura/escrita da planilha de entrada e a geração do
relatório usam arquivos reais (openpyxl), sem qualquer chamada de rede ou Selenium.
"""

import shutil
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from openpyxl import Workbook, load_workbook
from pydantic import SecretStr

import core.encerramento_iss_sem_movimento.controladores.encerramento_iss_multi_periodo as m
from core.encerramento_iss_sem_movimento.controladores.encerramento_iss import CodigoEncerramentoISS
from core.encerramento_iss_sem_movimento.iss.credenciais import CredenciaisISSFortaleza
from core.encerramento_iss_sem_movimento.iss.empresa import EmpresaSemMovimentoISSFortaleza

CAMINHO_TEMPLATE = Path("core/templates/template_relatorio_execucao.xlsx")


def _empresa(codigo=1, cnpj="12345678000199", nome="EMPRESA TESTE", linha=2) -> EmpresaSemMovimentoISSFortaleza:
    return EmpresaSemMovimentoISSFortaleza(
        codigo=codigo, nome=nome, cnpj=cnpj, responsavel="SM - FULANO", municipio="FORTALEZA", linha_planilha_fiscal=linha
    )


def _planilha_fiscal(pasta: Path, empresas: list[tuple[int, str, str]], nome: str = "planilha_fiscal.xlsx") -> Path:
    """Cria uma planilha fiscal mínima: cada item é (codigo, nome, cnpj); uma linha por empresa,
    a partir da linha 2 (linha 1 é cabeçalho, igual à planilha real)."""
    wb = Workbook()
    ws = wb.active
    ws["A1"], ws["B1"], ws["D1"], ws["H1"], ws["W1"], ws["Y1"] = "CODIGO", "NOME", "CNPJ", "RESPONSAVEL", "MUNICIPIO", "RESULTADO"
    for i, (codigo, nome_empresa, cnpj) in enumerate(empresas, start=2):
        ws[f"A{i}"] = codigo
        ws[f"B{i}"] = nome_empresa
        ws[f"D{i}"] = cnpj
        ws[f"H{i}"] = "SM - FULANO"
        ws[f"W{i}"] = "FORTALEZA"
    caminho = pasta / nome
    wb.save(caminho)
    return caminho


def _controlador(caminho_planilha: Path, pasta_saida: Path, competencias: list[date]) -> m.ControladorEncerramentoISSMultiPeriodo:
    return m.ControladorEncerramentoISSMultiPeriodo(
        caminho_planilha=caminho_planilha,
        caminho_saida=pasta_saida,
        caminho_webdriver=Path("chrome.exe"),
        url_iss_fortaleza="https://iss.fortaleza.ce.gov.br/grpfor/login.seam",
        credenciais=CredenciaisISSFortaleza(cpf=SecretStr("12345678901"), senha=SecretStr("segredo")),
        competencias=competencias,
        caminho_template_relatorio=CAMINHO_TEMPLATE,
        check_thread_stopped_callback=lambda: None,
    )


class TestConstrutor(unittest.TestCase):
    def test_sem_competencias_levanta_erro(self):
        with self.assertRaises(ValueError):
            _controlador(Path("planilha.xlsx"), Path("saida"), [])

    def test_competencias_ficam_ordenadas_e_sem_duplicatas(self):
        c = _controlador(
            Path("planilha.xlsx"),
            Path("saida"),
            [date(2026, 3, 1), date(2025, 12, 1), date(2026, 1, 1), date(2026, 3, 1)],
        )
        self.assertEqual(c.competencias, [date(2025, 12, 1), date(2026, 1, 1), date(2026, 3, 1)])


class TestResultados(unittest.TestCase):
    def test_contagens_da_classe_de_resultado(self):
        emp = _empresa()
        resultados = [
            m.ResultadoMesEmpresa(emp, date(2026, 1, 1), encerrada=True),
            m.ResultadoMesEmpresa(emp, date(2026, 2, 1), encerrada=False, problemas=["SERVIÇOS PENDENTES"]),
            m.ResultadoMesEmpresa(emp, date(2026, 3, 1), encerrada=True),
        ]
        resultado = m.ResultadoEncerramentoISSMultiPeriodo(status=None, resultados=resultados, report_paths=[])
        self.assertEqual(resultado.processadas, 3)
        self.assertEqual(resultado.encerradas, 2)
        self.assertEqual(resultado.problemas, 1)

    def test_contagens_com_lista_vazia(self):
        resultado = m.ResultadoEncerramentoISSMultiPeriodo(status=None, resultados=[], report_paths=[])
        self.assertEqual((resultado.processadas, resultado.encerradas, resultado.problemas), (0, 0, 0))


class TestAcrescentarNaPlanilhaFiscal(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199")])
        self.controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])

    def _ler_celula_y2(self) -> str:
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        valor = reader.get_cell("Y2").value
        reader.close_workbook()
        return valor

    def test_primeiro_acrescimo_nao_leva_ponto_e_virgula(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), "05/01/2026", row_number=2)
        reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "01/2026: 05/01/2026")

    def test_varias_competencias_ficam_concatenadas_na_mesma_celula_sem_apagar_a_anterior(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), "05/01/2026", row_number=2)
        reader.close_workbook()

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 2, 1), "SERVIÇOS PENDENTES", row_number=2)
        reader.close_workbook()

        self.assertEqual(self._ler_celula_y2(), "01/2026: 05/01/2026; 02/2026: SERVIÇOS PENDENTES")

    def test_sem_inscricao_nao_leva_rotulo_de_competencia(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, competencia=None, valor="S/ INSCRIÇÃO", row_number=2)
        reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "S/ INSCRIÇÃO")

    def test_nao_apaga_conteudo_que_ja_existia_manualmente_na_celula(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        reader.set_cell_value("Anotação manual do contador", coords="Y2")
        reader.save_spreadsheet(close_workbook=False)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), "05/01/2026", row_number=2)
        reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "Anotação manual do contador; 01/2026: 05/01/2026")

    def test_repetir_a_mesma_informacao_nao_empilha_na_celula(self):
        from python_spreadsheet_reader.readers import XLSXReader

        for _ in range(3):  # três execuções seguidas com o mesmo resultado
            reader = XLSXReader(self.caminho)
            self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), m.MSG_ERRO_PORTAL, row_number=2)
            reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "01/2026: ERRO NO PORTAL")

    def test_resultado_diferente_para_a_mesma_competencia_e_acrescentado(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), m.MSG_ERRO_PORTAL, row_number=2)
        reader.close_workbook()
        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), "05/02/2026", row_number=2)
        reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "01/2026: ERRO NO PORTAL; 01/2026: 05/02/2026")

    def test_competencia_diferente_com_o_mesmo_texto_nao_e_confundida(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 1, 1), m.MSG_ERRO_PORTAL, row_number=2)
        reader.close_workbook()
        reader = XLSXReader(self.caminho)
        self.controlador._acrescentar_na_planilha_fiscal(reader, date(2026, 2, 1), m.MSG_ERRO_PORTAL, row_number=2)
        reader.close_workbook()
        self.assertEqual(self._ler_celula_y2(), "01/2026: ERRO NO PORTAL; 02/2026: ERRO NO PORTAL")


@unittest.skipUnless(CAMINHO_TEMPLATE.exists(), "template do relatório não encontrado")
class TestRelatorioDeExecucao(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_rel_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.controlador = _controlador(self.pasta / "fiscal.xlsx", self.pasta / "saida", [date(2026, 1, 1), date(2026, 2, 1)])

    def test_sem_nenhum_resultado_levanta_excecao_e_nao_gera_arquivo(self):
        with self.assertRaises(m.SemEmpresasProcessadasException):
            self.controlador._gerar_relatorios_de_execucao()

    def test_gera_um_relatorio_por_competencia_na_pasta_ano_mes(self):
        emp_a = _empresa(codigo=1, cnpj="12345678000199", nome="EMPRESA A")
        emp_b = _empresa(codigo=2, cnpj="98765432000111", nome="EMPRESA B", linha=3)
        self.controlador.resultados = [
            m.ResultadoMesEmpresa(emp_a, date(2026, 1, 1), encerrada=True, dt_encerramento=date(2026, 1, 5)),
            m.ResultadoMesEmpresa(emp_b, date(2026, 1, 1), encerrada=False, problemas=["SERVIÇOS PENDENTES"]),
            m.ResultadoMesEmpresa(emp_a, date(2026, 2, 1), encerrada=True, dt_encerramento=date(2026, 2, 6)),
        ]
        caminhos = self.controlador._gerar_relatorios_de_execucao()

        self.assertEqual(len(caminhos), 2)  # uma competência sem resultado (só teria se pedisse 3 meses) não gera nada
        nomes_pastas = sorted(p.parent.name for p in caminhos)
        self.assertEqual(nomes_pastas, ["2026-01", "2026-02"])

        wb = load_workbook(caminhos[0])
        ws = wb.active
        self.assertEqual(ws.cell(row=7, column=2).value, "Competência: 01/2026")
        self.assertEqual(ws.cell(row=5, column=3).value, 2)  # processadas em jan/2026
        self.assertEqual(ws.cell(row=5, column=5).value, 1)  # encerradas em jan/2026
        self.assertEqual(ws.cell(row=5, column=7).value, 1)  # problemas em jan/2026
        nomes_no_relatorio = {ws.cell(row=10, column=2).value, ws.cell(row=11, column=2).value}
        self.assertEqual(nomes_no_relatorio, {"EMPRESA A", "EMPRESA B"})

    def test_relatorio_de_fevereiro_so_tem_a_empresa_daquele_mes(self):
        emp_a = _empresa(codigo=1, cnpj="12345678000199", nome="EMPRESA A")
        self.controlador.resultados = [
            m.ResultadoMesEmpresa(emp_a, date(2026, 1, 1), encerrada=True, dt_encerramento=date(2026, 1, 5)),
            m.ResultadoMesEmpresa(emp_a, date(2026, 2, 1), encerrada=True, dt_encerramento=date(2026, 2, 6)),
        ]
        caminhos = self.controlador._gerar_relatorios_de_execucao()
        caminho_fev = next(p for p in caminhos if p.parent.name == "2026-02")
        wb = load_workbook(caminho_fev)
        ws = wb.active
        self.assertEqual(ws.cell(row=5, column=3).value, 1)


class TestOrquestracaoPorEmpresaEPorMes(unittest.TestCase):
    """Exercita `_processar_linha` com os métodos que tocam o portal totalmente dublados —
    prova a ordem (por empresa, depois por mês) e o isolamento de erros, sem qualquer Selenium."""

    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_orq_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho_planilha = _planilha_fiscal(
            self.pasta, [(1, "EMPRESA A", "12345678000199"), (2, "EMPRESA B", "98765432000111")]
        )
        self.controlador = _controlador(
            self.caminho_planilha, self.pasta / "saida", [date(2026, 1, 1), date(2026, 2, 1)]
        )
        # Planilha com uma única empresa, para os testes que olham o resultado de UMA empresa
        # isoladamente (o comportamento entre empresas diferentes já é coberto pelos testes acima/abaixo).
        self.caminho_planilha_1_empresa = _planilha_fiscal(
            self.pasta, [(1, "EMPRESA A", "12345678000199")], nome="planilha_fiscal_1_empresa.xlsx"
        )
        self.controlador_1_empresa = _controlador(
            self.caminho_planilha_1_empresa, self.pasta / "saida", [date(2026, 1, 1), date(2026, 2, 1)]
        )

    def _reader(self):
        from python_spreadsheet_reader.readers import XLSXReader

        return XLSXReader(self.caminho_planilha)

    def _reader_1_empresa(self):
        from python_spreadsheet_reader.readers import XLSXReader

        return XLSXReader(self.caminho_planilha_1_empresa)

    def test_processa_todos_os_meses_de_uma_empresa_antes_de_abrir_o_modal_de_troca(self):
        chamadas = []
        with mock.patch.object(self.controlador, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(self.controlador, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(self.controlador, "_navegar_tela_escrituracao", lambda driver: chamadas.append("navegar")), \
             mock.patch.object(
                 self.controlador, "_processar_competencia",
                 lambda driver, reader, empresa, competencia: chamadas.append(("competencia", competencia)) or m.ResultadoMesEmpresa(empresa, competencia, encerrada=True),
             ), \
             mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao", lambda driver: chamadas.append("modal")):
            reader = self._reader()
            for row_number, row in reader.lazy_load_sheet():
                self.controlador._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()

        # 1 navegação e 1 abertura de modal por empresa, com as 2 competências entre elas, na ordem certa
        self.assertEqual(
            chamadas,
            [
                "navegar", ("competencia", date(2026, 1, 1)), ("competencia", date(2026, 2, 1)), "modal",
                "navegar", ("competencia", date(2026, 1, 1)), ("competencia", date(2026, 2, 1)), "modal",
            ],
        )
        self.assertEqual(len(self.controlador.resultados), 4)  # 2 empresas × 2 competências

    def test_empresa_sem_inscricao_marca_todas_as_competencias_e_nao_abre_o_modal(self):
        with mock.patch.object(self.controlador_1_empresa, "_procurar_inscricao_empresa", return_value=False), \
             mock.patch.object(self.controlador_1_empresa, "_abrir_modal_de_alteracao_de_inscricao") as modal, \
             mock.patch.object(self.controlador_1_empresa, "_navegar_tela_escrituracao") as navegar:
            reader = self._reader_1_empresa()
            # A linha 1 (cabeçalho) também é processada, mas é ignorada silenciosamente: os valores de
            # texto ("CODIGO", "NOME"...) falham na validação do Pydantic antes de chegar a qualquer
            # chamada ao portal — só a empresa da linha 2 realmente gera resultado.
            for row_number, row in reader.lazy_load_sheet():
                self.controlador_1_empresa._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()

        resultados = self.controlador_1_empresa.resultados
        self.assertEqual(len(resultados), 2)  # as 2 competências pedidas, ambas com problema
        self.assertTrue(all(r.problemas == ["S/ INSCRIÇÃO"] and not r.encerrada for r in resultados))
        modal.assert_not_called()
        navegar.assert_not_called()

        # a planilha recebeu o aviso, sem rótulo de mês (não é específico de uma competência)
        reader2 = self._reader_1_empresa()
        self.assertEqual(reader2.get_cell("Y2").value, "S/ INSCRIÇÃO")
        reader2.close_workbook()

    def test_erro_num_mes_nao_impede_os_demais_meses_da_mesma_empresa(self):
        def processar_competencia(driver, reader, empresa, competencia):
            if competencia == date(2026, 1, 1):
                raise m.NoSuchElementException("não achei a linha da competência")
            return m.ResultadoMesEmpresa(empresa, competencia, encerrada=True)

        with mock.patch.object(self.controlador_1_empresa, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(self.controlador_1_empresa, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(self.controlador_1_empresa, "_navegar_tela_escrituracao"), \
             mock.patch.object(self.controlador_1_empresa, "_processar_competencia", side_effect=processar_competencia), \
             mock.patch.object(self.controlador_1_empresa, "_abrir_modal_de_alteracao_de_inscricao") as modal:
            reader = self._reader_1_empresa()
            for row_number, row in reader.lazy_load_sheet():
                self.controlador_1_empresa._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()

        resultados = self.controlador_1_empresa.resultados
        self.assertEqual(len(resultados), 2)
        self.assertFalse(resultados[0].encerrada)
        self.assertEqual(resultados[0].problemas, [m.MSG_ERRO_PORTAL])  # texto curto, sem nome de exceção do Selenium
        self.assertTrue(resultados[1].encerrada)  # fevereiro seguiu normalmente
        modal.assert_called_once()  # ainda troca de empresa ao final, mesmo com uma falha no meio

        reader2 = self._reader_1_empresa()
        self.assertEqual(reader2.get_cell("Y2").value, "01/2026: ERRO NO PORTAL")
        reader2.close_workbook()

    def test_erro_inesperado_num_mes_refaz_o_caminho_pelo_menu_e_segue_para_o_proximo_mes(self):
        chamadas = []

        def processar_competencia(driver, reader, empresa, competencia):
            if competencia == date(2026, 1, 1):
                raise RuntimeError("falha no meio do fluxo")  # não é NoSuchElement/Timeout
            return m.ResultadoMesEmpresa(empresa, competencia, encerrada=True)

        with mock.patch.object(self.controlador_1_empresa, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(self.controlador_1_empresa, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(self.controlador_1_empresa, "_navegar_tela_escrituracao", lambda driver: chamadas.append("navegar")), \
             mock.patch.object(self.controlador_1_empresa, "_processar_competencia", side_effect=processar_competencia), \
             mock.patch.object(self.controlador_1_empresa, "_abrir_modal_de_alteracao_de_inscricao", lambda driver: chamadas.append("modal")):
            reader = self._reader_1_empresa()
            for row_number, row in reader.lazy_load_sheet():
                self.controlador_1_empresa._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()

        resultados = self.controlador_1_empresa.resultados
        self.assertEqual([r.encerrada for r in resultados], [False, True])
        self.assertEqual(resultados[0].problemas, [m.MSG_ERRO_PORTAL])
        # 1 navegação inicial + 1 de recuperação depois do erro em janeiro; o modal só no fim
        self.assertEqual(chamadas, ["navegar", "navegar", "modal"])

    def test_navegador_fechado_no_meio_de_um_mes_nao_e_engolido(self):
        def processar_competencia(driver, reader, empresa, competencia):
            raise m.InvalidSessionIdException("invalid session id: session deleted as the browser has closed the connection")

        with mock.patch.object(self.controlador_1_empresa, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(self.controlador_1_empresa, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(self.controlador_1_empresa, "_navegar_tela_escrituracao"), \
             mock.patch.object(self.controlador_1_empresa, "_processar_competencia", side_effect=processar_competencia), \
             mock.patch.object(self.controlador_1_empresa, "_abrir_modal_de_alteracao_de_inscricao"):
            reader = self._reader_1_empresa()
            with self.assertRaises(m.InvalidSessionIdException):
                for row_number, row in reader.lazy_load_sheet():
                    self.controlador_1_empresa._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()

        self.assertEqual(self.controlador_1_empresa.resultados, [])  # nada foi registrado como "erro de portal"

    def test_erro_inesperado_numa_empresa_nao_impede_a_proxima_empresa(self):
        chamadas_de_empresa = []

        def procurar_inscricao(driver, empresa):
            chamadas_de_empresa.append(empresa.codigo)
            if empresa.codigo == 1:
                raise RuntimeError("falha inesperada e não prevista no portal")
            return True

        with mock.patch.object(self.controlador, "_procurar_inscricao_empresa", side_effect=procurar_inscricao), \
             mock.patch.object(self.controlador, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(self.controlador, "_navegar_tela_escrituracao"), \
             mock.patch.object(
                 self.controlador, "_processar_competencia",
                 lambda driver, reader, empresa, competencia: m.ResultadoMesEmpresa(empresa, competencia, encerrada=True),
             ), \
             mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao"):
            reader = self._reader()
            for row_number, row in reader.lazy_load_sheet():
                try:
                    self.controlador._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
                except Exception:  # o executar_processo() de verdade também blinda cada linha; replicamos aqui
                    pass
            reader.close_workbook()

        self.assertEqual(chamadas_de_empresa, [1, 2])  # a empresa 2 foi tentada mesmo com a falha na 1
        self.assertEqual(len(self.controlador.resultados), 2)  # só os 2 resultados da empresa 2 (1 por competência)


SLEEP_RETRY = "core.encerramento_iss_sem_movimento.python_webdriver.functions.sleep"


class _ElementoFalso:
    def __init__(self, texto="", visivel=True, classe=""):
        self._texto = texto
        self._visivel = visivel
        self._classe = classe

    @property
    def text(self):
        return self._texto if self._visivel else ""  # como o Selenium: texto de elemento oculto vem vazio

    def is_displayed(self):
        return self._visivel

    def get_attribute(self, nome):
        return self._classe if nome == "class" else None


class _SeleniumFalso:
    def __init__(self, elementos=None, campos_existentes=None):
        self.elementos = elementos or {}  # id -> lista de elementos
        self.campos_existentes = campos_existentes  # None = todos existem
        self.scripts = []  # (campo, valor)

    def find_elements(self, by, valor):
        return list(self.elementos.get(valor, []))

    def execute_script(self, script, campo, valor):
        if self.campos_existentes is not None and campo not in self.campos_existentes:
            return False
        self.scripts.append((campo, valor))
        return True


class _DriverPortalFalso:
    """Só o necessário para os métodos de leitura/período: `get_driver()` e `find_element().by_xpath()`."""

    def __init__(self, selenium=None, celula_somatorio=None):
        self._selenium = selenium or _SeleniumFalso()
        self.celula_somatorio = celula_somatorio
        self.cliques = []

    def get_driver(self):
        return self._selenium

    def click(self, elemento, **_k):
        self.cliques.append(elemento)

    def find_element(self, *_a):
        driver = self

        class _Localizador:
            def by_id(self, id_):
                return ("id", id_)

            def by_xpath(self, xpath):
                class _Achado:
                    def get_element(self_inner):
                        return driver.celula_somatorio()

                return _Achado()

        return _Localizador()


class TestPeriodoEEsperasDaEtapa1(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_e1_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199")])
        self.controlador = _controlador(self.caminho, self.pasta / "saida", [date(2025, 5, 1)])

    def test_periodo_fixa_data_inicial_e_final_na_mesma_competencia(self):
        selenium = _SeleniumFalso()
        self.controlador._definir_periodo_consulta(_DriverPortalFalso(selenium), date(2025, 5, 1))
        self.assertEqual(
            selenium.scripts,
            [(campo, "05/2025") for campo in m.CAMPOS_PERIODO_CONSULTA],
        )
        # as quatro pontas: início/fim visíveis + os "CurrentDate" do calendário
        self.assertEqual(len(m.CAMPOS_PERIODO_CONSULTA), 4)
        self.assertTrue(any("dataFinal" in c for c in m.CAMPOS_PERIODO_CONSULTA))

    def test_periodo_com_campo_ausente_levanta_erro_claro(self):
        selenium = _SeleniumFalso(campos_existentes={m.CAMPOS_PERIODO_CONSULTA[0]})
        with self.assertRaises(m.NoSuchElementException):
            self.controlador._definir_periodo_consulta(_DriverPortalFalso(selenium), date(2025, 5, 1))

    def test_garantir_tela_de_lista_nao_navega_se_ja_esta_na_lista(self):
        selenium = _SeleniumFalso({m.ID_BOTAO_CONSULTAR: [_ElementoFalso()]})
        with mock.patch.object(self.controlador, "_navegar_tela_escrituracao") as navegar:
            self.controlador._garantir_tela_de_lista(_DriverPortalFalso(selenium))
        navegar.assert_not_called()

    def test_garantir_tela_de_lista_volta_pelo_menu_se_ficou_na_tela_da_competencia(self):
        selenium = _SeleniumFalso({m.ID_BOTAO_CONSULTAR: []})  # botão Consultar não existe na tela de encerramento
        with mock.patch.object(self.controlador, "_navegar_tela_escrituracao") as navegar:
            self.controlador._garantir_tela_de_lista(_DriverPortalFalso(selenium))
        navegar.assert_called_once()

    def test_modal_de_troca_ja_visivel_nao_faz_nada(self):
        selenium = _SeleniumFalso({m.ID_CAMPO_PESQUISA_INSCRICAO: [_ElementoFalso()]})
        with mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao") as abrir:
            self.controlador._garantir_modal_de_troca(_DriverPortalFalso(selenium))
        abrir.assert_not_called()

    def test_modal_que_demora_um_pouco_e_esperado_sem_reabrir(self):
        campo = _ElementoFalso(visivel=False)
        selenium = _SeleniumFalso({m.ID_CAMPO_PESQUISA_INSCRICAO: [campo]})
        chamadas = {"n": 0}

        def dormir(_s):
            chamadas["n"] += 1
            if chamadas["n"] == 3:
                campo._visivel = True

        with mock.patch.object(m.time, "sleep", dormir), \
             mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao") as abrir:
            self.controlador._garantir_modal_de_troca(_DriverPortalFalso(selenium))
        abrir.assert_not_called()

    def test_modal_que_nao_abre_sozinho_e_aberto_pelo_botao_do_topo(self):
        campo = _ElementoFalso(visivel=False)
        selenium = _SeleniumFalso({m.ID_CAMPO_PESQUISA_INSCRICAO: [campo]})

        def abrir(_driver):
            campo._visivel = True

        with mock.patch.object(m.time, "sleep"), \
             mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao", side_effect=abrir) as abrir_mock:
            self.controlador._garantir_modal_de_troca(_DriverPortalFalso(selenium))
        abrir_mock.assert_called_once()

    def test_modal_que_nunca_aparece_levanta_erro(self):
        selenium = _SeleniumFalso({m.ID_CAMPO_PESQUISA_INSCRICAO: [_ElementoFalso(visivel=False)]})
        with mock.patch.object(m.time, "sleep"), \
             mock.patch.object(self.controlador, "_abrir_modal_de_alteracao_de_inscricao"):
            with self.assertRaises(m.NoSuchElementException):
                self.controlador._garantir_modal_de_troca(_DriverPortalFalso(selenium))

    # -- serviços prestados: o Somatório fica na aba Encerramento --------------------------------

    def _selenium_com_aba(self, classe_aba):
        return _SeleniumFalso({"abaEncerramento_lbl": [_ElementoFalso(classe=classe_aba)]})

    def test_le_o_somatorio_sem_trocar_para_a_aba_servicos_prestados(self):
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-active"), lambda: _ElementoFalso("0"))
        motivo = self.controlador._verificar_servicos_prestados(driver, _empresa())
        self.assertIsNone(motivo)
        self.assertEqual(driver.cliques, [])  # nenhum clique: nem na aba Serviços Prestados, nem "de volta"

    def test_quantidade_maior_que_zero_e_problema(self):
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-active"), lambda: _ElementoFalso("1234"))
        self.assertEqual(self.controlador._verificar_servicos_prestados(driver, _empresa()), "SERVIÇOS PRESTADOS")

    def test_espera_o_texto_do_somatorio_aparecer_em_vez_de_ler_vazio(self):
        estados = iter([_ElementoFalso("", visivel=False), _ElementoFalso("", visivel=False), _ElementoFalso("0")])
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-active"), lambda: next(estados))
        with mock.patch(SLEEP_RETRY):
            motivo = self.controlador._verificar_servicos_prestados(driver, _empresa())
        self.assertIsNone(motivo)  # antes lia '' na primeira tentativa e marcava ERRO AO LER SERVIÇOS PRESTADOS

    def test_somatorio_que_nunca_aparece_vira_erro_de_leitura_e_nao_encerra(self):
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-active"), lambda: _ElementoFalso("", visivel=False))
        with mock.patch(SLEEP_RETRY):
            motivo = self.controlador._verificar_servicos_prestados(driver, _empresa())
        self.assertEqual(motivo, "ERRO AO LER SERVIÇOS PRESTADOS")

    def test_valor_nao_numerico_vira_erro_de_leitura(self):
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-active"), lambda: _ElementoFalso("abc"))
        self.assertEqual(self.controlador._verificar_servicos_prestados(driver, _empresa()), "ERRO AO LER SERVIÇOS PRESTADOS")

    def test_se_a_aba_encerramento_nao_esta_ativa_clica_nela_antes_de_ler(self):
        driver = _DriverPortalFalso(self._selenium_com_aba("rich-tab-header rich-tab-inactive"), lambda: _ElementoFalso("0"))
        self.controlador._verificar_servicos_prestados(driver, _empresa())
        self.assertEqual(driver.cliques, [("id", "abaEncerramento_lbl")])

    def test_situacao_do_portal_e_lida_sem_o_rotulo_da_competencia(self):
        class _Linha:
            def get_element(self_inner):
                return _ElementoFalso("09/2026 Aberta - Normal 01/09/2026")

        self.assertEqual(m.ControladorEncerramentoISSMultiPeriodo._ler_situacao(_Linha(), "09/2026"), "Aberta - Normal 01/09/2026")

    def test_situacao_que_falha_ao_ler_devolve_vazio_sem_levantar(self):
        class _LinhaQuebrada:
            def get_element(self_inner):
                raise RuntimeError("stale")

        self.assertEqual(m.ControladorEncerramentoISSMultiPeriodo._ler_situacao(_LinhaQuebrada(), "09/2026"), "")

    def test_busca_de_inscricao_garante_o_modal_antes_de_qualquer_clique(self):
        driver = mock.MagicMock()
        empresa = _empresa()
        empresa.cnpj_m = "12.345.678/0001-99"
        with mock.patch.object(self.controlador, "_garantir_modal_de_troca", side_effect=m.NoSuchElementException("sem modal")):
            with self.assertRaises(m.NoSuchElementException):
                self.controlador._procurar_inscricao_empresa(driver, empresa)
        driver.click.assert_not_called()

    def test_navegar_para_a_escrituracao_espera_a_tela_de_lista_carregar(self):
        ordem = []
        driver = mock.MagicMock()
        driver.click.side_effect = lambda *_a, **_k: ordem.append("clique")
        with mock.patch.object(self.controlador, "_aguardar_tela_de_lista", lambda d: ordem.append("aguardou")):
            self.controlador._navegar_tela_escrituracao(driver)
        self.assertEqual(ordem, ["clique", "clique", "aguardou"])

    def test_processar_competencia_garante_a_lista_e_fixa_o_periodo_antes_de_consultar(self):
        ordem = []
        with mock.patch.object(self.controlador, "_garantir_tela_de_lista", lambda d: ordem.append("garantir")), \
             mock.patch.object(self.controlador, "_definir_periodo_consulta", lambda d, c: ordem.append(("periodo", c))), \
             mock.patch.object(m, "retry_on_exception", side_effect=m.NoSuchElementException("sem linha")):
            driver = mock.MagicMock()
            driver.click.side_effect = lambda *_a, **_k: ordem.append("consultar")
            from python_spreadsheet_reader.readers import XLSXReader

            reader = XLSXReader(self.caminho)
            self.controlador._processar_competencia(driver, reader, _empresa(linha=2), date(2025, 5, 1))
            reader.close_workbook()
        self.assertEqual(ordem, ["garantir", ("periodo", date(2025, 5, 1)), "consultar"])


def _controlador_modo(caminho_planilha, pasta_saida, competencias, **extra):
    return m.ControladorEncerramentoISSMultiPeriodo(
        caminho_planilha=caminho_planilha,
        caminho_saida=pasta_saida,
        caminho_webdriver=Path("chrome.exe"),
        url_iss_fortaleza="https://iss.fortaleza.ce.gov.br/grpfor/login.seam",
        credenciais=CredenciaisISSFortaleza(cpf=SecretStr("12345678901"), senha=SecretStr("segredo")),
        competencias=competencias,
        caminho_template_relatorio=CAMINHO_TEMPLATE,
        check_thread_stopped_callback=lambda: None,
        **extra,
    )


class TestResultadosPorCategoria(unittest.TestCase):
    def test_categoria_deduzida_de_encerrada_quando_nao_informada(self):
        emp = _empresa()
        self.assertEqual(m.ResultadoMesEmpresa(emp, date(2026, 1, 1), encerrada=True).categoria, m.ACAO_ENCERRADA_AGORA)
        self.assertEqual(m.ResultadoMesEmpresa(emp, date(2026, 1, 1), encerrada=False).categoria, m.ACAO_PROBLEMA)

    def test_contagens_separam_agora_ja_encerradas_aptas_e_problemas(self):
        emp = _empresa()
        r = m.ResultadoEncerramentoISSMultiPeriodo(
            None,
            [
                m.ResultadoMesEmpresa(emp, date(2026, 1, 1), encerrada=True, acao=m.ACAO_ENCERRADA_AGORA),
                m.ResultadoMesEmpresa(emp, date(2026, 2, 1), encerrada=True, acao=m.ACAO_JA_ENCERRADA_CERT_EXISTENTE),
                m.ResultadoMesEmpresa(emp, date(2026, 3, 1), encerrada=True, acao=m.ACAO_JA_ENCERRADA_CERT_BAIXADO),
                m.ResultadoMesEmpresa(emp, date(2026, 4, 1), encerrada=False, acao=m.ACAO_APTA),
                m.ResultadoMesEmpresa(emp, date(2026, 5, 1), encerrada=False, problemas=["SERVIÇOS PENDENTES"]),
            ],
            [],
        )
        self.assertEqual(r.processadas, 5)
        self.assertEqual(r.encerradas_agora, 1)
        self.assertEqual(r.ja_encerradas, 2)
        self.assertEqual(r.encerradas, 3)  # agora + já estavam (o relatório continua contando assim)
        self.assertEqual(r.aptas, 1)
        self.assertEqual(r.problemas, 1)  # a "apta" da verificação NÃO é problema

    def test_modo_invalido_levanta_erro(self):
        with self.assertRaises(ValueError):
            _controlador_modo(Path("p.xlsx"), Path("saida"), [date(2026, 1, 1)], modo="qualquer")


class TestPausaEParada(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_pausa_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199"), (2, "EMPRESA B", "98765432000111")])

    def _rodar_linhas(self, controlador):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        try:
            for row_number, row in reader.lazy_load_sheet():
                controlador._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
        finally:
            reader.close_workbook()

    def _dubles(self, controlador, processados):
        return (
            mock.patch.object(controlador, "_procurar_inscricao_empresa", return_value=True),
            mock.patch.object(controlador, "_dar_ciencia_nas_mensagens_nao_lidas"),
            mock.patch.object(controlador, "_navegar_tela_escrituracao"),
            mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao"),
            mock.patch.object(
                controlador, "_processar_competencia",
                lambda d, r, e, c: processados.append((e.codigo, c)) or m.ResultadoMesEmpresa(e, c, encerrada=True),
            ),
        )

    def test_ponto_seguro_consulta_pausa_e_depois_parada(self):
        ordem = []
        c = _controlador_modo(self.caminho, self.pasta / "s", [date(2026, 1, 1)],
                              callback_pausa=lambda: ordem.append("pausa"), callback_parar=lambda: ordem.append("parar") or False)
        c._ponto_seguro()
        self.assertEqual(ordem, ["pausa", "parar"])

    def test_ponto_seguro_levanta_o_sinal_de_parada(self):
        c = _controlador_modo(self.caminho, self.pasta / "s", [date(2026, 1, 1)], callback_parar=lambda: True)
        with self.assertRaises(m.UserStoppedThreadException):
            c._ponto_seguro()

    def test_ponto_seguro_sem_callbacks_nao_faz_nada(self):
        _controlador_modo(self.caminho, self.pasta / "s", [date(2026, 1, 1)])._ponto_seguro()

    def test_pausa_e_consultada_antes_de_cada_competencia(self):
        pausas, processados = [], []
        c = _controlador_modo(self.caminho, self.pasta / "s", [date(2026, 1, 1), date(2026, 2, 1)],
                              callback_pausa=lambda: pausas.append(len(processados)))
        with self.subTest():
            d1, d2, d3, d4, d5 = self._dubles(c, processados)
            with d1, d2, d3, d4, d5:
                self._rodar_linhas(c)
        # 2 empresas x 2 competências: a pausa é consultada ANTES de cada uma (nunca no meio de uma)
        self.assertEqual(pausas, [0, 1, 2, 3])

    def test_parar_termina_a_competencia_em_andamento_e_nao_comeca_a_proxima(self):
        processados = []
        c = _controlador_modo(self.caminho, self.pasta / "s", [date(2026, 1, 1), date(2026, 2, 1)],
                              callback_parar=lambda: len(processados) >= 1)
        d1, d2, d3, d4, d5 = self._dubles(c, processados)
        with d1, d2, d3, d4, d5:
            with self.assertRaises(m.UserStoppedThreadException):
                self._rodar_linhas(c)
        self.assertEqual(processados, [(1, date(2026, 1, 1))])  # a 1ª competência foi até o fim; a 2ª nem começou
        self.assertEqual(len(c.resultados), 1)

    def test_parar_pelo_executar_processo_gera_relatorio_do_que_ja_foi_feito(self):
        if not CAMINHO_TEMPLATE.exists():
            self.skipTest("template do relatório não encontrado")
        processados = []
        c = _controlador_modo(self.caminho, self.pasta / "saida", [date(2026, 1, 1)],
                              callback_parar=lambda: len(processados) >= 1)
        driver_falso = mock.Mock(name="ChromeDriverFalso")
        d1, d2, d3, d4, d5 = self._dubles(c, processados)
        with mock.patch.object(m, "ChromeDriver", mock.Mock(return_value=driver_falso)), \
             mock.patch.object(c, "_login_iss_fortaleza"), d1, d2, d3, d4, d5:
            resultado = c.executar_processo()
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.USER_ENDED_PROCESS)
        self.assertEqual(resultado.processadas, 1)
        self.assertEqual(len(resultado.report_paths), 1)  # relatório do que foi feito antes de parar
        driver_falso.quit_driver.assert_called_once()


class TestCertificadoExistenteEJaEncerrada(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_cert_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(7, "EMPRESA A", "12345678000199")])
        self.saida = self.pasta / "saida"
        self.competencia = date(2025, 10, 1)
        self.empresa = _empresa(codigo=7, linha=2)

    def _criar_certificado(self, nome="CERTIFICADO ISS_EMP7_20251101.pdf", conteudo=b"%PDF-1.4 conteudo"):
        pasta = self.saida / "2025-10" / "EMP_7"
        pasta.mkdir(parents=True, exist_ok=True)
        (pasta / nome).write_bytes(conteudo)
        return pasta / nome

    def _driver_com_data(self, texto="07/11/2025"):
        driver = mock.MagicMock()
        driver.find_element.return_value.by_xpath.return_value.get_element.return_value.text = texto
        return driver

    def _ler_y2(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        valor = reader.get_cell("Y2").value
        reader.close_workbook()
        return valor

    def _processar(self, controlador, driver):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        try:
            return controlador._processar_ja_encerrada(driver, reader, self.empresa, self.competencia, object())
        finally:
            reader.close_workbook()

    def test_certificado_existente_e_encontrado(self):
        pdf = self._criar_certificado()
        c = _controlador(self.caminho, self.saida, [self.competencia])
        self.assertEqual(c._certificado_existente(self.empresa, self.competencia), pdf)

    def test_pdf_vazio_nao_conta_como_certificado(self):
        self._criar_certificado(conteudo=b"")
        c = _controlador(self.caminho, self.saida, [self.competencia])
        self.assertIsNone(c._certificado_existente(self.empresa, self.competencia))

    def test_sem_pasta_nao_ha_certificado(self):
        c = _controlador(self.caminho, self.saida, [self.competencia])
        self.assertIsNone(c._certificado_existente(self.empresa, self.competencia))

    def test_certificado_de_outra_competencia_ou_empresa_nao_conta(self):
        outra = self.saida / "2025-09" / "EMP_7"
        outra.mkdir(parents=True)
        (outra / "CERTIFICADO ISS_EMP7_x.pdf").write_bytes(b"pdf")
        outra2 = self.saida / "2025-10" / "EMP_8"
        outra2.mkdir(parents=True)
        (outra2 / "CERTIFICADO ISS_EMP8_x.pdf").write_bytes(b"pdf")
        c = _controlador(self.caminho, self.saida, [self.competencia])
        self.assertIsNone(c._certificado_existente(self.empresa, self.competencia))

    def test_ja_encerrada_com_certificado_nao_baixa_de_novo(self):
        pdf = self._criar_certificado()
        c = _controlador(self.caminho, self.saida, [self.competencia])
        driver = self._driver_com_data()
        with mock.patch.object(c, "_imprimir_declaracao_fechamento_iss") as imprimir:
            resultado = self._processar(c, driver)
        driver.click.assert_not_called()  # nem abriu o certificado
        imprimir.assert_not_called()
        self.assertEqual(resultado.acao, m.ACAO_JA_ENCERRADA_CERT_EXISTENTE)
        self.assertTrue(resultado.encerrada)
        self.assertEqual(resultado.caminho_certificado, pdf)
        self.assertEqual(resultado.dt_encerramento, date(2025, 11, 7))
        self.assertEqual(self._ler_y2(), "10/2025: 07/11/2025")  # continua registrando a data na planilha

    def test_ja_encerrada_sem_certificado_baixa(self):
        c = _controlador(self.caminho, self.saida, [self.competencia])
        driver = self._driver_com_data()
        with mock.patch.object(c, "_imprimir_declaracao_fechamento_iss", return_value="saida/x.pdf") as imprimir:
            resultado = self._processar(c, driver)
        imprimir.assert_called_once()
        self.assertGreaterEqual(driver.click.call_count, 2)
        driver.get_driver.return_value.back.assert_called_once()
        self.assertEqual(resultado.acao, m.ACAO_JA_ENCERRADA_CERT_BAIXADO)

    def test_verificacao_ja_encerrada_sem_certificado_nao_baixa_nem_grava(self):
        c = _controlador_modo(self.caminho, self.saida, [self.competencia], modo=m.MODO_VERIFICAR)
        driver = self._driver_com_data()
        with mock.patch.object(c, "_imprimir_declaracao_fechamento_iss") as imprimir:
            resultado = self._processar(c, driver)
        driver.click.assert_not_called()
        imprimir.assert_not_called()
        self.assertEqual(resultado.acao, m.ACAO_JA_ENCERRADA_SEM_CERT)
        self.assertIsNone(self._ler_y2())  # verificação não deixa rastro na planilha

    def test_verificacao_ja_encerrada_com_certificado(self):
        self._criar_certificado()
        c = _controlador_modo(self.caminho, self.saida, [self.competencia], modo=m.MODO_VERIFICAR)
        resultado = self._processar(c, self._driver_com_data())
        self.assertEqual(resultado.acao, m.ACAO_JA_ENCERRADA_CERT_EXISTENTE)

    def test_caminho_do_certificado_baixado_usa_a_mesma_pasta_da_checagem(self):
        c = _controlador(self.caminho, self.saida, [self.competencia])
        driver = mock.MagicMock()
        c._imprimir_declaracao_fechamento_iss(driver, self.empresa, self.competencia)
        pasta_usada = Path(driver.page_to_pdf.call_args.kwargs["file_path"])
        self.assertEqual(pasta_usada, c._diretorio_certificados(self.empresa, self.competencia))
        self.assertEqual(pasta_usada.parts[-2:], ("2025-10", "EMP_7"))


class TestModoVerificar(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_verif_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199")])
        self.competencia = date(2026, 1, 1)
        self.empresa = _empresa(linha=2)

    def _reader(self):
        from python_spreadsheet_reader.readers import XLSXReader

        return XLSXReader(self.caminho)

    def _a_encerrar(self, controlador, driver):
        reader = self._reader()
        try:
            return controlador._processar_a_encerrar(driver, reader, self.empresa, self.competencia, "LINK")
        finally:
            reader.close_workbook()

    def test_verificar_apta_nao_clica_em_encerrar_nem_no_sim(self):
        c = _controlador_modo(self.caminho, self.pasta / "s", [self.competencia], modo=m.MODO_VERIFICAR)
        driver = mock.MagicMock()
        driver.scroll_element_into_view.side_effect = AssertionError("TENTATIVA DE CLICAR EM ENCERRAR NA VERIFICAÇÃO")
        with mock.patch.object(c, "_verificar_servicos_prestados", return_value=None), \
             mock.patch.object(c, "_verificar_servicos_pendentes", return_value=None):
            resultado = self._a_encerrar(c, driver)
        self.assertEqual(resultado.categoria, m.ACAO_APTA)
        self.assertFalse(resultado.encerrada)
        self.assertEqual(resultado.problemas, [])
        self.assertEqual(driver.click.call_args_list, [mock.call("LINK")])  # só abriu a tela; nada além disso

    def test_verificar_com_problema_registra_o_motivo_sem_gravar_na_planilha(self):
        c = _controlador_modo(self.caminho, self.pasta / "s", [self.competencia], modo=m.MODO_VERIFICAR)
        with mock.patch.object(c, "_verificar_servicos_prestados", return_value="SERVIÇOS PRESTADOS"):
            resultado = self._a_encerrar(c, mock.MagicMock())
        self.assertEqual(resultado.categoria, m.ACAO_PROBLEMA)
        self.assertEqual(resultado.problemas, ["SERVIÇOS PRESTADOS"])
        reader = self._reader()
        self.assertIsNone(reader.get_cell("Y2").value)
        reader.close_workbook()

    def test_modo_encerrar_apta_realmente_encerra_e_marca_encerrada_agora(self):
        c = _controlador(self.caminho, self.pasta / "s", [self.competencia])
        driver = mock.MagicMock()
        with mock.patch.object(c, "_verificar_servicos_prestados", return_value=None), \
             mock.patch.object(c, "_verificar_servicos_pendentes", return_value=None), \
             mock.patch.object(c, "_imprimir_declaracao_fechamento_iss", return_value="saida/x.pdf"):
            resultado = self._a_encerrar(c, driver)
        driver.scroll_element_into_view.assert_called_once()  # o clique em Encerrar aconteceu
        self.assertEqual(resultado.categoria, m.ACAO_ENCERRADA_AGORA)
        self.assertTrue(resultado.encerrada)

    def test_verificacao_nunca_gera_relatorio_de_encerramento(self):
        c = _controlador_modo(self.caminho, self.pasta / "saida", [self.competencia], modo=m.MODO_VERIFICAR)
        driver_falso = mock.Mock(name="ChromeDriverFalso")
        with mock.patch.object(m, "ChromeDriver", mock.Mock(return_value=driver_falso)), \
             mock.patch.object(c, "_login_iss_fortaleza"), \
             mock.patch.object(c, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(c, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(c, "_navegar_tela_escrituracao"), \
             mock.patch.object(c, "_abrir_modal_de_alteracao_de_inscricao"), \
             mock.patch.object(c, "_gerar_relatorios_de_execucao") as gerar, \
             mock.patch.object(
                 c, "_processar_competencia",
                 lambda d, r, e, comp: m.ResultadoMesEmpresa(e, comp, encerrada=False, acao=m.ACAO_APTA),
             ):
            resultado = c.executar_processo()
        gerar.assert_not_called()
        self.assertEqual(resultado.report_paths, [])
        self.assertEqual((resultado.aptas, resultado.problemas), (1, 0))
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.SUCCESS)

    def test_planilha_nao_e_gravada_em_nenhum_caminho_da_verificacao(self):
        c = _controlador_modo(self.caminho, self.pasta / "s", [self.competencia], modo=m.MODO_VERIFICAR)
        reader = self._reader()
        c._acrescentar_na_planilha_fiscal(reader, self.competencia, "QUALQUER", row_number=2)
        c._acrescentar_na_planilha_fiscal(reader, None, "S/ INSCRIÇÃO", row_number=2)
        self.assertIsNone(reader.get_cell("Y2").value)
        reader.close_workbook()


class TestExecucaoRestritaAoQueFoiVerificado(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_alvos_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199"), (2, "EMPRESA B", "98765432000111")])

    def test_competencias_da_empresa_filtra_pelos_alvos(self):
        c = _controlador_modo(
            self.caminho, self.pasta / "s", [date(2026, 1, 1), date(2026, 2, 1), date(2026, 3, 1)],
            alvos={"12345678000199": {date(2026, 1, 1), date(2026, 3, 1)}},
        )
        self.assertEqual(c._competencias_da_empresa(_empresa(cnpj="12345678000199")), [date(2026, 1, 1), date(2026, 3, 1)])
        self.assertEqual(c._competencias_da_empresa(_empresa(codigo=2, cnpj="98765432000111")), [])

    def test_sem_alvos_processa_todas_as_competencias(self):
        c = _controlador(self.caminho, self.pasta / "s", [date(2026, 1, 1), date(2026, 2, 1)])
        self.assertEqual(c._competencias_da_empresa(_empresa()), [date(2026, 1, 1), date(2026, 2, 1)])

    def test_empresa_fora_da_lista_aprovada_nem_e_procurada_no_portal(self):
        from python_spreadsheet_reader.readers import XLSXReader

        processados = []
        c = _controlador_modo(
            self.caminho, self.pasta / "s", [date(2026, 1, 1), date(2026, 2, 1)],
            alvos={"12345678000199": {date(2026, 2, 1)}},  # só a empresa A, só fevereiro
        )
        procurar = mock.Mock(return_value=True)
        with mock.patch.object(c, "_procurar_inscricao_empresa", procurar), \
             mock.patch.object(c, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(c, "_navegar_tela_escrituracao"), \
             mock.patch.object(c, "_abrir_modal_de_alteracao_de_inscricao"), \
             mock.patch.object(
                 c, "_processar_competencia",
                 lambda d, r, e, comp: processados.append((e.codigo, comp)) or m.ResultadoMesEmpresa(e, comp, encerrada=True),
             ):
            reader = XLSXReader(self.caminho)
            for row_number, row in reader.lazy_load_sheet():
                c._processar_linha(driver=object(), reader=reader, row_number=row_number, row=row)
            reader.close_workbook()
        self.assertEqual(procurar.call_count, 1)  # a empresa B nem foi buscada
        self.assertEqual(processados, [(1, date(2026, 2, 1))])


def _dv_cnpj(base12: str) -> str:
    """Implementação independente do algoritmo dos dígitos verificadores (não usa o código testado)."""
    def dv(base, pesos):
        r = sum(int(a) * b for a, b in zip(base, pesos)) % 11
        return "0" if r < 2 else str(11 - r)

    p1 = [5, 4, 3, 2, 9, 8, 7, 6, 5, 4, 3, 2]
    d1 = dv(base12, p1)
    d2 = dv(base12 + d1, [6] + p1)
    return base12 + d1 + d2


CNPJ_A = _dv_cnpj("112223330001")  # 11222333000181
CNPJ_B = _dv_cnpj("456789120001")


class TestCnpjValido(unittest.TestCase):
    def test_gerador_independente_bate_com_um_cnpj_conhecido(self):
        self.assertEqual(CNPJ_A, "11222333000181")

    def test_aceita_cnpj_valido_com_e_sem_mascara(self):
        self.assertTrue(m.cnpj_valido("11222333000181"))
        self.assertTrue(m.cnpj_valido("11.222.333/0001-81"))
        self.assertTrue(m.cnpj_valido(CNPJ_B))

    def test_rejeita_digito_verificador_errado(self):
        self.assertFalse(m.cnpj_valido("11222333000182"))
        self.assertFalse(m.cnpj_valido("12345678000199"))

    def test_rejeita_todos_iguais_tamanho_errado_e_vazio(self):
        self.assertFalse(m.cnpj_valido("11111111111111"))
        self.assertFalse(m.cnpj_valido("00000000000000"))
        self.assertFalse(m.cnpj_valido("1122233300018"))
        self.assertFalse(m.cnpj_valido(""))
        self.assertFalse(m.cnpj_valido(None))


class TestEmpresaAlvo(unittest.TestCase):
    def test_texto_sem_nome_mostra_o_cnpj_com_mascara(self):
        self.assertEqual(str(m.EmpresaAlvo(cnpj=CNPJ_A)), "CNPJ 11.222.333/0001-81")

    def test_texto_com_nome(self):
        self.assertEqual(str(m.EmpresaAlvo(cnpj=CNPJ_A, nome="EMPRESA X")), f"EMPRESA X (CNPJ: {CNPJ_A})")

    def test_identificador_usa_o_codigo_na_planilha_e_o_cnpj_no_manual(self):
        self.assertEqual(m.identificador_empresa(_empresa(codigo=7)), "EMP_7")
        self.assertEqual(m.identificador_empresa(m.EmpresaAlvo(cnpj=CNPJ_A)), f"CNPJ_{CNPJ_A}")


class TestConstrutorOrigem(unittest.TestCase):
    def test_exige_exatamente_uma_origem(self):
        with self.assertRaises(ValueError):
            _controlador_modo(None, Path("s"), [date(2026, 1, 1)])  # nenhuma
        with self.assertRaises(ValueError):
            _controlador_modo(Path("p.xlsx"), Path("s"), [date(2026, 1, 1)], cnpjs=[CNPJ_A])  # as duas

    def test_cnpj_invalido_no_construtor_levanta(self):
        with self.assertRaises(ValueError):
            _controlador_modo(None, Path("s"), [date(2026, 1, 1)], cnpjs=["123"])

    def test_cnpjs_sao_normalizados_e_sem_duplicatas_na_ordem(self):
        c = _controlador_modo(None, Path("s"), [date(2026, 1, 1)], cnpjs=["11.222.333/0001-81", CNPJ_B, CNPJ_A])
        self.assertEqual(c.cnpjs, [CNPJ_A, CNPJ_B])

    def test_planilha_continua_funcionando_sem_cnpjs(self):
        c = _controlador_modo(Path("p.xlsx"), Path("s"), [date(2026, 1, 1)])
        self.assertIsNone(c.cnpjs)


class TestOrigemManual(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_manual_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)

    def _controlador(self, cnpjs=None, competencias=None, **extra):
        return _controlador_modo(None, self.pasta / "saida", competencias or [date(2026, 1, 1)], cnpjs=cnpjs or [CNPJ_A, CNPJ_B], **extra)

    def _rodar(self, controlador, **patches):
        driver_falso = mock.Mock(name="ChromeDriverFalso")
        real = m.XLSXReader
        self.planilhas_abertas = []

        def leitor(*a, **k):  # registra o que foi aberto; só o modelo do relatório é permitido na origem manual
            self.planilhas_abertas.append(str(a[0] if a else k.get("workbook_path")))
            return real(*a, **k)

        with mock.patch.object(m, "ChromeDriver", mock.Mock(return_value=driver_falso)), \
             mock.patch.object(m, "XLSXReader", leitor), \
             mock.patch.object(controlador, "_login_iss_fortaleza"), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao"), \
             mock.patch.multiple(controlador, **patches):
            return controlador.executar_processo(), driver_falso

    def _ok(self, empresa_falsa=None):
        return dict(
            _procurar_inscricao_empresa=mock.Mock(return_value=True),
            _dar_ciencia_nas_mensagens_nao_lidas=mock.Mock(),
            _navegar_tela_escrituracao=mock.Mock(),
            _processar_competencia=mock.Mock(
                side_effect=lambda d, r, e, c: m.ResultadoMesEmpresa(e, c, encerrada=True, dt_encerramento=date(2026, 1, 5))
            ),
        )

    def test_percorre_os_cnpjs_sem_abrir_planilha(self):
        c = self._controlador()
        patches = self._ok()
        resultado, driver = self._rodar(c, **patches)
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.SUCCESS)
        self.assertIn("CNPJs informados", resultado.status.message)
        self.assertEqual(resultado.processadas, 2)  # 2 empresas x 1 competência
        # nenhuma planilha fiscal foi aberta (só o modelo do relatório) e o reader é None em todas as chamadas
        self.assertTrue(all(p == str(CAMINHO_TEMPLATE) for p in self.planilhas_abertas), self.planilhas_abertas)
        leitores = {chamada.args[1] for chamada in patches["_processar_competencia"].call_args_list}
        self.assertEqual(leitores, {None})
        driver.quit_driver.assert_called_once()

    def test_empresas_processadas_na_ordem_informada_como_empresa_alvo(self):
        c = self._controlador(cnpjs=[CNPJ_B, CNPJ_A])
        patches = self._ok()
        self._rodar(c, **patches)
        empresas = [chamada.args[1] for chamada in patches["_procurar_inscricao_empresa"].call_args_list]
        self.assertTrue(all(isinstance(e, m.EmpresaAlvo) for e in empresas))
        self.assertEqual([e.cnpj for e in empresas], [CNPJ_B, CNPJ_A])
        self.assertTrue(all(e.linha_planilha_fiscal is None and e.codigo is None for e in empresas))

    def test_gera_relatorio_mesmo_sem_codigo_nem_responsavel(self):
        if not CAMINHO_TEMPLATE.exists():
            self.skipTest("template do relatório não encontrado")
        c = self._controlador()
        resultado, _ = self._rodar(c, **self._ok())
        self.assertEqual(len(resultado.report_paths), 1)
        self.assertTrue(resultado.report_paths[0].exists())

    def test_sem_inscricao_no_manual_marca_problema_sem_tocar_em_planilha(self):
        c = self._controlador(cnpjs=[CNPJ_A])
        patches = self._ok()
        patches["_procurar_inscricao_empresa"] = mock.Mock(return_value=False)
        resultado, _ = self._rodar(c, **patches)
        self.assertEqual([r.problemas for r in resultado.resultados], [["S/ INSCRIÇÃO"]])

    def test_erro_inesperado_numa_empresa_recupera_e_segue_para_a_proxima(self):
        c = self._controlador()
        tentadas = []

        def procurar(driver, empresa):
            tentadas.append(empresa.cnpj)
            if empresa.cnpj == CNPJ_A:
                raise RuntimeError("falha inesperada")
            return True

        patches = self._ok()
        patches["_procurar_inscricao_empresa"] = mock.Mock(side_effect=procurar)
        recuperar = mock.Mock()
        patches["_recuperar_para_proxima_empresa"] = recuperar
        resultado, _ = self._rodar(c, **patches)
        self.assertEqual(tentadas, [CNPJ_A, CNPJ_B])
        recuperar.assert_called_once()
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.SUCCESS)

    def test_navegador_fechado_para_a_execucao_no_manual_tambem(self):
        c = self._controlador()
        tentadas = []

        def procurar(driver, empresa):
            tentadas.append(empresa.cnpj)
            raise m.InvalidSessionIdException("invalid session id")

        patches = self._ok()
        patches["_procurar_inscricao_empresa"] = mock.Mock(side_effect=procurar)
        resultado, _ = self._rodar(c, **patches)
        self.assertEqual(tentadas, [CNPJ_A])
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.ERROR_UNHANDLED_EXCEPTION)

    def test_parar_no_manual_termina_a_empresa_em_andamento_e_nao_comeca_a_proxima(self):
        processados = []
        c = self._controlador(callback_parar=lambda: len(processados) >= 1)
        patches = self._ok()
        patches["_processar_competencia"] = mock.Mock(
            side_effect=lambda d, r, e, comp: processados.append(e.cnpj) or m.ResultadoMesEmpresa(e, comp, encerrada=True)
        )
        resultado, _ = self._rodar(c, **patches)
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.USER_ENDED_PROCESS)
        self.assertEqual(processados, [CNPJ_A])

    def test_pausa_e_consultada_a_cada_empresa_e_competencia_no_manual(self):
        pausas = []
        c = self._controlador(competencias=[date(2026, 1, 1), date(2026, 2, 1)], callback_pausa=lambda: pausas.append(1))
        self._rodar(c, **self._ok())
        self.assertEqual(len(pausas), 2 + 4)  # 2 empresas + 2 empresas x 2 competências

    def test_alvos_no_manual_pulam_empresas_fora_da_lista_sem_buscar_no_portal(self):
        c = self._controlador(alvos={CNPJ_B: {date(2026, 1, 1)}})
        patches = self._ok()
        resultado, _ = self._rodar(c, **patches)
        buscadas = [chamada.args[1].cnpj for chamada in patches["_procurar_inscricao_empresa"].call_args_list]
        self.assertEqual(buscadas, [CNPJ_B])

    def test_acrescentar_na_planilha_sem_reader_nao_faz_nada(self):
        c = self._controlador()
        c._acrescentar_na_planilha_fiscal(None, date(2026, 1, 1), "X", None)  # não levanta
        c._acrescentar_na_planilha_fiscal(None, None, "S/ INSCRIÇÃO", None)

    def test_mensagem_de_erro_inesperado_cita_o_cnpj(self):
        c = self._controlador(cnpjs=[CNPJ_A])
        patches = self._ok()
        patches["_procurar_inscricao_empresa"] = mock.Mock(side_effect=RuntimeError("x"))
        patches["_recuperar_para_proxima_empresa"] = mock.Mock()
        mensagens = []
        ident = m.logger.add(lambda msg: mensagens.append(msg.record["message"]), level="ERROR")
        try:
            self._rodar(c, **patches)
        finally:
            m.logger.remove(ident)
        self.assertTrue(any("Erro inesperado na empresa 11.222.333/0001-81" in t for t in mensagens))


class TestNomeDaLinhaEPastaManual(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_nome_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)

    def _empresa_manual(self):
        e = m.EmpresaAlvo(cnpj=CNPJ_A)
        e.cnpj_m = "11.222.333/0001-81"
        return e

    def _driver_com_linha(self, texto):
        driver = mock.MagicMock()
        driver.get_driver.return_value.find_elements.return_value = [_ElementoFalso(texto)] if texto is not None else []
        return driver

    def test_nome_vem_da_razao_social_da_linha_sem_cnpj_nem_inscricao(self):
        e = self._empresa_manual()
        m.ControladorEncerramentoISSMultiPeriodo._ler_nome_da_linha(
            self._driver_com_linha("11.222.333/0001-81\n0000075-2\nSOCIEDADE MEDICO CIRURGICO LTDA"), "//xpath", e)
        self.assertEqual(e.nome, "SOCIEDADE MEDICO CIRURGICO LTDA")

    def test_linha_ausente_deixa_o_nome_vazio_sem_levantar(self):
        e = self._empresa_manual()
        with mock.patch.object(m.time, "sleep"):
            m.ControladorEncerramentoISSMultiPeriodo._ler_nome_da_linha(self._driver_com_linha(None), "//xpath", e)
        self.assertEqual(e.nome, "")

    def test_erro_ao_ler_o_nome_nunca_derruba_a_busca(self):
        e = self._empresa_manual()
        driver = mock.MagicMock()
        driver.get_driver.return_value.find_elements.side_effect = RuntimeError("stale")
        m.ControladorEncerramentoISSMultiPeriodo._ler_nome_da_linha(driver, "//xpath", e)
        self.assertEqual(e.nome, "")

    def _controlador(self):
        return _controlador_modo(None, self.pasta / "saida", [date(2025, 10, 1)], cnpjs=[CNPJ_A])

    def test_busca_le_o_nome_so_na_origem_manual(self):
        c = self._controlador()
        driver = mock.MagicMock()
        with mock.patch.object(c, "_garantir_modal_de_troca"), mock.patch.object(c, "_confirmar_alteracao_de_inscricao"), \
             mock.patch.object(c, "_ler_nome_da_linha") as ler:
            manual = self._empresa_manual()
            self.assertTrue(c._procurar_inscricao_empresa(driver, manual))
            ler.assert_called_once()
            ler.reset_mock()
            planilha = _empresa()
            planilha.cnpj_m = "12.345.678/0001-99"
            self.assertTrue(c._procurar_inscricao_empresa(driver, planilha))
            ler.assert_not_called()  # na planilha o nome já veio da própria planilha

    def test_certificado_do_manual_vai_para_pasta_por_cnpj(self):
        c = self._controlador()
        e = self._empresa_manual()
        self.assertEqual(c._diretorio_certificados(e, date(2025, 10, 1)).parts[-2:], ("2025-10", f"CNPJ_{CNPJ_A}"))

    def test_nome_do_arquivo_do_certificado_por_origem(self):
        c = self._controlador()
        driver = mock.MagicMock()
        c._imprimir_declaracao_fechamento_iss(driver, self._empresa_manual(), date(2025, 10, 1))
        self.assertIn(f"CERTIFICADO ISS_CNPJ{CNPJ_A}_", driver.page_to_pdf.call_args.kwargs["file_name"])
        driver2 = mock.MagicMock()
        c._imprimir_declaracao_fechamento_iss(driver2, _empresa(codigo=7), date(2025, 10, 1))
        self.assertIn("CERTIFICADO ISS_EMP7_", driver2.page_to_pdf.call_args.kwargs["file_name"])  # planilha: igual a antes

    def test_certificado_existente_no_manual_e_reconhecido_e_nao_baixa_de_novo(self):
        c = self._controlador()
        e = self._empresa_manual()
        pasta = c._diretorio_certificados(e, date(2025, 10, 1))
        pasta.mkdir(parents=True)
        pdf = pasta / f"CERTIFICADO ISS_CNPJ{CNPJ_A}_x.pdf"
        pdf.write_bytes(b"%PDF conteudo")
        driver = mock.MagicMock()
        driver.find_element.return_value.by_xpath.return_value.get_element.return_value.text = "07/11/2025"
        resultado = c._processar_ja_encerrada(driver, None, e, date(2025, 10, 1), object())
        driver.click.assert_not_called()
        self.assertEqual(resultado.acao, m.ACAO_JA_ENCERRADA_CERT_EXISTENTE)
        self.assertEqual(resultado.caminho_certificado, pdf)


class TestNavegadorIndisponivel(unittest.TestCase):
    def test_reconhece_sessao_perdida(self):
        self.assertTrue(m.navegador_indisponivel(m.InvalidSessionIdException("invalid session id")))
        self.assertTrue(m.navegador_indisponivel(m.NoSuchWindowException("no such window")))
        self.assertTrue(m.navegador_indisponivel(
            m.WebDriverException("disconnected: not connected to DevTools (Session info: chrome=153)")
        ))
        self.assertTrue(m.navegador_indisponivel(ConnectionError("conexão recusada")))

    def test_erros_comuns_de_portal_nao_sao_confundidos_com_navegador_morto(self):
        self.assertFalse(m.navegador_indisponivel(m.NoSuchElementException("sem elemento")))
        self.assertFalse(m.navegador_indisponivel(m.TimeoutException("demorou")))
        self.assertFalse(m.navegador_indisponivel(m.StaleElementReferenceException("velho")))
        self.assertFalse(m.navegador_indisponivel(m.WebDriverException("qualquer outro problema")))
        self.assertFalse(m.navegador_indisponivel(RuntimeError("erro de programação")))


class TestCompetenciaNaoEncontrada(unittest.TestCase):
    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_nao_enc_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199")])
        self.controlador = _controlador(self.caminho, self.pasta / "saida", [date(2025, 5, 1)])

    def test_linha_do_mes_ausente_vira_problema_legivel_sem_nome_de_excecao(self):
        from python_spreadsheet_reader.readers import XLSXReader

        reader = XLSXReader(self.caminho)
        with mock.patch.object(self.controlador, "_garantir_tela_de_lista"), \
             mock.patch.object(self.controlador, "_definir_periodo_consulta"), \
             mock.patch.object(m, "retry_on_exception", side_effect=m.NoSuchElementException("linha não veio")):
            resultado = self.controlador._processar_competencia(mock.MagicMock(), reader, _empresa(linha=2), date(2025, 5, 1))
        reader.close_workbook()

        self.assertFalse(resultado.encerrada)
        self.assertEqual(resultado.problemas, [m.MSG_COMPETENCIA_NAO_ENCONTRADA])

        reader2 = XLSXReader(self.caminho)
        self.assertEqual(reader2.get_cell("Y2").value, "05/2025: COMPETÊNCIA NÃO ENCONTRADA NO PORTAL")
        reader2.close_workbook()


class TestRecuperacaoDeErros(unittest.TestCase):
    """Comportamento de `executar_processo()` quando o navegador morre ou o portal fica numa tela
    desconhecida — sempre com ChromeDriver falso."""

    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_rec_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199"), (2, "EMPRESA B", "98765432000111")])

    def _rodar(self, controlador, driver_falso, **patches):
        classe = mock.Mock(return_value=driver_falso)
        with mock.patch.object(m, "ChromeDriver", classe), \
             mock.patch.object(controlador, "_login_iss_fortaleza"), \
             mock.patch.object(controlador, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(controlador, "_navegar_tela_escrituracao"), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao"):
            with mock.patch.multiple(controlador, **patches):
                return controlador.executar_processo()

    def test_navegador_fechado_para_a_execucao_e_nao_tenta_as_empresas_restantes(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        empresas_tentadas = []

        def procurar(driver, empresa):
            empresas_tentadas.append(empresa.codigo)
            raise m.InvalidSessionIdException("invalid session id: session deleted as the browser has closed the connection")

        driver_falso = mock.Mock(name="ChromeDriverFalso")
        driver_falso.quit_driver.side_effect = m.InvalidSessionIdException("já fechado")  # quit() também falha: não pode estourar
        resultado = self._rodar(controlador, driver_falso, _procurar_inscricao_empresa=mock.Mock(side_effect=procurar))

        self.assertEqual(empresas_tentadas, [1])  # a empresa 2 NÃO foi tentada com o navegador morto
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.ERROR_UNHANDLED_EXCEPTION)
        self.assertIn("navegador", resultado.status.message.lower())
        driver_falso.quit_driver.assert_called_once()

    def test_navegador_fechado_ainda_gera_relatorio_do_que_ja_foi_feito(self):
        if not CAMINHO_TEMPLATE.exists():
            self.skipTest("template do relatório não encontrado")
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])

        def procurar(driver, empresa):
            if empresa.codigo == 2:
                raise m.InvalidSessionIdException("invalid session id")
            return True

        resultado = self._rodar(
            controlador,
            mock.Mock(name="ChromeDriverFalso"),
            _procurar_inscricao_empresa=mock.Mock(side_effect=procurar),
            _processar_competencia=mock.Mock(
                side_effect=lambda driver, reader, empresa, competencia: m.ResultadoMesEmpresa(empresa, competencia, encerrada=True)
            ),
        )

        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.ERROR_UNHANDLED_EXCEPTION)
        self.assertEqual(resultado.processadas, 1)  # só a empresa 1 concluiu
        self.assertEqual(len(resultado.report_paths), 1)

    def test_erro_inesperado_por_empresa_recupera_o_portal_e_segue_para_a_proxima(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        empresas_tentadas = []

        def procurar(driver, empresa):
            empresas_tentadas.append(empresa.codigo)
            if empresa.codigo == 1:
                raise RuntimeError("falha inesperada")
            return True

        recuperar = mock.Mock()
        resultado = self._rodar(
            controlador,
            mock.Mock(name="ChromeDriverFalso"),
            _procurar_inscricao_empresa=mock.Mock(side_effect=procurar),
            _recuperar_para_proxima_empresa=recuperar,
            _processar_competencia=mock.Mock(
                side_effect=lambda driver, reader, empresa, competencia: m.ResultadoMesEmpresa(empresa, competencia, encerrada=True)
            ),
        )

        self.assertEqual(empresas_tentadas, [1, 2])
        recuperar.assert_called_once()
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.SUCCESS)

    def test_se_nao_da_para_recuperar_o_portal_a_execucao_para_com_mensagem_clara(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        empresas_tentadas = []

        def procurar(driver, empresa):
            empresas_tentadas.append(empresa.codigo)
            raise RuntimeError("falha inesperada")

        resultado = self._rodar(
            controlador,
            mock.Mock(name="ChromeDriverFalso"),
            _procurar_inscricao_empresa=mock.Mock(side_effect=procurar),
            _recuperar_para_proxima_empresa=mock.Mock(side_effect=m.RecuperacaoImpossivelException("não consegui voltar à troca de empresa")),
        )

        self.assertEqual(empresas_tentadas, [1])  # parou em vez de gerar falhas em cascata
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.ERROR_UNHANDLED_EXCEPTION)
        self.assertIn("não consegui voltar", resultado.status.message)

    def test_recuperacao_nao_faz_nada_se_a_troca_de_empresa_ja_esta_aberta(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        with mock.patch.object(controlador, "_modal_de_troca_de_inscricao_esta_aberto", return_value=True), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao") as abrir:
            controlador._recuperar_para_proxima_empresa(mock.Mock())
        abrir.assert_not_called()

    def test_recuperacao_abre_a_troca_de_empresa_quando_ela_nao_esta_aberta(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        with mock.patch.object(controlador, "_modal_de_troca_de_inscricao_esta_aberto", return_value=False), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao") as abrir:
            controlador._recuperar_para_proxima_empresa(mock.Mock())
        abrir.assert_called_once()

    def test_recuperacao_que_falha_levanta_recuperacao_impossivel(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        with mock.patch.object(controlador, "_modal_de_troca_de_inscricao_esta_aberto", return_value=False), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao", side_effect=m.NoSuchElementException("sem botão")):
            with self.assertRaises(m.RecuperacaoImpossivelException):
                controlador._recuperar_para_proxima_empresa(mock.Mock())

    def test_recuperacao_com_navegador_morto_vira_navegador_indisponivel(self):
        controlador = _controlador(self.caminho, self.pasta / "saida", [date(2026, 1, 1)])
        with mock.patch.object(controlador, "_modal_de_troca_de_inscricao_esta_aberto", side_effect=m.InvalidSessionIdException("invalid session id")):
            with self.assertRaises(m.NavegadorIndisponivelException):
                controlador._recuperar_para_proxima_empresa(mock.Mock())


class TestExecutarProcessoIntegrado(unittest.TestCase):
    """`executar_processo()` de ponta a ponta, com um ChromeDriver totalmente falso (nunca abre
    navegador) e a leitura/escrita real da planilha e do relatório."""

    def setUp(self):
        self.pasta = Path(tempfile.mkdtemp(prefix="encerr_multi_full_"))
        self.addCleanup(shutil.rmtree, self.pasta, ignore_errors=True)
        self.caminho_planilha = _planilha_fiscal(self.pasta, [(1, "EMPRESA A", "12345678000199")])

    def test_execucao_completa_gera_resultado_e_relatorio_sem_abrir_navegador(self):
        controlador = _controlador(self.caminho_planilha, self.pasta / "saida", [date(2026, 1, 1)])

        driver_falso = mock.Mock(name="ChromeDriverFalso")
        classe_driver_falsa = mock.Mock(return_value=driver_falso)

        with mock.patch.object(m, "ChromeDriver", classe_driver_falsa), \
             mock.patch.object(controlador, "_login_iss_fortaleza"), \
             mock.patch.object(controlador, "_procurar_inscricao_empresa", return_value=True), \
             mock.patch.object(controlador, "_dar_ciencia_nas_mensagens_nao_lidas"), \
             mock.patch.object(controlador, "_navegar_tela_escrituracao"), \
             mock.patch.object(
                 controlador, "_processar_competencia",
                 lambda driver, reader, empresa, competencia: m.ResultadoMesEmpresa(
                     empresa, competencia, encerrada=True, dt_encerramento=date(2026, 1, 5)
                 ),
             ), \
             mock.patch.object(controlador, "_abrir_modal_de_alteracao_de_inscricao"):
            resultado = controlador.executar_processo()

        driver_falso.start_driver.assert_called_once()
        driver_falso.quit_driver.assert_called_once()
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.SUCCESS)
        self.assertEqual((resultado.processadas, resultado.encerradas, resultado.problemas), (1, 1, 0))
        self.assertEqual(len(resultado.report_paths), 1)
        self.assertTrue(resultado.report_paths[0].exists())

    def test_cancelamento_fecha_o_navegador_e_nao_levanta(self):
        controlador = _controlador(self.caminho_planilha, self.pasta / "saida", [date(2026, 1, 1)])

        driver_falso = mock.Mock(name="ChromeDriverFalso")
        classe_driver_falsa = mock.Mock(return_value=driver_falso)

        def login_e_cancela(driver):
            raise m.UserStoppedThreadException()

        with mock.patch.object(m, "ChromeDriver", classe_driver_falsa), \
             mock.patch.object(controlador, "_login_iss_fortaleza", side_effect=login_e_cancela):
            resultado = controlador.executar_processo()

        driver_falso.quit_driver.assert_called_once()
        self.assertEqual(resultado.status.codigo, CodigoEncerramentoISS.USER_ENDED_PROCESS)
        self.assertEqual(resultado.processadas, 0)


if __name__ == "__main__":
    unittest.main()
