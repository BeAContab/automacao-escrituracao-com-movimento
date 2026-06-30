from __future__ import annotations

import argparse
import sys
import re
import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    import fitz  # PyMuPDF
except ImportError:  # pragma: no cover - fallback local
    fitz = None

from PyPDF2 import PdfReader

from iss_fortaleza_automacao import (
    executar_automacao_iss,
    interpretar_competencia,
    solicitar_competencia,
)
from cnae_final import enriquecer_registros_cnae_final
from gemini_extracao import (
    extrair_campos_gemini,
    extrair_campos_gemini_somente_texto,
    extrair_campos_groq,
    ErroCotaGemini,
    ErroCotaGroq,
)
from tratamento_erros import registrar_erro, registrar_evento_execucao, configurar_pasta_logs

def _eh_documento_ou_telefone(valor: str) -> bool:
    """Retorna True se o valor tiver características ou tamanho de CPF, CNPJ ou telefone."""
    if not valor:
        return False
    # Remove tudo exceto dígitos
    digitos = re.sub(r"\D+", "", valor)
    
    # CNPJ (14 dígitos)
    if len(digitos) == 14:
        return True
        
    # CPF (11 dígitos) ou celular com DDD (11 dígitos)
    if len(digitos) == 11:
        return True
        
    # Telefone fixo com DDD (10 dígitos)
    if len(digitos) == 10:
        return True
        
    # Telefone sem DDD (8 ou 9 dígitos)
    if len(digitos) in (8, 9):
        return True
        
    # Outras validações por expressões regulares
    if re.search(r"\d{3}\.\d{3}\.\d{3}-\d{2}", valor): # CPF
        return True
    if re.search(r"\d{2}\.\d{3}\.\d{3}/\d{4}-\d{2}", valor): # CNPJ
        return True
    if re.search(r"\(\d{2}\)\s?\d{4,5}-\d{4}", valor): # Telefone formatado
        return True
        
    return False


@dataclass
class NotaFiscalExtraida:
    arquivo_pdf: str
    cnpj_prestador: str = ""
    numero_nf: str = ""
    data_emissao: str = ""
    id_cnae: str = ""
    desc_cnae: str = ""
    descricao_servico: str = ""
    uf_local_prestacao: str = ""
    cidade_local_prestacao: str = ""
    natureza_operacao: str = ""
    iss_retido: str = ""
    valor_servico: str = ""
    aliquota: str = ""
    id_cnae_final: str = ""
    desc_cnae_final: str = ""
    # Novos campos
    prefeitura: str = ""
    valor_deducoes: str = "0,00"
    descontos_incondicionados: str = "0,00"
    descontos_condicionados: str = "0,00"
    outras_retencoes: str = "0,00"
    ir: str = "0,00"
    pis_nao_retido: str = "0,00"
    cofins_nao_retido: str = "0,00"
    csrf: str = "0,00"
    inss: str = "0,00"
    modelo_ia: str = ""

    def __post_init__(self) -> None:
        # Garante que o ID CNAE sempre contenha numeração e não seja CPF, CNPJ ou telefone
        if self.id_cnae:
            if not any(c.isdigit() for c in self.id_cnae) or _eh_documento_ou_telefone(self.id_cnae):
                self.id_cnae = ""
        if self.id_cnae_final:
            if not any(c.isdigit() for c in self.id_cnae_final) or _eh_documento_ou_telefone(self.id_cnae_final):
                self.id_cnae_final = ""

    def como_linha(self) -> list[str]:
        return [
            self.arquivo_pdf,
            self.prefeitura,
            self.cnpj_prestador,
            self.numero_nf,
            self.data_emissao,
            self.id_cnae,
            self.desc_cnae,
            self.descricao_servico,
            self.uf_local_prestacao,
            self.cidade_local_prestacao,
            self.natureza_operacao,
            self.iss_retido,
            self.valor_servico,
            self.valor_deducoes,
            self.descontos_incondicionados,
            self.descontos_condicionados,
            self.outras_retencoes,
            self.ir,
            self.pis_nao_retido,
            self.cofins_nao_retido,
            self.csrf,
            self.inss,
            self.aliquota,
            self.id_cnae_final,
            self.desc_cnae_final,
            self.modelo_ia,
        ]

    def campos_vazios(self) -> list[str]:
        """Lista as colunas que não foram preenchidas para esta NF."""
        campos = {
            "PREFEITURA": self.prefeitura,
            "CNPJ_PRESTADOR": self.cnpj_prestador,
            "NUMERO_NF": self.numero_nf,
            "DATA_EMISSAO": self.data_emissao,
            "ID_CNAE": self.id_cnae,
            "DESC_CNAE": self.desc_cnae,
            "DESCRICAO_SERVICO": self.descricao_servico,
            "UF_LOCAL_PRESTACAO": self.uf_local_prestacao,
            "CIDADE_LOCAL_PRESTACAO": self.cidade_local_prestacao,
            "NATUREZA_OPERACAO": self.natureza_operacao,
            "ISS_RETIDO": self.iss_retido,
            "VALOR_SERVICO": self.valor_servico,
            "ALIQUOTA": self.aliquota,
        }
        return [nome for nome, valor in campos.items() if not str(valor).strip()]

    def resumo_campos(self) -> str:
        """Devolve um resumo curto para facilitar o diagnóstico no log."""
        return (
            f"Prefeitura={self.prefeitura or 'vazio'} | "
            f"CNPJ={self.cnpj_prestador or 'vazio'} | "
            f"NF={self.numero_nf or 'vazio'} | "
            f"Data={self.data_emissao or 'vazio'} | "
            f"Valor={self.valor_servico or 'vazio'}"
        )


def normalizar_texto(texto: str) -> str:
    """Consolida espaços para facilitar a extração por expressão regular."""
    return " ".join(texto.split())


def texto_para_busca(texto: str) -> str:
    """Remove acentos mantendo o tamanho do texto para preservar os índices.

    Os PDFs das três prefeituras usam rótulos parecidos, mas com acentos,
    espaços e quebras diferentes. Esta versão facilita a busca sem perder a
    possibilidade de recortar o conteúdo original.
    """
    equivalencias = {"º": "O", "ª": "A", "–": "-", "—": "-", "\u00a0": " "}
    caracteres: list[str] = []

    for caractere in texto:
        caractere = equivalencias.get(caractere, caractere)
        decomposicao = unicodedata.normalize("NFKD", caractere)
        base = "".join(c for c in decomposicao if not unicodedata.combining(c))
        caracteres.append((base[:1] or " ").upper())

    return "".join(caracteres)


def ler_texto_pdf(caminho_pdf: Path) -> str:
    """Lê o texto de todas as páginas do PDF."""
    if fitz is not None:
        documento = fitz.open(str(caminho_pdf))
        partes = [pagina.get_text("text") or "" for pagina in documento]
        return normalizar_texto("\n".join(partes))

    leitor = PdfReader(str(caminho_pdf))
    partes: list[str] = []
    for pagina in leitor.pages:
        partes.append(pagina.extract_text() or "")
    return normalizar_texto(" ".join(partes))
















def limpar_cnpj(cnpj: str) -> str:
    return re.sub(r"\s+", "", cnpj)


def normalizar_numero_nf(numero: str) -> str:
    """Remove zeros à esquerda sem alterar números já normalizados."""
    numero = re.sub(r"\D+", "", numero or "")
    if not numero:
        return ""
    return str(int(numero))



def data_valida_recente(data_str: str) -> bool:
    if not data_str:
        return False
    try:
        from datetime import datetime
        from datetime import timedelta
        data_parsed = datetime.strptime(data_str, "%d/%m/%Y")
        limite_inferior = datetime.now() - timedelta(days=90)
        return data_parsed >= limite_inferior
    except Exception:
        return False

def normalizar_data_br(valor: str) -> str:
    """Retorna apenas a data DD/MM/AAAA, desde que não seja muito antiga."""
    if not valor:
        return ""
    encontrado = re.search(r"\d{2}/\d{2}/\d{4}", valor)
    if encontrado:
        data_str = encontrado.group(0)
        if data_valida_recente(data_str):
            return data_str
    return ""





































def formatar_moeda_br(valor: Decimal) -> str:
    valor = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    texto = f"{valor:,.2f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")


def formatar_percentual(valor: Decimal) -> str:
    valor = Decimal(valor).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    texto = f"{valor:,.2f}"
    return texto.replace(",", "X").replace(".", ",").replace("X", ".")
















def determinar_natureza_operacao(cidade_local_prestacao: str) -> str:
    return (
        "Tributação no Município"
        if cidade_local_prestacao.strip().lower() == "fortaleza"
        else "Tributação Fora do Município"
    )









def extrair_nota_fiscal(
    caminho_pdf: Path,
    usar_gemini: bool = True,
    usar_groq: bool = True,
    gemini_key: str | None = None,
    groq_key: str | None = None,
) -> NotaFiscalExtraida:
    texto = ler_texto_pdf(caminho_pdf)
    
    dados_ia = {}
    modelo_ia = ""
    
    if usar_gemini:
        try:
            dados_ia = extrair_campos_gemini(caminho_pdf, texto, gemini_key)
            if dados_ia:
                modelo_ia = "Gemini"
        except ErroCotaGemini as exc:
            if not usar_groq:
                raise exc
        except Exception:
            pass

    if not dados_ia and usar_groq:
        dados_ia = extrair_campos_groq(texto, groq_key)
        if dados_ia:
            modelo_ia = "Groq"

    if not dados_ia:
        raise ValueError("O documento não pôde ser analisado por nenhuma IA (Gemini ou Groq).")
    
    prefeitura = dados_ia.get("prefeitura", "")
    cnpj_prestador = dados_ia.get("cnpj_prestador", "")
    numero_nf = dados_ia.get("numero_nf", "")
    data_emissao = dados_ia.get("data_emissao", "")
    id_cnae = dados_ia.get("id_cnae", "")
    desc_cnae = dados_ia.get("desc_cnae", "")
    descricao_servico = dados_ia.get("descricao_servico", "")
    uf_local_prestacao = dados_ia.get("uf_local_prestacao", "")
    cidade_local_prestacao = dados_ia.get("cidade_local_prestacao", "")
    iss_retido = dados_ia.get("iss_retido", "").strip().upper()
    valor_servico = dados_ia.get("valor_servico", "")
    aliquota = dados_ia.get("aliquota", "")
    valor_deducoes = dados_ia.get("valor_deducoes", "0,00")
    descontos_incondicionados = dados_ia.get("descontos_incondicionados", "0,00")
    descontos_condicionados = dados_ia.get("descontos_condicionados", "0,00")
    outras_retencoes = dados_ia.get("outras_retencoes", "0,00")
    ir = dados_ia.get("ir", "0,00")
    pis_nao_retido = dados_ia.get("pis_nao_retido", "0,00")
    cofins_nao_retido = dados_ia.get("cofins_nao_retido", "0,00")
    csrf = dados_ia.get("csrf", "0,00")
    inss = dados_ia.get("inss", "0,00")

    # Anti-alucinação: validar se os dígitos de CNPJ e Número da NF constam no texto bruto
    digitos_texto = "".join(c for c in texto if c.isdigit())
    
    digitos_cnpj = "".join(c for c in cnpj_prestador if c.isdigit())
    if digitos_cnpj and digitos_cnpj not in digitos_texto:
        cnpj_prestador = ""
        
    digitos_nf = "".join(c for c in numero_nf if c.isdigit())
    if digitos_nf and digitos_nf not in digitos_texto:
        numero_nf = ""

    natureza_operacao = determinar_natureza_operacao(cidade_local_prestacao)

    return NotaFiscalExtraida(
        arquivo_pdf=caminho_pdf.name,
        prefeitura=prefeitura,
        cnpj_prestador=limpar_cnpj(cnpj_prestador),
        numero_nf=normalizar_numero_nf(numero_nf),
        data_emissao=normalizar_data_br(data_emissao),
        id_cnae=id_cnae,
        desc_cnae=desc_cnae,
        descricao_servico=descricao_servico,
        uf_local_prestacao=uf_local_prestacao,
        cidade_local_prestacao=cidade_local_prestacao,
        natureza_operacao=natureza_operacao,
        iss_retido=iss_retido,
        valor_servico=valor_servico,
        aliquota=aliquota,
        valor_deducoes=valor_deducoes,
        descontos_incondicionados=descontos_incondicionados,
        descontos_condicionados=descontos_condicionados,
        outras_retencoes=outras_retencoes,
        ir=ir,
        pis_nao_retido=pis_nao_retido,
        cofins_nao_retido=cofins_nao_retido,
        csrf=csrf,
        inss=inss,
        modelo_ia=modelo_ia,
    )


def extrair_nota_fiscal_somente_texto(
    caminho_pdf: Path,
    usar_gemini: bool = True,
    usar_groq: bool = True,
    gemini_key: str | None = None,
    groq_key: str | None = None,
) -> NotaFiscalExtraida:
    """Lê o texto do PDF localmente e envia apenas o texto para a IA (Função 3)."""
    texto = ler_texto_pdf(caminho_pdf)
    
    dados_ia = {}
    modelo_ia = ""
    
    if usar_gemini:
        try:
            dados_ia = extrair_campos_gemini_somente_texto(caminho_pdf, texto, gemini_key)
            if dados_ia:
                modelo_ia = "Gemini"
        except ErroCotaGemini as exc:
            if not usar_groq:
                raise exc
        except Exception:
            pass

    if not dados_ia and usar_groq:
        dados_ia = extrair_campos_groq(texto, groq_key)
        if dados_ia:
            modelo_ia = "Groq"

    if not dados_ia:
        raise ValueError("O documento não pôde ser analisado por nenhuma IA (Gemini ou Groq).")
    
    prefeitura = dados_ia.get("prefeitura", "")
    cnpj_prestador = dados_ia.get("cnpj_prestador", "")
    numero_nf = dados_ia.get("numero_nf", "")
    data_emissao = dados_ia.get("data_emissao", "")
    id_cnae = dados_ia.get("id_cnae", "")
    desc_cnae = dados_ia.get("desc_cnae", "")
    descricao_servico = dados_ia.get("descricao_servico", "")
    uf_local_prestacao = dados_ia.get("uf_local_prestacao", "")
    cidade_local_prestacao = dados_ia.get("cidade_local_prestacao", "")
    iss_retido = dados_ia.get("iss_retido", "").strip().upper()
    valor_servico = dados_ia.get("valor_servico", "")
    aliquota = dados_ia.get("aliquota", "")
    valor_deducoes = dados_ia.get("valor_deducoes", "0,00")
    descontos_incondicionados = dados_ia.get("descontos_incondicionados", "0,00")
    descontos_condicionados = dados_ia.get("descontos_condicionados", "0,00")
    outras_retencoes = dados_ia.get("outras_retencoes", "0,00")
    ir = dados_ia.get("ir", "0,00")
    pis_nao_retido = dados_ia.get("pis_nao_retido", "0,00")
    cofins_nao_retido = dados_ia.get("cofins_nao_retido", "0,00")
    csrf = dados_ia.get("csrf", "0,00")
    inss = dados_ia.get("inss", "0,00")

    # Anti-alucinação: validar se os dígitos de CNPJ e Número da NF constam no texto bruto
    digitos_texto = "".join(c for c in texto if c.isdigit())
    
    digitos_cnpj = "".join(c for c in cnpj_prestador if c.isdigit())
    if digitos_cnpj and digitos_cnpj not in digitos_texto:
        cnpj_prestador = ""
        
    digitos_nf = "".join(c for c in numero_nf if c.isdigit())
    if digitos_nf and digitos_nf not in digitos_texto:
        numero_nf = ""

    natureza_operacao = determinar_natureza_operacao(cidade_local_prestacao)

    return NotaFiscalExtraida(
        arquivo_pdf=caminho_pdf.name,
        prefeitura=prefeitura,
        cnpj_prestador=limpar_cnpj(cnpj_prestador),
        numero_nf=normalizar_numero_nf(numero_nf),
        data_emissao=normalizar_data_br(data_emissao),
        id_cnae=id_cnae,
        desc_cnae=desc_cnae,
        descricao_servico=descricao_servico,
        uf_local_prestacao=uf_local_prestacao,
        cidade_local_prestacao=cidade_local_prestacao,
        natureza_operacao=natureza_operacao,
        iss_retido=iss_retido,
        valor_servico=valor_servico,
        aliquota=aliquota,
        valor_deducoes=valor_deducoes,
        descontos_incondicionados=descontos_incondicionados,
        descontos_condicionados=descontos_condicionados,
        outras_retencoes=outras_retencoes,
        ir=ir,
        pis_nao_retido=pis_nao_retido,
        cofins_nao_retido=cofins_nao_retido,
        csrf=csrf,
        inss=inss,
        modelo_ia=modelo_ia,
    )



def listar_pdfs(pasta: Path) -> Iterable[Path]:
    return sorted(pasta.glob("*.pdf"))


def ajustar_largura_colunas(planilha) -> None:
    for coluna in planilha.columns:
        maior = 0
        letra_coluna = get_column_letter(coluna[0].column)
        for celula in coluna:
            if celula.value is not None:
                maior = max(maior, len(str(celula.value)))
        planilha.column_dimensions[letra_coluna].width = min(maior + 2, 80)


def _converter_valor_br_para_float(valor_str: str) -> float:
    """Converte valores monetários formatados no padrão BR para float."""
    if not valor_str:
        return 0.0
    try:
        limpo = str(valor_str).replace("R$", "").strip()
        if "," in limpo:
            limpo = limpo.replace(".", "").replace(",", ".")
        return float(limpo)
    except ValueError:
        return 0.0


def _aliquota_diferente_de_5(aliquota_str: str) -> bool:
    """Retorna True se a alíquota for diferente de 5% (com tolerância decimal)."""
    if not aliquota_str:
        return True  # Alíquota ausente/vazia é diferente de 5%
    val = _converter_valor_br_para_float(aliquota_str)
    # Se for menor que 1.0 (ex: 0.05 em vez de 5.0), multiplicamos por 100
    if 0.0 < val < 1.0:
        val = val * 100.0
    return abs(val - 5.0) > 0.01


def gerar_xlsx(registros: list[NotaFiscalExtraida], destino: Path) -> None:
    # Enriquecemos os registros com o CNAE oficial antes de gravar a planilha final.
    enriquecer_registros_cnae_final(registros)

    wb = Workbook()
    ws = wb.active
    ws.title = "NF Extraídas"

    cabecalhos = [
        "ARQUIVO_PDF",
        "PREFEITURA",
        "CNPJ_PRESTADOR",
        "NUMERO_NF",
        "DATA_EMISSAO",
        "ID_CNAE",
        "DESC_CNAE",
        "DESCRICAO_SERVICO",
        "UF_LOCAL_PRESTACAO",
        "CIDADE_LOCAL_PRESTACAO",
        "NATUREZA_OPERACAO",
        "ISS_RETIDO",
        "VALOR_SERVICO",
        "VALOR_DEDUCOES",
        "DESCONTOS_INCONDICIONADOS",
        "DESCONTOS_CONDICIONADOS",
        "OUTRAS_RETENCOES",
        "IR",
        "PIS_NAO_RETIDO",
        "COFINS_NAO_RETIDO",
        "CSRF (CSLL + PIS + Cofins Retidos)",
        "INSS",
        "ALIQUOTA",
        "ID_CNAE_FINAL",
        "DESC_CNAE_FINAL",
        "MODELO_IA",
    ]

    ws.append(cabecalhos)
    for registro in registros:
        ws.append(registro.como_linha())

    titulos_vermelhos = {
        "CNPJ_PRESTADOR",
        "NUMERO_NF",
        "DATA_EMISSAO",
        "DESCRICAO_SERVICO",
        "UF_LOCAL_PRESTACAO",
        "CIDADE_LOCAL_PRESTACAO",
        "NATUREZA_OPERACAO",
        "ISS_RETIDO",
        "VALOR_SERVICO",
        "VALOR_DEDUCOES",
        "DESCONTOS_INCONDICIONADOS",
        "DESCONTOS_CONDICIONADOS",
        "OUTRAS_RETENCOES",
        "IR",
        "PIS_NAO_RETIDO",
        "COFINS_NAO_RETIDO",
        "CSRF (CSLL + PIS + Cofins Retidos)",
        "INSS",
        "ALIQUOTA",
        "ID_CNAE_FINAL",
        "DESC_CNAE_FINAL",
    }

    for celula in ws[1]:
        celula.font = Font(bold=True, color="FFFFFF")
        if celula.value in titulos_vermelhos:
            celula.fill = PatternFill("solid", fgColor="C00000") # Vermelho escuro profissional
        else:
            celula.fill = PatternFill("solid", fgColor="1F4E78") # Azul original

    # Mapeamento dinâmico de colunas para formatação condicional
    mapa_colunas = {nome: idx + 1 for idx, nome in enumerate(cabecalhos)}
    
    colunas_contabeis = {
        "VALOR_SERVICO",
        "VALOR_DEDUCOES",
        "DESCONTOS_INCONDICIONADOS",
        "DESCONTOS_CONDICIONADOS",
        "OUTRAS_RETENCOES",
        "IR",
        "PIS_NAO_RETIDO",
        "COFINS_NAO_RETIDO",
        "CSRF (CSLL + PIS + Cofins Retidos)",
        "INSS",
    }
    indices_contabeis = [mapa_colunas[c] for c in colunas_contabeis if c in mapa_colunas]
    indice_aliquota = mapa_colunas.get("ALIQUOTA")
    
    formato_contabil = '_("R$"* #,##0.00_);_("R$"* (#,##0.00);_("R$"* "-"??_);_(@_)'
    
    # Primeiro passo: converter valores monetários e alíquota de texto para numérico com formatação adequada
    for row_idx in range(2, ws.max_row + 1):
        for col_idx in indices_contabeis:
            celula = ws.cell(row=row_idx, column=col_idx)
            val_float = _converter_valor_br_para_float(celula.value)
            celula.value = val_float
            celula.number_format = formato_contabil

        if indice_aliquota:
            celula = ws.cell(row=row_idx, column=indice_aliquota)
            val_float = _converter_valor_br_para_float(celula.value)
            if val_float > 1.0:
                val_float = val_float / 100.0
            celula.value = val_float
            celula.number_format = '0.00%'

    colunas_valores_retidos = [
        "VALOR_DEDUCOES",
        "DESCONTOS_INCONDICIONADOS",
        "DESCONTOS_CONDICIONADOS",
        "OUTRAS_RETENCOES",
        "IR",
        "PIS_NAO_RETIDO",
        "COFINS_NAO_RETIDO",
        "CSRF (CSLL + PIS + Cofins Retidos)",
        "INSS"
    ]
    indices_valores_retidos = [mapa_colunas[c] for c in colunas_valores_retidos if c in mapa_colunas]
    
    preenchimento_vermelho_alerta = PatternFill("solid", fgColor="FFC7CE") # Vermelho claro/pastel
    fonte_vermelha_alerta = Font(color="9C0006") # Vermelho escuro para contraste
    
    # Segundo passo: aplica coloração condicional de dados
    for row_idx in range(2, ws.max_row + 1):
        # Alerta se impostos ou deduções forem maiores que zero
        for col_idx in indices_valores_retidos:
            celula = ws.cell(row=row_idx, column=col_idx)
            val = celula.value
            if not isinstance(val, (int, float)):
                val = _converter_valor_br_para_float(val)
            if val > 0.01:
                celula.fill = preenchimento_vermelho_alerta
                celula.font = fonte_vermelha_alerta
                
        # Alerta se alíquota for diferente de 5%
        if indice_aliquota:
            celula = ws.cell(row=row_idx, column=indice_aliquota)
            val = celula.value
            if not isinstance(val, (int, float)):
                val = _converter_valor_br_para_float(val)
            # Converte para base percentual cheia se for decimal puro
            if 0.0 < val < 1.0:
                val = val * 100.0
            if abs(val - 5.0) > 0.01:
                celula.fill = preenchimento_vermelho_alerta
                celula.font = fonte_vermelha_alerta

    ajustar_largura_colunas(ws)
    destino.parent.mkdir(parents=True, exist_ok=True)
    wb.save(destino)


def registrar_log_incompletos(
    registros: list[NotaFiscalExtraida],
    destino_xlsx: Path,
    origem_dir: Path | None = None,
) -> Path:
    """Cria um log detalhado na pasta 'log' quando alguma NF não preencher todas as colunas.

    O objetivo é permitir auditoria rápida dos arquivos que precisam de ajuste
    manual ou de novas regras de extração.
    """
    linhas_log: list[str] = []

    for registro in registros:
        campos_vazios = registro.campos_vazios()
        if not campos_vazios:
            continue

        linhas_log.append(
            f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {registro.arquivo_pdf}"
        )
        linhas_log.append(f"Campos ausentes: {', '.join(campos_vazios)}")
        linhas_log.append(f"Resumo: {registro.resumo_campos()}")
        linhas_log.append("")

    caminho_log = destino_xlsx.parent / "log" / f"{destino_xlsx.stem}_log_extracao.txt"
    caminho_log.parent.mkdir(parents=True, exist_ok=True)
    
    if linhas_log:
        caminho_log.write_text("\n".join(linhas_log), encoding="utf-8")
    elif caminho_log.exists():
        try:
            caminho_log.unlink()
        except Exception:
            pass

    if origem_dir:
        caminho_log_origem = origem_dir / "log" / f"{destino_xlsx.stem}_log_extracao.txt"
        try:
            caminho_log_origem.parent.mkdir(parents=True, exist_ok=True)
            if linhas_log:
                caminho_log_origem.write_text("\n".join(linhas_log), encoding="utf-8")
            elif caminho_log_origem.exists():
                caminho_log_origem.unlink()
        except Exception:
            pass

    return caminho_log


def registrar_incompleto_realtime(
    registro: NotaFiscalExtraida,
    destino_xlsx: Path,
    origem_dir: Path | None = None,
) -> None:
    """Registra uma NF incompleta no arquivo de log em tempo real (append)."""
    campos_vazios = registro.campos_vazios()
    if not campos_vazios:
        return

    linhas_log = [
        f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] {registro.arquivo_pdf}",
        f"Campos ausentes: {', '.join(campos_vazios)}",
        f"Resumo: {registro.resumo_campos()}",
        ""
    ]
    texto = "\n".join(linhas_log) + "\n"

    # Salva na pasta log do destino
    caminho_log = destino_xlsx.parent / "log" / f"{destino_xlsx.stem}_log_extracao.txt"
    try:
        caminho_log.parent.mkdir(parents=True, exist_ok=True)
        with open(caminho_log, "a", encoding="utf-8") as f:
            f.write(texto)
    except Exception:
        pass

    # Salva na pasta log da origem se fornecida
    if origem_dir:
        caminho_log_origem = origem_dir / "log" / f"{destino_xlsx.stem}_log_extracao.txt"
        try:
            caminho_log_origem.parent.mkdir(parents=True, exist_ok=True)
            with open(caminho_log_origem, "a", encoding="utf-8") as f:
                f.write(texto)
        except Exception:
            pass


def abrir_site_iss_fortaleza() -> None:
    """Mantido apenas por compatibilidade histórica com versões anteriores."""
    pass



def solicitar_opcao() -> str:
    """Pede ao usuário que escolha entre gerar o XLSX (Apenas IA), a automação ou a extração otimizada por texto."""
    while True:
        print("Escolha uma opção:")
        print("1 - Gerar o XLSX a partir dos PDFs (Apenas IA)")
        print("2 - Executar a automação visível da ISS de Fortaleza")
        print("3 - Gerar o XLSX a partir dos PDFs usando apenas texto local (Função 3 - Otimizada)")
        resposta = input("Digite 1, 2 ou 3: ").strip()
        if resposta in {"1", "2", "3"}:
            return resposta
        print("Opção inválida. Por favor, digite 1, 2 ou 3.")
        print()



def solicitar_planilha_automacao() -> Path:
    """Pergunta ao usuário qual planilha será usada na automação da ISS."""

    entrada = input(
        "Informe o caminho da planilha XLSX que será usada na automação da ISS "
        "(Enter para `nf_compilado.xlsx`): "
    ).strip().strip('"')
    if not entrada:
        return Path("nf_compilado.xlsx")
    return Path(entrada)


def oferecer_proximo_passo_apos_funcao1(destino: Path, args: argparse.Namespace) -> int:
    """Pergunta se o usuário quer seguir para a função 2 ou encerrar o programa.

    A planilha recém-gerada vira a entrada padrão da automação da ISS para
    evitar retrabalho e manter o fluxo contínuo entre as duas etapas.
    """

    if not sys.stdin.isatty():
        print(
            "Próximo passo sugerido: escolha a opção 2 para iniciar a automação visível da ISS de Fortaleza."
        )
        return 0

    while True:
        print()
        print("A função 1 foi concluída com sucesso.")
        print(f"Planilha gerada: {destino.resolve()}")
        print("O que você deseja fazer agora?")
        print("1 - Iniciar a função 2 (automação visível da ISS de Fortaleza)")
        print("2 - Encerrar o programa")
        resposta = input("Digite 1 ou 2: ").strip()

        if resposta == "1":
            registrar_evento_execucao(
                "Usuário optou por iniciar a função 2 após a geração do XLSX",
                "extrair_nf_pdfs.py",
            )
            try:
                competencia = (
                    interpretar_competencia(args.competencia)
                    if args.competencia is not None
                    else solicitar_competencia()
                )
            except ValueError as exc:
                raise SystemExit(f"Competência inválida: {exc}") from exc

            # A planilha gerada na função 1 é a base natural para a automação da ISS.
            executar_automacao_iss(destino, competencia, depuracao=args.debug_funcao2)
            return 0

        if resposta == "2":
            registrar_evento_execucao(
                "Usuário optou por encerrar o programa após a função 1",
                "extrair_nf_pdfs.py",
            )
            print("Programa encerrado.")
            return 0

        print("Opção inválida. Digite 1 para iniciar a função 2 ou 2 para encerrar.")
        print()



def construir_argumentos() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Extrai dados de NFS-e em PDF, gera XLSX ou executa a automação da ISS."
    )
    parser.add_argument(
        "-m",
        "--modo",
        choices=("1", "2", "3"),
        default=None,
        help="1 para gerar o XLSX (Apenas IA); 2 para executar a automação; 3 para extração por texto (Otimizada).",
    )
    parser.add_argument(
        "-c",
        "--competencia",
        default=None,
        help="Competência a trabalhar na ISS, por exemplo '5/2026', '5 2026' ou 'maio 2026'.",
    )
    parser.add_argument(
        "--planilha",
        type=Path,
        default=None,
        help="Planilha XLSX usada pela automação da ISS.",
    )
    parser.add_argument(
        "--debug-funcao2",
        action="store_true",
        help="Abre a função 2 com pausas extras e navegador em modo de teste visível.",
    )
    parser.add_argument(
        "pasta",
        nargs="?",
        type=Path,
        default=None,
        help="Pasta que contém os arquivos PDF. Se omitido, o script solicita no terminal.",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=Path("nf_compilado.xlsx"),
        help="Caminho do arquivo XLSX de saída.",
    )
    return parser.parse_args()



def main() -> int:
    args = construir_argumentos()
    registrar_evento_execucao("Fluxo principal iniciado no extrair_nf_pdfs.py", "extrair_nf_pdfs.py")

    modo = args.modo
    if modo is None:
        if args.pasta is not None:
            modo = "1"
        else:
            modo = solicitar_opcao()

    if modo == "2":
        try:
            competencia = (
                interpretar_competencia(args.competencia)
                if args.competencia is not None
                else solicitar_competencia()
            )
        except ValueError as exc:
            raise SystemExit(f"Competência inválida: {exc}") from exc
        planilha = args.planilha if args.planilha is not None else solicitar_planilha_automacao()
        if not planilha.exists():
            raise SystemExit(
                f"A planilha informada para a automação não existe: {planilha}"
            )
        registrar_evento_execucao(
            f"Modo 2 selecionado com planilha {planilha.resolve()}",
            "extrair_nf_pdfs.py",
        )
        executar_automacao_iss(planilha, competencia, depuracao=args.debug_funcao2)
        return 0

    pasta = args.pasta
    if pasta is None:
        entrada = input("Informe o caminho da pasta com os PDFs: ").strip().strip('"')
        pasta = Path(entrada)

    destino = args.output

    if not pasta.exists() or not pasta.is_dir():
        raise SystemExit(f"A pasta informada não existe ou não é válida: {pasta}")

    configurar_pasta_logs(pasta)

    if not destino.is_absolute():
        destino = (pasta / destino.name).resolve()

    pdfs = list(listar_pdfs(pasta))
    if not pdfs:
        raise SystemExit("Nenhum PDF foi encontrado na pasta informada.")

    registrar_evento_execucao(
        f"Modo {modo} selecionado com {len(pdfs)} PDF(s) em {pasta.resolve()}",
        "extrair_nf_pdfs.py",
    )

    registros: list[NotaFiscalExtraida] = []
    notas_nao_analisadas: list[str] = []
    for indice, pdf in enumerate(pdfs, start=1):
        print(f"[{indice}/{len(pdfs)}] Processando PDF: {pdf.name}", flush=True)
        registrar_evento_execucao(
            f"Iniciando processamento do PDF {indice}/{len(pdfs)}: {pdf.name}",
            "extrair_nf_pdfs.py",
        )
        try:
            if modo == "3":
                registros.append(extrair_nota_fiscal_somente_texto(pdf))
            else:
                registros.append(extrair_nota_fiscal(pdf))
        except ValueError as exc:
            print(f"[{indice}/{len(pdfs)}] Falha ao processar {pdf.name}: {exc}", flush=True)
            registrar_evento_execucao(
                f"Falha de IA dupla para {pdf.name}: {exc}",
                "extrair_nf_pdfs.py",
            )
            notas_nao_analisadas.append(pdf.name)
        except Exception as exc:
            print(f"[{indice}/{len(pdfs)}] Erro inesperado ao processar {pdf.name}: {exc}", flush=True)
            registrar_evento_execucao(
                f"Erro inesperado em {pdf.name}: {exc}",
                "extrair_nf_pdfs.py",
            )
        else:
            registrar_evento_execucao(
                f"Processamento concluído com sucesso para {pdf.name}",
                "extrair_nf_pdfs.py",
            )

    registros_completos = [registro for registro in registros if not registro.campos_vazios()]
    gerar_xlsx(registros_completos, destino)
    caminho_log = registrar_log_incompletos(registros, destino, origem_dir=pasta)
    print(f"Arquivo XLSX gerado com sucesso: {destino.resolve()}")
    print(f"Total de PDFs processados: {len(pdfs)}")
    print(f"Linhas completas exportadas: {len(registros_completos)}")
    if caminho_log.exists():
        print(f"Log de extração gerado em: {caminho_log.resolve()}")

    if notas_nao_analisadas:
        caminho_log_na = pasta / "log" / "log_nao_analisados.txt"
        try:
            with open(caminho_log_na, "w", encoding="utf-8") as f:
                f.write(f"[{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}] Ocorreram falhas ao tentar extrair dados com as IAs (Gemini e Groq) para as seguintes notas fiscais:\n")
                for nome_pdf in notas_nao_analisadas:
                    f.write(f"- {nome_pdf}\n")
            print(f"Aviso: Algumas notas não puderam ser analisadas. Detalhes gravados em: {caminho_log_na.resolve()}", flush=True)
        except Exception as exc_log:
            registrar_evento_execucao(
                f"Erro ao salvar log {caminho_log_na.name}: {exc_log}",
                "extrair_nf_pdfs.py",
            )

    registrar_evento_execucao(
        f"XLSX gerado em {destino.resolve()} com {len(registros_completos)} linha(s) exportada(s)",
        "extrair_nf_pdfs.py",
    )
    return oferecer_proximo_passo_apos_funcao1(destino, args)


if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except SystemExit as exc:
        if exc.code not in (0, None):
            registrar_erro(exc, "extrair_nf_pdfs.py")
        raise
    except Exception as exc:
        registrar_erro(exc, "extrair_nf_pdfs.py")
        raise

