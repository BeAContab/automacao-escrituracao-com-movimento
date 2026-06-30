"""Classificacao do CNAE oficial para enriquecimento do XLSX da funcao 1."""

from __future__ import annotations

import json
import os
import re
import unicodedata
from dataclasses import dataclass
from difflib import SequenceMatcher
from functools import lru_cache
from pathlib import Path
from typing import Any
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

from openpyxl import load_workbook

try:
    from google import genai
    from google.genai import types
except ImportError:  # pragma: no cover
    genai = None
    types = None

GEMINI_MODEL_PADRAO = "gemini-2.5-flash"
GEMINI_MODELOS_FALLBACK = [
    "gemini-2.5-pro",
    "gemini-2.5-flash-lite",
]
LIMITE_CANDIDATOS_GEMINI = 10


@dataclass(frozen=True)
class CnaeOficial:
    """Representa uma linha do arquivo oficial de CNAE."""

    codigo: str
    descricao: str
    descricao_normalizada: str
    tokens: frozenset[str]


def _carregar_env_local(caminho: Path = Path(".env")) -> None:
    """Carrega variáveis locais do .env sem depender de bibliotecas extras."""

    if not caminho.exists():
        return

    for linha in caminho.read_text(encoding="utf-8").splitlines():
        texto = linha.strip()
        if not texto or texto.startswith("#") or "=" not in texto:
            continue

        chave, valor = texto.split("=", 1)
        chave = chave.strip()
        valor = valor.strip().strip('"').strip("'")
        if chave and chave not in os.environ:
            os.environ[chave] = valor


def _normalizar_texto(texto: str) -> str:
    """Normaliza acentos, pontuacao e espacos para comparar descricoes."""

    texto = unicodedata.normalize("NFD", texto or "")
    texto = "".join(caractere for caractere in texto if unicodedata.category(caractere) != "Mn")
    texto = re.sub(r"[^0-9A-Za-z]+", " ", texto)
    return " ".join(texto.lower().split())


def _tokenizar(texto: str) -> frozenset[str]:
    """Separa a descricao em tokens uteis para pontuacao local."""

    palavras = {
        "de",
        "da",
        "do",
        "das",
        "dos",
        "e",
        "em",
        "para",
        "com",
        "sem",
        "ou",
        "a",
        "o",
        "as",
        "os",
        "na",
        "no",
        "nas",
        "nos",
        "servico",
        "servicos",
        "atividade",
        "atividades",
        "outro",
        "outros",
        "especificado",
        "especificados",
        "anteriormente",
    }
    tokens = [token for token in _normalizar_texto(texto).split() if token not in palavras]
    return frozenset(tokens)


IDS_QUE_FORCAM_ISS_RETIDO = {"932989910", "900190201"}
CODIGO_CNAE_EVENTOS = "932989910"


@lru_cache(maxsize=1)
def carregar_cnaes_oficiais(caminho: str = "cnae_oficial.xlsx") -> tuple[CnaeOficial, ...]:
    """Lê o arquivo oficial de CNAE e devolve a tabela pronta para busca."""
    import sys
    arquivo = Path(caminho)
    if not arquivo.exists():
        if hasattr(sys, '_MEIPASS'):
            arquivo = Path(sys._MEIPASS) / caminho
        else:
            arquivo = Path(__file__).parent / caminho

    if not arquivo.exists():
        raise FileNotFoundError(f"Arquivo oficial de CNAE não encontrado: {arquivo}")

    workbook = load_workbook(arquivo, data_only=True)
    planilha = workbook.active
    cabecalhos = [str(celula.value or "").strip() for celula in next(planilha.iter_rows(min_row=1, max_row=1))]
    mapa = {cabecalho.lower(): indice for indice, cabecalho in enumerate(cabecalhos)}

    indice_codigo = mapa.get("cnae", 0)
    indice_descricao = mapa.get("descricao", 1)

    cnaes: list[CnaeOficial] = []
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if not linha or indice_codigo >= len(linha) or indice_descricao >= len(linha):
            continue

        codigo = str(linha[indice_codigo] or "").strip()
        descricao = str(linha[indice_descricao] or "").strip()
        if not codigo or not descricao:
            continue

        descricao_normalizada = _normalizar_texto(descricao)
        tokens = _tokenizar(descricao)
        cnaes.append(
            CnaeOficial(
                codigo=codigo,
                descricao=descricao,
                descricao_normalizada=descricao_normalizada,
                tokens=tokens,
            )
        )

    return tuple(cnaes)


def _pontuar_texto(consulta: str, candidata: CnaeOficial) -> float:
    """Pontua uma descricao oficial contra o texto enviado pelo PDF."""

    consulta_normalizada = _normalizar_texto(consulta)
    if not consulta_normalizada:
        return 0.0

    score = 0.0
    if consulta_normalizada == candidata.descricao_normalizada:
        score += 120.0
    if consulta_normalizada in candidata.descricao_normalizada:
        score += 60.0
    if candidata.descricao_normalizada in consulta_normalizada:
        score += 40.0

    tokens_consulta = _tokenizar(consulta)
    interseccao = len(tokens_consulta & candidata.tokens)
    score += interseccao * 10.0

    score += SequenceMatcher(None, consulta_normalizada, candidata.descricao_normalizada).ratio() * 25.0
    return score


def _pontuar_candidato(desc_cnae: str, descricao_servico: str, id_cnae: str, candidata: CnaeOficial) -> float:
    """Combina o texto do PDF e a descrição do serviço para ranquear o CNAE."""

    score = 0.0
    score += _pontuar_texto(desc_cnae, candidata) * 0.75
    score += _pontuar_texto(descricao_servico, candidata) * 0.25

    codigo_extraido = re.sub(r"\D+", "", id_cnae or "")
    codigo_oficial = re.sub(r"\D+", "", candidata.codigo)
    if codigo_extraido and codigo_extraido == codigo_oficial:
        score += 80.0
    elif codigo_extraido and codigo_extraido in codigo_oficial:
        score += 20.0

    return score


def _selecionar_candidatos(
    desc_cnae: str,
    descricao_servico: str,
    id_cnae: str,
    oficiais: tuple[CnaeOficial, ...],
    limite: int = LIMITE_CANDIDATOS_GEMINI,
) -> list[CnaeOficial]:
    """Escolhe um subconjunto pequeno de CNAEs para a IA analisar com mais foco."""

    if not oficiais:
        return []

    ranqueados = sorted(
        oficiais,
        key=lambda candidata: _pontuar_candidato(desc_cnae, descricao_servico, id_cnae, candidata),
        reverse=True,
    )
    return ranqueados[:limite]


def _obter_chave_gemini() -> str:
    """Lê a chave do Gemini do ambiente ou do arquivo .env local."""

    _carregar_env_local()
    return (
        os.getenv("GEMINI_API_KEY", "").strip()
        or os.getenv("GOOGLE_API_KEY", "").strip()
    )


def _obter_modelo_gemini() -> str:
    """Retorna o modelo padrão do Gemini para a classificação de CNAE."""

    _carregar_env_local()
    return os.getenv("GEMINI_MODEL", GEMINI_MODEL_PADRAO).strip() or GEMINI_MODEL_PADRAO


def _modelos_candidatos_gemini() -> list[str]:
    """Lista os modelos candidatos do Gemini para classificação de CNAE."""

    candidatos = [_obter_modelo_gemini()]
    for modelo in GEMINI_MODELOS_FALLBACK:
        if modelo not in candidatos:
            candidatos.append(modelo)
    return candidatos


def _normalizar_codigo_cnae(codigo: str) -> str:
    """Remove pontuação do código para facilitar regras exatas de negócio."""

    return re.sub(r"\D+", "", codigo or "")


def _pontuar_descricao_oficial(desc_cnae: str, candidata: CnaeOficial) -> float:
    """Mete o quanto a descrição extraída parece a descrição oficial da CNAE."""

    score = _pontuar_texto(desc_cnae, candidata)
    descricao_normalizada = _normalizar_texto(desc_cnae)
    if descricao_normalizada == candidata.descricao_normalizada:
        score += 200.0
    elif descricao_normalizada in candidata.descricao_normalizada:
        score += 120.0
    elif candidata.descricao_normalizada in descricao_normalizada:
        score += 90.0
    return score


def _resolver_por_descricao_explicita(
    desc_cnae: str,
    oficiais: tuple[CnaeOficial, ...],
) -> tuple[str, str] | None:
    """Tenta resolver o CNAE apenas pela descrição visível no PDF."""

    if not desc_cnae.strip() or not oficiais:
        return None

    tokens_desc = _tokenizar(desc_cnae)
    if {"organizacao", "feiras", "congressos"}.issubset(tokens_desc) or {"organizacao", "feiras", "exposicoes"}.issubset(tokens_desc):
        for candidata in oficiais:
            if _normalizar_codigo_cnae(candidata.codigo) == CODIGO_CNAE_EVENTOS:
                return candidata.codigo, candidata.descricao
        return CODIGO_CNAE_EVENTOS, desc_cnae.strip()

    ranqueados = sorted(
        oficiais,
        key=lambda candidata: _pontuar_descricao_oficial(desc_cnae, candidata),
        reverse=True,
    )
    melhor = ranqueados[0]
    score_melhor = _pontuar_descricao_oficial(desc_cnae, melhor)
    segundo = _pontuar_descricao_oficial(desc_cnae, ranqueados[1]) if len(ranqueados) > 1 else 0.0

    interseccao = len(tokens_desc & melhor.tokens)
    cobertura = interseccao / max(len(melhor.tokens), 1)

    if score_melhor >= 170.0 or (cobertura >= 0.7 and score_melhor >= segundo + 20.0):
        return melhor.codigo, melhor.descricao

    return None


def _classificar_localmente(
    desc_cnae: str,
    descricao_servico: str,
    id_cnae: str,
    oficiais: tuple[CnaeOficial, ...],
) -> tuple[str, str]:
    """Fallback local quando a IA não estiver disponível ou falhar."""

    candidatos = _selecionar_candidatos(desc_cnae, descricao_servico, id_cnae, oficiais, limite=1)
    if not candidatos:
        return id_cnae.strip(), desc_cnae.strip()
    melhor = candidatos[0]
    return melhor.codigo, melhor.descricao


def _chamar_gemini(
    desc_cnae: str,
    descricao_servico: str,
    id_cnae: str,
    candidatos: list[CnaeOficial],
) -> tuple[str, str]:
    """Consulta o Gemini para escolher um único CNAE oficial a partir de candidatos."""

    if genai is None or types is None:
        raise RuntimeError("SDK do Gemini (google-genai) não instalado.")

    chave = _obter_chave_gemini()
    if not chave:
        raise RuntimeError("GEMINI_API_KEY não configurada.")

    cliente = genai.Client(api_key=chave)

    esquema_resposta = {
        "type": "object",
        "properties": {
            "id_cnae_final": {"type": "string", "description": "Código CNAE final escolhido da lista oficial."},
            "desc_cnae_final": {"type": "string", "description": "Descrição oficial do CNAE escolhido."},
        },
        "required": ["id_cnae_final", "desc_cnae_final"],
        "additionalProperties": False,
    }

    prompt = (
        "Você é um especialista em CNAE. "
        "Escolha exatamente um item da lista oficial que melhor corresponda à NFS-e. "
        "Baseie-se principalmente em DESC_CNAE_PDF e use DESCRICAO_SERVICO apenas como apoio.\n\n"
        f"Dados da nota fiscal:\n"
        f"- desc_cnae_pdf: {desc_cnae}\n"
        f"- descricao_servico: {descricao_servico}\n"
        f"- id_cnae_extraido: {id_cnae}\n\n"
        f"Candidatos oficiais:\n"
        f"{json.dumps([{'codigo': c.codigo, 'descricao': c.descricao} for c in candidatos], ensure_ascii=False, indent=2)}"
    )

    configuracao = {
        "temperature": 0,
        "response_mime_type": "application/json",
        "response_json_schema": esquema_resposta,
    }

    for modelo in _modelos_candidatos_gemini():
        try:
            resposta = cliente.models.generate_content(
                model=modelo,
                contents=prompt,
                config=configuracao,
            )
            texto_resposta = getattr(resposta, "text", "") or ""
            if not texto_resposta:
                continue

            resultado = json.loads(texto_resposta)
            id_cnae_final = str(resultado.get("id_cnae_final", "")).strip()
            desc_cnae_final = str(resultado.get("desc_cnae_final", "")).strip()
            if id_cnae_final and desc_cnae_final:
                return id_cnae_final, desc_cnae_final
        except Exception:
            continue

    raise RuntimeError("Todas as chamadas aos modelos do Gemini falharam ou retornaram dados inválidos.")


def resolver_cnae_final(
    desc_cnae: str,
    descricao_servico: str = "",
    id_cnae: str = "",
    caminho_oficial: str = "cnae_oficial.xlsx",
) -> tuple[str, str]:
    """Resolve o CNAE final com Gemini quando possível e fallback local quando necessário."""

    try:
        oficiais = carregar_cnaes_oficiais(caminho_oficial)
    except FileNotFoundError:
        return id_cnae.strip(), desc_cnae.strip()

    resolvido_por_descricao = _resolver_por_descricao_explicita(desc_cnae, oficiais)
    if resolvido_por_descricao is not None:
        return resolvido_por_descricao

    candidatos = _selecionar_candidatos(desc_cnae, descricao_servico, id_cnae, oficiais)
    if not candidatos:
        return id_cnae.strip(), desc_cnae.strip()

    try:
        return _chamar_gemini(desc_cnae, descricao_servico, id_cnae, candidatos)
    except Exception:
        return _classificar_localmente(desc_cnae, descricao_servico, id_cnae, oficiais)


def enriquecer_registros_cnae_final(registros: list[Any]) -> None:
    """Preenche os campos finais de CNAE nos registros já extraídos dos PDFs."""

    cache: dict[tuple[str, str, str], tuple[str, str]] = {}
    for registro in registros:
        desc_cnae = str(getattr(registro, "desc_cnae", "") or "").strip()
        descricao_servico = str(getattr(registro, "descricao_servico", "") or "").strip()
        id_cnae = str(getattr(registro, "id_cnae", "") or "").strip()
        chave = (_normalizar_texto(desc_cnae), _normalizar_texto(descricao_servico), _normalizar_texto(id_cnae))

        if chave not in cache:
            cache[chave] = resolver_cnae_final(desc_cnae, descricao_servico, id_cnae)

        id_cnae_final, desc_cnae_final = cache[chave]
        setattr(registro, "id_cnae_final", id_cnae_final)
        setattr(registro, "desc_cnae_final", desc_cnae_final)

        if _normalizar_codigo_cnae(id_cnae_final) in IDS_QUE_FORCAM_ISS_RETIDO:
            setattr(registro, "iss_retido", "SIM")
