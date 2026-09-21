"""Escrituração Multi-CNPJ — acesso às empresas no portal ISS Fortaleza.

Fluxo por empresa (grupo de linhas com o mesmo CNPJ_TOMADOR da planilha gerada por
"Processar XMLs — Multi-CNPJ"): login com o CPF/senha da própria planilha, seleção e
CONFERÊNCIA da empresa pelo CNPJ, (etapa seguinte: escrituração das notas) e logout
completo — incluindo a sessão do login único (idp2) — antes da próxima empresa.

Este módulo separa:
- a lógica pura (planilha, credenciais, orquestração, retomada), testável sem navegador;
- `AcessoIssSelenium`, a implementação real que conduz o Chrome.

Segurança: a senha vive só em memória. Nunca vai para log, mensagem de erro, tela ou
arquivo de retomada; nos logs o login aparece mascarado.
"""

from __future__ import annotations

import json
import re
import time
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Optional, Protocol

from openpyxl import load_workbook

URL_HOME = "https://iss.fortaleza.ce.gov.br/grpfor/home.seam"
URL_LOGOUT_LOGIN_UNICO = (
    "https://idp2.sefin.fortaleza.ce.gov.br/realms/sefin/protocol/openid-connect/logout"
    "?post_logout_redirect_uri=https://iss.fortaleza.ce.gov.br/grpfor"
)

COLUNAS_MULTI_CNPJ = ("CNPJ_TOMADOR", "LOGIN", "SENHA")

STATUS_ACESSO_VALIDADO = "ACESSO_VALIDADO"
STATUS_CONCLUIDA = "CONCLUIDA"
STATUS_JA_CONCLUIDA = "JA_CONCLUIDA"
STATUS_PLANILHA_INVALIDA = "PLANILHA_INVALIDA"
STATUS_LOGIN_RECUSADO = "LOGIN_RECUSADO"
STATUS_EMPRESA_NAO_ENCONTRADA = "EMPRESA_NAO_ENCONTRADA"
STATUS_CNPJ_DIVERGENTE = "CNPJ_DIVERGENTE"
STATUS_ERRO = "ERRO"
STATUS_CANCELADA = "CANCELADA"

ROTULOS_STATUS = {
    STATUS_ACESSO_VALIDADO: "Acesso validado",
    STATUS_CONCLUIDA: "Concluída",
    STATUS_JA_CONCLUIDA: "Já concluída (pulada)",
    STATUS_PLANILHA_INVALIDA: "Dados da planilha inválidos",
    STATUS_LOGIN_RECUSADO: "Login recusado",
    STATUS_EMPRESA_NAO_ENCONTRADA: "Empresa não encontrada neste login",
    STATUS_CNPJ_DIVERGENTE: "CNPJ divergente",
    STATUS_ERRO: "Erro",
    STATUS_CANCELADA: "Cancelada",
}


# ---------------------------------------------------------------------------
# Exceções
# ---------------------------------------------------------------------------


class ErroAcessoIss(RuntimeError):
    """Falha ao acessar uma empresa no portal (a execução segue para a próxima empresa)."""

    status = STATUS_ERRO


class LoginRecusadoError(ErroAcessoIss):
    status = STATUS_LOGIN_RECUSADO


class EmpresaNaoEncontradaError(ErroAcessoIss):
    status = STATUS_EMPRESA_NAO_ENCONTRADA


class CnpjDivergenteError(ErroAcessoIss):
    status = STATUS_CNPJ_DIVERGENTE


class OperacaoCanceladaError(RuntimeError):
    """O operador pediu para cancelar a execução."""


# ---------------------------------------------------------------------------
# Utilitários puros
# ---------------------------------------------------------------------------


def somente_alfanumericos(texto: Any) -> str:
    """Remove pontuação e espaços; maiúsculas (cobre também o novo CNPJ alfanumérico)."""
    return re.sub(r"[^0-9A-Za-z]", "", str(texto or "")).upper()


def formatar_cnpj(cnpj: str) -> str:
    c = somente_alfanumericos(cnpj)
    if len(c) != 14:
        return cnpj
    return f"{c[0:2]}.{c[2:5]}.{c[5:8]}/{c[8:12]}-{c[12:14]}"


def normalizar_cnpj_celula(valor: Any) -> str:
    """CNPJ da célula, só com caracteres alfanuméricos. Se o Excel guardou como número e
    perdeu os zeros à esquerda, repõe até 14 dígitos."""
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(int(valor)) if isinstance(valor, int) else str(valor)
    limpo = somente_alfanumericos(texto)
    if limpo.isdigit() and len(limpo) < 14:
        limpo = limpo.zfill(14)
    return limpo


def normalizar_login(valor: Any) -> str:
    """Login (CPF) da célula. O portal aceita o CPF com ou sem pontuação; se o Excel guardou
    como número e perdeu zeros à esquerda, repõe até 11 dígitos. Logins que não parecem
    CPF são mantidos como estão."""
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    texto = str(valor).strip()
    if re.fullmatch(r"[\d.\-\s]+", texto):
        digitos = re.sub(r"\D", "", texto)
        return digitos.zfill(11) if len(digitos) <= 11 else digitos
    return texto


def normalizar_senha(valor: Any) -> str:
    """Senha da célula, sem qualquer alteração (nem strip) — exceto números inteiros que o
    Excel devolve como float/int."""
    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, float) and valor.is_integer():
        return str(int(valor))
    return str(valor)


def mascarar_login(login: str) -> str:
    """Máscara para logs: CPF de 11 dígitos vira `***.***.***-NN`; outros, `***` + 2 últimos."""
    digitos = re.sub(r"\D", "", login or "")
    if len(digitos) == 11:
        return f"***.***.***-{digitos[-2:]}"
    return f"***{(login or '')[-2:]}"


def _sanitizar(mensagem: str, *segredos: str) -> str:
    """Remove valores sensíveis de uma mensagem antes de logá-la."""
    for segredo in segredos:
        if segredo:
            mensagem = mensagem.replace(segredo, "***")
    return mensagem


# ---------------------------------------------------------------------------
# Planilha
# ---------------------------------------------------------------------------


@dataclass
class EmpresaMultiCnpj:
    cnpj: str
    login: str = ""
    senha: str = field(default="", repr=False)
    qtd_notas: int = 0
    problemas: list[str] = field(default_factory=list)
    avisos: list[str] = field(default_factory=list)

    @property
    def valida(self) -> bool:
        return not self.problemas


@dataclass
class PlanilhaMultiCnpj:
    empresas: list[EmpresaMultiCnpj]
    avisos_gerais: list[str] = field(default_factory=list)


def carregar_empresas_xlsx(caminho_xlsx: Path) -> PlanilhaMultiCnpj:
    """Lê a planilha Multi-CNPJ e agrupa as linhas por CNPJ_TOMADOR (ordem da primeira
    aparição), validando LOGIN e SENHA: preenchidos e IGUAIS em todas as linhas do CNPJ
    (valores diferentes costumam vir do autopreenchimento do Excel, que incrementa números
    no fim do texto). As mensagens de problema nunca contêm os valores."""
    workbook = load_workbook(caminho_xlsx, data_only=True)
    try:
        planilha = workbook.active
        linhas = planilha.iter_rows(values_only=True)
        cabecalhos = [str(c or "").strip().upper() for c in next(linhas, [])]
        faltantes = [c for c in COLUNAS_MULTI_CNPJ if c not in cabecalhos]
        if faltantes:
            raise ValueError(
                f"A planilha não tem as colunas {', '.join(faltantes)}. Use a planilha gerada por "
                "'Processar XMLs — Multi-CNPJ'."
            )
        ix = {nome: cabecalhos.index(nome) for nome in COLUNAS_MULTI_CNPJ}

        grupos: dict[str, dict[str, Any]] = {}
        avisos_gerais: list[str] = []
        linhas_sem_cnpj = 0
        for linha in linhas:
            if not any(v is not None and str(v).strip() for v in linha):
                continue
            cnpj = normalizar_cnpj_celula(linha[ix["CNPJ_TOMADOR"]])
            if not cnpj:
                linhas_sem_cnpj += 1
                continue
            grupo = grupos.setdefault(cnpj, {"logins": [], "senhas": [], "qtd": 0})
            grupo["qtd"] += 1
            grupo["logins"].append(normalizar_login(linha[ix["LOGIN"]]))
            grupo["senhas"].append(normalizar_senha(linha[ix["SENHA"]]))
        if linhas_sem_cnpj:
            avisos_gerais.append(
                f"{linhas_sem_cnpj} linha(s) sem CNPJ_TOMADOR foram ignoradas (nenhuma empresa para acessar)."
            )
    finally:
        workbook.close()

    empresas: list[EmpresaMultiCnpj] = []
    for cnpj, grupo in grupos.items():
        empresa = EmpresaMultiCnpj(cnpj=cnpj, qtd_notas=grupo["qtd"])
        if len(cnpj) != 14:
            empresa.problemas.append("CNPJ_TOMADOR não tem 14 caracteres")
        for nome, valores in (("LOGIN", grupo["logins"]), ("SENHA", grupo["senhas"])):
            preenchidos = [v for v in valores if v != ""]
            distintos = set(preenchidos)
            if not preenchidos:
                empresa.problemas.append(f"{nome} em branco")
            elif len(preenchidos) < len(valores):
                empresa.problemas.append(f"{nome} preenchido só em algumas linhas deste CNPJ")
            elif len(distintos) > 1:
                empresa.problemas.append(
                    f"{nome} diferente entre as linhas deste CNPJ (confira o autopreenchimento do Excel)"
                )
        if empresa.valida:
            empresa.login = grupo["logins"][0]
            empresa.senha = grupo["senhas"][0]
            if empresa.senha != empresa.senha.strip():
                empresa.avisos.append("SENHA tem espaço no início ou no fim (confira se é isso mesmo)")
        empresas.append(empresa)
    return PlanilhaMultiCnpj(empresas=empresas, avisos_gerais=avisos_gerais)


# ---------------------------------------------------------------------------
# Retomada (empresas já concluídas)
# ---------------------------------------------------------------------------


def caminho_estado(caminho_xlsx: Path) -> Path:
    return caminho_xlsx.with_name(caminho_xlsx.stem + ".escrituracao_multi_estado.json")


def ler_empresas_concluidas(caminho_xlsx: Path) -> set[str]:
    try:
        dados = json.loads(caminho_estado(caminho_xlsx).read_text(encoding="utf-8"))
        return set(dados.get("empresas_concluidas", []))
    except (OSError, ValueError):
        return set()


def marcar_empresa_concluida(caminho_xlsx: Path, cnpj: str) -> None:
    concluidas = ler_empresas_concluidas(caminho_xlsx)
    concluidas.add(cnpj)
    try:
        caminho_estado(caminho_xlsx).write_text(
            json.dumps({"empresas_concluidas": sorted(concluidas)}, ensure_ascii=False),
            encoding="utf-8",
        )
    except OSError:
        pass


def limpar_estado(caminho_xlsx: Path) -> None:
    try:
        caminho_estado(caminho_xlsx).unlink()
    except OSError:
        pass


# ---------------------------------------------------------------------------
# Orquestração
# ---------------------------------------------------------------------------


class AcessoIss(Protocol):
    def login(self, login: str, senha: str) -> None: ...

    def garantir_empresa(self, cnpj: str) -> None: ...

    def logout(self) -> bool: ...

    def encerrar(self) -> None: ...


@dataclass
class ResultadoEmpresa:
    cnpj: str
    status: str
    mensagem: str = ""
    qtd_notas: int = 0


@dataclass
class ResultadoMultiCnpj:
    empresas: list[ResultadoEmpresa] = field(default_factory=list)
    cancelado: bool = False

    def contagem(self) -> dict[str, int]:
        contagem: dict[str, int] = {}
        for r in self.empresas:
            contagem[r.status] = contagem.get(r.status, 0) + 1
        return contagem

    def para_json(self) -> dict[str, Any]:
        """Resumo para a GUI (sem nenhum dado sensível)."""
        return {
            "cancelado": self.cancelado,
            "total": len(self.empresas),
            "contagem": {ROTULOS_STATUS[k]: v for k, v in self.contagem().items()},
            "empresas": [
                {
                    "cnpj": formatar_cnpj(r.cnpj),
                    "status": ROTULOS_STATUS.get(r.status, r.status),
                    "ok": r.status in (STATUS_ACESSO_VALIDADO, STATUS_CONCLUIDA, STATUS_JA_CONCLUIDA),
                    "mensagem": r.mensagem,
                }
                for r in self.empresas
            ],
        }


def executar_escrituracao_multi_cnpj(
    caminho_xlsx: Path,
    criar_acesso: Callable[[], AcessoIss],
    processar_empresa: Optional[Callable[[EmpresaMultiCnpj, AcessoIss], int]] = None,
    callback_log: Optional[Callable[[str, bool], None]] = None,
    callback_progresso: Optional[Callable[[int, str], None]] = None,
    callback_pausa: Optional[Callable[[], None]] = None,
    callback_cancelamento: Optional[Callable[[], bool]] = None,
) -> ResultadoMultiCnpj:
    """Percorre as empresas da planilha: login → seleção/conferência da empresa → (se
    `processar_empresa` for informado) escrituração das notas → logout.

    Sem `processar_empresa` (Etapa 1) roda em modo "apenas validar acessos": não escritura
    nada e não usa/grava o arquivo de retomada. Com ele, empresas já concluídas numa
    execução interrompida são puladas; empresas que falharam são tentadas de novo.

    Falha em uma empresa (planilha inválida, login recusado, empresa não encontrada, CNPJ
    divergente, erro) não interrompe as demais. Cancelar encerra tudo (com logout).
    """

    def log(msg: str, is_error: bool = False) -> None:
        if callback_log:
            callback_log(msg, is_error)
        else:
            print(f"{'[ERRO] ' if is_error else ''}{msg}", flush=True)

    def progresso(pct: int, status: str) -> None:
        if callback_progresso:
            callback_progresso(pct, status)

    def cancelado() -> bool:
        return bool(callback_cancelamento and callback_cancelamento())

    modo_validacao = processar_empresa is None
    planilha = carregar_empresas_xlsx(caminho_xlsx)
    for aviso in planilha.avisos_gerais:
        log(aviso, True)
    log(
        f"Planilha lida: {len(planilha.empresas)} empresa(s) encontrada(s)."
        + (" Modo: apenas validar acessos (nada será escriturado)." if modo_validacao else "")
    )

    resultado = ResultadoMultiCnpj()
    concluidas = set() if modo_validacao else ler_empresas_concluidas(caminho_xlsx)
    a_processar: list[EmpresaMultiCnpj] = []

    for empresa in planilha.empresas:
        prefixo = f"[{formatar_cnpj(empresa.cnpj)}] "
        if not empresa.valida:
            mensagem = "; ".join(empresa.problemas)
            log(f"{prefixo}Não vou acessar esta empresa: {mensagem}.", True)
            resultado.empresas.append(ResultadoEmpresa(empresa.cnpj, STATUS_PLANILHA_INVALIDA, mensagem, empresa.qtd_notas))
        elif empresa.cnpj in concluidas:
            log(f"{prefixo}Já foi concluída numa execução anterior — vou pular.")
            resultado.empresas.append(ResultadoEmpresa(empresa.cnpj, STATUS_JA_CONCLUIDA, "", empresa.qtd_notas))
        else:
            for aviso in empresa.avisos:
                log(f"{prefixo}Aviso: {aviso}.")
            a_processar.append(empresa)

    if not a_processar:
        log("Nenhuma empresa para acessar.")
        progresso(100, "Concluído!")
        return resultado

    total = len(a_processar)
    acesso: Optional[AcessoIss] = None
    try:
        acesso = criar_acesso()
        for posicao, empresa in enumerate(a_processar, start=1):
            prefixo = f"[{formatar_cnpj(empresa.cnpj)}] "
            if callback_pausa:
                callback_pausa()
            if cancelado():
                resultado.cancelado = True
                break

            progresso(int(((posicao - 1) / total) * 100), f"Empresa {posicao} de {total} — {formatar_cnpj(empresa.cnpj)}")
            log(f"{prefixo}Empresa {posicao} de {total}: entrando no portal com o login {mascarar_login(empresa.login)}...")
            status, mensagem, qtd = STATUS_ERRO, "", 0
            try:
                acesso.login(empresa.login, empresa.senha)
                log(f"{prefixo}Login realizado.")
                acesso.garantir_empresa(empresa.cnpj)
                log(f"{prefixo}Empresa confirmada no portal (CNPJ conferido).")
                if processar_empresa is None:
                    status, mensagem = STATUS_ACESSO_VALIDADO, "Login, seleção da empresa e conferência do CNPJ ok."
                else:
                    qtd = processar_empresa(empresa, acesso)
                    marcar_empresa_concluida(caminho_xlsx, empresa.cnpj)
                    status, mensagem = STATUS_CONCLUIDA, f"{qtd} nota(s) escriturada(s)."
                log(f"{prefixo}{ROTULOS_STATUS[status]}.")
            except OperacaoCanceladaError:
                status, mensagem = STATUS_CANCELADA, "Cancelada pelo operador."
                resultado.cancelado = True
                log(f"{prefixo}Cancelado pelo operador.", True)
            except ErroAcessoIss as erro:
                status = erro.status
                mensagem = _sanitizar(str(erro), empresa.senha, empresa.login)
                log(f"{prefixo}{ROTULOS_STATUS[status]}: {mensagem} — vou pular esta empresa.", True)
            except Exception as erro:  # noqa: BLE001 — qualquer falha inesperada não pode derrubar as outras empresas
                status = STATUS_ERRO
                mensagem = _sanitizar(f"{type(erro).__name__}: {erro}", empresa.senha, empresa.login)
                log(f"{prefixo}Erro inesperado: {mensagem} — vou pular esta empresa.", True)
            resultado.empresas.append(ResultadoEmpresa(empresa.cnpj, status, mensagem, qtd))

            # Logout SEMPRE (mesmo com erro/cancelamento): a próxima empresa não pode herdar a sessão.
            log(f"{prefixo}Saindo do portal (logout)...")
            try:
                saiu = acesso.logout()
            except Exception:  # noqa: BLE001
                saiu = False
            if saiu:
                log(f"{prefixo}Logout confirmado.")
            else:
                restam = posicao < total and not resultado.cancelado
                detalhe = str(getattr(acesso, "detalhe_logout", "") or "")
                destino = (
                    "vou reabrir o navegador para garantir uma sessão limpa"
                    if restam
                    else "vou apenas fechar o navegador"
                )
                log(
                    f"{prefixo}Não consegui confirmar o logout — {destino}."
                    + (f" Detalhe: {_sanitizar(detalhe, empresa.senha, empresa.login)}" if detalhe else ""),
                    True,
                )
                try:
                    acesso.encerrar()
                except Exception:  # noqa: BLE001
                    pass
                acesso = None
                if restam:
                    acesso = criar_acesso()

            if resultado.cancelado:
                break
    finally:
        if acesso is not None:
            try:
                acesso.encerrar()
            except Exception:  # noqa: BLE001
                pass

    if resultado.cancelado:
        feitas = {r.cnpj for r in resultado.empresas}
        for empresa in a_processar:
            if empresa.cnpj not in feitas:
                resultado.empresas.append(ResultadoEmpresa(empresa.cnpj, STATUS_CANCELADA, "Não chegou a ser acessada.", empresa.qtd_notas))
    elif not modo_validacao:
        # Tudo concluído: não há mais nada a retomar
        if all(r.status in (STATUS_CONCLUIDA, STATUS_JA_CONCLUIDA) for r in resultado.empresas):
            limpar_estado(caminho_xlsx)

    contagem = resultado.contagem()
    resumo = ", ".join(f"{ROTULOS_STATUS[k]}: {v}" for k, v in contagem.items())
    log(f"Resumo — {resumo}." if not resultado.cancelado else f"Execução cancelada. Resumo até aqui — {resumo}.", resultado.cancelado)
    progresso(100, "Cancelado" if resultado.cancelado else "Concluído!")
    return resultado


# ---------------------------------------------------------------------------
# Implementação real: Selenium / Chrome
# ---------------------------------------------------------------------------


class AcessoIssSelenium:
    """Conduz o Chrome no portal ISS Fortaleza. Um navegador novo por instância (sessão limpa)."""

    def __init__(
        self,
        abrir_navegador: Optional[Callable[[], Any]] = None,
        callback_log: Optional[Callable[[str, bool], None]] = None,
        callback_cancelamento: Optional[Callable[[], bool]] = None,
        timeout: int = 30,
    ) -> None:
        self._abrir_navegador = abrir_navegador
        self._log = callback_log or (lambda msg, is_error=False: None)
        self._cancelamento = callback_cancelamento
        self.timeout = timeout
        self.driver: Any = None
        self.tempo_maximo_logout = 45.0  # segundos para concluir o logout completo
        self.pausa_logout = 0.6  # intervalo entre as verificações do logout
        self.detalhe_logout = ""  # diagnóstico da última falha de logout (etapas e página), sem dados sensíveis

    # -- infraestrutura -----------------------------------------------------

    def _navegador(self) -> Any:
        if self.driver is None:
            if self._abrir_navegador is not None:
                self.driver = self._abrir_navegador()
            else:
                from iss_fortaleza_automacao import abrir_navegador_visivel  # import tardio: traz o Selenium

                self.driver = abrir_navegador_visivel()
        return self.driver

    def _checar_cancelamento(self) -> None:
        if self._cancelamento and self._cancelamento():
            raise OperacaoCanceladaError()

    def _esperar(
        self,
        condicao: Callable[[Any], Any],
        timeout: Optional[float] = None,
        intervalo: float = 0.4,
        cancelavel: bool = True,
    ) -> Any:
        """Espera até `condicao(driver)` devolver algo verdadeiro (ou False no timeout). Atende ao
        cancelamento, exceto com `cancelavel=False` (usado no logout, que deve acontecer sempre)."""
        limite = time.time() + (timeout if timeout is not None else self.timeout)
        while time.time() < limite:
            if cancelavel:
                self._checar_cancelamento()
            try:
                valor = condicao(self._navegador())
            except Exception:  # noqa: BLE001 — DOM em transição (elemento obsoleto etc.)
                valor = False
            if valor:
                return valor
            time.sleep(intervalo)
        return False

    @staticmethod
    def _visivel(by: str, seletor: str) -> Callable[[Any], Any]:
        """Condição de espera: primeiro elemento visível que casa com o seletor (ou False)."""

        def condicao(driver: Any) -> Any:
            elementos = [e for e in driver.find_elements(by, seletor) if e.is_displayed()]
            return elementos[0] if elementos else False

        return condicao

    def _logado(self) -> bool:
        from selenium.webdriver.common.by import By

        try:
            return bool(self._navegador().find_elements(By.CSS_SELECTOR, 'a[href*="identity.logout"]'))
        except Exception:  # noqa: BLE001
            return False

    def _modal_selecao_visivel(self) -> bool:
        from selenium.webdriver.common.by import By

        try:
            return any(e.is_displayed() for e in self._navegador().find_elements(By.ID, "alteraInscricaoForm:cpfPesquisa"))
        except Exception:  # noqa: BLE001
            return False

    def _na_tela_de_login_do_idp(self) -> bool:
        from selenium.webdriver.common.by import By

        try:
            return "idp2" in self._url_sem_consulta() and self._visivel(By.ID, "username")(self._navegador()) is not False
        except Exception:  # noqa: BLE001 — DOM em transição
            return False

    def _inscricao_atual_definida(self) -> bool:
        """True quando o cabeçalho mostra 'Inscrição Atual: ...' (e não 'Selecione uma inscrição')."""
        from selenium.webdriver.common.by import By

        try:
            link = self._visivel(By.ID, "j_id157:j_id158")(self._navegador())
            return bool(link) and (link.text or "").strip().startswith("Inscrição Atual")
        except Exception:  # noqa: BLE001
            return False

    def _url_sem_consulta(self) -> str:
        try:
            return self._navegador().current_url.split("?")[0]
        except Exception:  # noqa: BLE001
            return "(desconhecida)"

    # -- login ---------------------------------------------------------------

    def _abrir_tela_de_login(self) -> None:
        """Abre o portal e chega à tela de CPF/senha do login único."""
        from selenium.webdriver.common.by import By

        driver = self._navegador()
        driver.get(URL_HOME)
        if self._logado():
            # Sessão anterior ainda aberta: sai antes de entrar com outro login.
            self._log("Sessão anterior ainda aberta — vou sair antes de continuar.", False)
            self.logout()
            driver.get(URL_HOME)
        link = self._esperar(self._visivel(By.CSS_SELECTOR, 'a[href="/grpfor/oauth2/login"]'))
        if not link:
            raise ErroAcessoIss(f"Não encontrei o botão 'Fazer login' do portal (página: {self._url_sem_consulta()}).")
        link.click()
        if not self._esperar(self._visivel(By.ID, "username")):
            raise ErroAcessoIss(f"A tela de login do portal não abriu (página: {self._url_sem_consulta()}).")

    def _mensagem_erro_login(self) -> str:
        from selenium.webdriver.common.by import By

        textos: list[str] = []
        for seletor in (".alert", ".alert-error", ".kc-feedback-text", "#input-error", ".error-message", "span.error"):
            try:
                for e in self._navegador().find_elements(By.CSS_SELECTOR, seletor):
                    t = (e.text or "").strip()
                    if t and t not in textos:
                        textos.append(t)
            except Exception:  # noqa: BLE001
                continue
        return " | ".join(textos)[:200]

    def login(self, login: str, senha: str) -> None:
        from selenium.webdriver.common.by import By

        self._abrir_tela_de_login()
        driver = self._navegador()
        campo_login = driver.find_element(By.ID, "username")
        campo_senha = driver.find_element(By.ID, "password")
        campo_login.click()
        campo_login.clear()
        campo_login.send_keys(login)
        campo_senha.click()
        campo_senha.clear()
        campo_senha.send_keys(senha)
        driver.find_element(By.ID, "botao-entrar").click()

        # Aguarda a página de resposta ao envio: ou entra no portal, ou volta à tela de login com erro.
        limite = time.time() + self.timeout
        while time.time() < limite:
            self._checar_cancelamento()
            if self._logado():
                return
            if self._na_tela_de_login_do_idp() and self._pagina_mudou_apos_envio(campo_senha):
                mensagem = self._mensagem_erro_login()
                # Sem nova tentativa: repetir uma senha errada pode bloquear a conta.
                raise LoginRecusadoError(
                    "o portal recusou o login" + (f" ({mensagem})" if mensagem else " (confira CPF e senha)")
                )
            time.sleep(0.5)
        raise ErroAcessoIss(f"Tempo esgotado aguardando o login (página: {self._url_sem_consulta()}).")

    @staticmethod
    def _pagina_mudou_apos_envio(campo_senha_antigo: Any) -> bool:
        """True quando o formulário enviado foi substituído (a resposta do servidor já chegou)."""
        try:
            campo_senha_antigo.is_enabled()
            return False  # ainda é o mesmo DOM: o envio não terminou
        except Exception:  # noqa: BLE001 — elemento obsoleto: a página foi recarregada
            return True

    # -- empresa -------------------------------------------------------------

    def _clicar_selecionar_da_linha(self, cnpj: str) -> bool:
        """Clica em 'Selecionar' na linha da tabela cujo CPF/CNPJ é `cnpj`. False se não estiver listada."""
        from selenium.webdriver.common.by import By

        linhas = self._navegador().find_elements(By.XPATH, '//tbody[@id="alteraInscricaoForm:empresaDataTable:tb"]/tr')
        for linha in linhas:
            documentos = linha.find_elements(By.XPATH, './/a[contains(@id,"linkDocumento")]')
            if documentos and somente_alfanumericos(documentos[0].text) == cnpj:
                linha.find_element(By.XPATH, './/a[@title="Selecionar"]').click()
                return True
        return False

    def _pesquisar_inscricao_por_cnpj(self, cnpj: str) -> None:
        """Radio CNPJ → digitar → sugestão (ou Enter) → clicar na sugestão (filtra a tabela)."""
        from selenium.webdriver.common.by import By
        from selenium.webdriver.common.keys import Keys

        driver = self._navegador()
        driver.find_element(By.ID, "alteraInscricaoForm:tipoPesquisa:1").click()  # opção CNPJ (padrão é CPF)
        time.sleep(1.5)  # a troca dispara AJAX que reaplica a máscara; digitar antes apaga o texto
        campo = driver.find_element(By.ID, "alteraInscricaoForm:cpfPesquisa")
        campo.click()
        campo.clear()
        campo.send_keys(cnpj)
        seletor_sugestao = '//div[@id="alteraInscricaoForm:sugestaoPesquisa"]//tr[contains(@class,"richfaces_suggestionEntry")]'

        def sugestao(drv: Any) -> Any:
            itens = [e for e in drv.find_elements(By.XPATH, seletor_sugestao) if e.is_displayed()]
            return itens[0] if itens else False

        item = self._esperar(sugestao, timeout=6)
        if not item:
            campo.send_keys(Keys.ENTER)  # na prática a lista só apareceu depois do Enter
            item = self._esperar(sugestao, timeout=8)
        if not item:
            raise EmpresaNaoEncontradaError(
                f"o CNPJ {formatar_cnpj(cnpj)} não aparece na lista de empresas deste login"
            )
        item.click()

    def _selecionar_inscricao(self, cnpj: str) -> None:
        from selenium.webdriver.common.by import By

        self._log("Selecionando a empresa pelo CNPJ...", False)
        if not self._modal_selecao_visivel():
            raise ErroAcessoIss("a tela de seleção de inscrição não está aberta")
        # Atalho: empresa já listada na página atual da tabela
        if not self._clicar_selecionar_da_linha(cnpj):
            self._pesquisar_inscricao_por_cnpj(cnpj)
            if not self._esperar(lambda d: self._clicar_selecionar_da_linha(cnpj), timeout=10):
                raise EmpresaNaoEncontradaError(
                    f"depois da pesquisa, o CNPJ {formatar_cnpj(cnpj)} não apareceu na tabela de empresas"
                )
        # Troca de uma inscrição já selecionada: o portal pede confirmação ("A operação atual será abortada")
        confirmar = self._esperar(
            self._visivel(
                By.XPATH,
                '//*[@id="alteraInscricaoForm:confirmaAlteraInscricaoAtualModalContainer"]'
                '//*[(self::input and @value="Sim") or ((self::a or self::button) and normalize-space(.)="Sim")]',
            ),
            timeout=3,
        )
        if confirmar:
            confirmar.click()
        if not self._esperar(lambda d: not self._modal_selecao_visivel(), timeout=self.timeout):
            raise ErroAcessoIss("a tela de seleção de inscrição não fechou depois de selecionar a empresa")

    def _fechar_modal_detalhes(self) -> None:
        driver = self._navegador()
        try:
            driver.execute_script("Richfaces.hideModalPanel('detalhesPrestadorModal');")
        except Exception:  # noqa: BLE001
            from selenium.webdriver.common.keys import Keys

            driver.switch_to.active_element.send_keys(Keys.ESCAPE)

    def _ler_cnpj_inscricao_atual(self) -> str:
        """Abre 'Dados do contribuinte' (link da Inscrição Atual) e devolve o CPF/CNPJ mostrado."""
        from selenium.webdriver.common.by import By

        if not self._esperar(lambda d: self._inscricao_atual_definida()):
            raise ErroAcessoIss("não encontrei a 'Inscrição Atual' no cabeçalho do portal")
        self._navegador().find_element(By.ID, "j_id157:j_id158").click()

        def texto_do_modal(drv: Any) -> Any:
            e = drv.find_element(By.ID, "detalhesPrestadorModalContainer")
            if not e.is_displayed():
                return False
            achado = re.search(r"CPF/CNPJ:\s*([0-9A-Za-z./\-]+)", e.text or "")
            return achado.group(1) if achado else False

        bruto = self._esperar(texto_do_modal)
        if not bruto:
            raise ErroAcessoIss("não consegui ler o CNPJ da empresa em 'Dados do contribuinte'")
        self._fechar_modal_detalhes()
        self._esperar(lambda d: not d.find_element(By.ID, "detalhesPrestadorModalContainer").is_displayed(), timeout=5)
        return somente_alfanumericos(bruto)

    def garantir_empresa(self, cnpj: str) -> None:
        from selenium.webdriver.common.by import By

        # Aguarda o portal pós-login: ou abre sozinho a seleção de inscrição, ou já mostra a inscrição atual
        pronto = self._esperar(
            lambda d: self._modal_selecao_visivel() or self._inscricao_atual_definida(),
            timeout=self.timeout,
        )
        if not pronto:
            raise ErroAcessoIss(f"o portal não mostrou a inscrição nem a seleção de empresas (página: {self._url_sem_consulta()})")

        if self._modal_selecao_visivel():
            self._selecionar_inscricao(cnpj)

        atual = self._ler_cnpj_inscricao_atual()
        if atual != cnpj:
            self._log(
                f"A inscrição atual é do CNPJ {formatar_cnpj(atual)}, e não {formatar_cnpj(cnpj)} — vou trocar de inscrição.",
                False,
            )
            botao_trocar = self._esperar(self._visivel(By.ID, "j_id157:j_id159"))
            if not botao_trocar:
                raise CnpjDivergenteError(
                    f"a inscrição atual é do CNPJ {formatar_cnpj(atual)}, mas a planilha pede {formatar_cnpj(cnpj)}, "
                    "e não achei o botão de trocar de inscrição"
                )
            botao_trocar.click()
            if not self._esperar(lambda d: self._modal_selecao_visivel()):
                raise ErroAcessoIss("a tela de troca de inscrição não abriu")
            self._selecionar_inscricao(cnpj)
            atual = self._ler_cnpj_inscricao_atual()
            if atual != cnpj:
                raise CnpjDivergenteError(
                    f"a inscrição atual é do CNPJ {formatar_cnpj(atual)}, mas a planilha pede {formatar_cnpj(cnpj)}"
                )

    # -- logout / encerramento -----------------------------------------------

    def logout(self) -> bool:
        """Sai do ISS E do login único (idp2). Devolve False se não conseguiu confirmar.

        O logout é uma pequena máquina de estados: a cada volta olha em que tela está e faz o passo
        seguinte, sem depender da ordem — o login único pode encerrar a sessão e o ISS continuar
        logado (ou o contrário), e um modal aberto pode bloquear o "Sair"."""
        self.detalhe_logout = ""
        if self.driver is None:
            return True
        etapas: list[str] = []
        try:
            if self._sair_do_portal(etapas):
                return True
        except Exception as erro:  # noqa: BLE001
            etapas.append(f"erro {type(erro).__name__}")
        self.detalhe_logout = "; ".join(etapas) + f"; página: {self._url_sem_consulta()}"
        return False

    def _aceitar_alerta(self) -> bool:
        """Aceita um alerta do navegador (ex.: 'sair desta página?'), se houver."""
        try:
            self._navegador().switch_to.alert.accept()
            return True
        except Exception:  # noqa: BLE001 — não há alerta
            return False

    def _fechar_modais_abertos(self) -> int:
        """Fecha qualquer modal RichFaces aberto (ex.: seleção de empresa, dados do contribuinte):
        um modal aberto bloqueia o clique em 'Sair'. Devolve quantos foram fechados."""
        try:
            return int(
                self._navegador().execute_script(
                    """
                    var n = 0, vistos = {};
                    var candidatos = document.querySelectorAll('.rich-modalpanel, [id$="ModalContainer"]');
                    candidatos.forEach(function (el) {
                        var r = el.getBoundingClientRect();
                        if (!(r.width > 0 && r.height > 0) || !el.id) return;
                        var id = el.id.replace(/Container$/, '');
                        if (vistos[id]) return;
                        vistos[id] = true;
                        try { Richfaces.hideModalPanel(id); n++; } catch (e) {}
                    });
                    return n;
                    """
                )
                or 0
            )
        except Exception:  # noqa: BLE001
            return 0

    def _clicar_sair_do_iss(self, etapas: list[str]) -> None:
        """Clica em 'Sair' do ISS. Se algo cobrir o botão, registra o motivo e força o clique por script."""
        from selenium.webdriver.common.by import By

        driver = self._navegador()
        try:
            driver.execute_script("window.onbeforeunload = null;")  # evita o aviso "sair desta página?"
        except Exception:  # noqa: BLE001
            self._aceitar_alerta()
        fechados = self._fechar_modais_abertos()
        if fechados:
            etapas.append(f"{fechados} modal(is) fechado(s)")
        botao = driver.find_element(By.CSS_SELECTOR, 'a[href*="identity.logout"]')
        try:
            botao.click()
            etapas.append("clicou em Sair do ISS")
        except Exception as erro:  # noqa: BLE001 — algo cobre o botão
            motivo = " ".join(str(erro).split())[:160]
            etapas.append(f"clique em Sair bloqueado ({type(erro).__name__}: {motivo})")
            self._aceitar_alerta()
            driver.execute_script("arguments[0].click();", botao)
            etapas.append("clique por script")

    def _sair_do_portal(self, etapas: list[str]) -> bool:
        from selenium.webdriver.common.by import By

        driver = self._navegador()
        limite = time.time() + self.tempo_maximo_logout
        ultima: list[str] = []

        def passo(descricao: str) -> None:
            if not ultima or ultima[-1] != descricao:
                etapas.append(descricao)
            ultima.append(descricao)

        acesso_direto = False
        while time.time() < limite:
            try:
                confirmacao = self._visivel(By.ID, "kc-logout")(driver)
                if confirmacao:
                    # O login único pede confirmação ("Você realmente deseja sair?"): sem ela a sessão continua ativa
                    confirmacao.click()
                    passo("confirmou no login único")
                elif driver.find_elements(By.CSS_SELECTOR, 'a[href="/grpfor/oauth2/login"]'):
                    return True  # tela inicial do portal, deslogada
                elif self._logado():
                    passo("ISS ainda logado")
                    self._clicar_sair_do_iss(etapas)
                elif not acesso_direto and "idp2" not in self._url_sem_consulta():
                    passo("logout direto no login único")
                    acesso_direto = True
                    driver.get(URL_LOGOUT_LOGIN_UNICO)
                else:
                    passo("aguardando a página")
            except Exception:  # noqa: BLE001 — DOM em transição ou alerta do navegador travando a leitura
                passo("página em transição")
                self._aceitar_alerta()
            time.sleep(self.pausa_logout)
        passo("tempo esgotado")
        return False

    def encerrar(self) -> None:
        driver, self.driver = self.driver, None
        if driver is not None:
            try:
                driver.quit()
            except Exception:  # noqa: BLE001
                pass
