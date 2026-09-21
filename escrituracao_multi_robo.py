"""Robô de escrituração do portal ISS Fortaleza para o modo Multi-CNPJ.

CÓPIA INDEPENDENTE de `iss_fortaleza_automacao.py` (a função "Escrituração" individual, que
já funciona e NÃO deve ser alterada). Foi feita por poda mecânica: ficaram os helpers de
campos, prestador, CNAE, documento, gravação e navegação; saíram tudo o que era só da
função individual (pedidos por terminal, login manual, abertura de navegador, extração
avulsa, ponto de entrada da CLI).

Diferenças deliberadas em relação ao original:
- não abre nem fecha o navegador: recebe um `driver` já logado na empresa certa (quem cuida
  do login, da empresa e do logout é `escrituracao_multi_cnpj.py`);
- nada é impresso no terminal nem gravado por `tratamento_erros`: tudo sai pelos ganchos
  `configurar_saida` (mensagens) e `configurar_saida(callback_evento=...)` (marcos), que o
  orquestrador liga ao log único da execução (com o CNPJ da empresa em cada linha);
- pausa e cancelamento próprios do módulo (`configurar_controle`), zerados ao terminar;
- GRAVAÇÃO BLOQUEADA POR PADRÃO: o único clique em "Gravar" (`_clicar_gravar_documento`)
  só funciona dentro de `gravacao_habilitada(True)`. No modo simulação e em qualquer teste
  a gravação fica bloqueada e a tentativa levanta `GravacaoBloqueadaError`.

ATENÇÃO — duplicação consciente: uma correção futura nos campos do formulário precisa ser
feita nos DOIS módulos (este e `iss_fortaleza_automacao.py`).
"""

from __future__ import annotations

import re
import sys
import threading
import time
import unicodedata
from contextlib import contextmanager
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any, Callable, Iterator

from openpyxl import load_workbook
from selenium.common.exceptions import (
    ElementNotInteractableException,
    JavascriptException,
    NoSuchElementException,
    StaleElementReferenceException,
    TimeoutException,
    WebDriverException,
)
from selenium.webdriver.common.by import By
from selenium.webdriver.common.keys import Keys
from selenium.webdriver.remote.webdriver import WebDriver
from selenium.webdriver.remote.webelement import WebElement
from selenium.webdriver.support import expected_conditions as EC
from selenium.webdriver.support.select import Select
from selenium.webdriver.support.wait import WebDriverWait

# ---------------------------------------------------------------------------
# Constantes
# ---------------------------------------------------------------------------

URL_ISS_FORTALEZA = "https://iss.fortaleza.ce.gov.br/grpfor/home.seam"

# Mapeamentos de mês por nome (normalizado sem acento)
MESES_POR_NOME = {
    "janeiro": 1,
    "fevereiro": 2,
    "marco": 3,
    "abril": 4,
    "maio": 5,
    "junho": 6,
    "julho": 7,
    "agosto": 8,
    "setembro": 9,
    "outubro": 10,
    "novembro": 11,
    "dezembro": 12,
}

# Abreviações dos meses usadas no calendário do portal
MESES_ABREVIADOS_PORTAL = {
    1: "jan",
    2: "fev",
    3: "mar",
    4: "abr",
    5: "mai",
    6: "jun",
    7: "jul",
    8: "ago",
    9: "set",
    10: "out",
    11: "nov",
    12: "dez",
}


# ---------------------------------------------------------------------------
# Estruturas de dados
# ---------------------------------------------------------------------------


@dataclass(frozen=True)
class CompetenciaTrabalho:
    """Representa a competência trabalhada no portal."""

    mes: int
    ano: int
    rotulo: str

    @property
    def nome_mes(self) -> str:
        nomes = [
            "",
            "Janeiro",
            "Fevereiro",
            "Março",
            "Abril",
            "Maio",
            "Junho",
            "Julho",
            "Agosto",
            "Setembro",
            "Outubro",
            "Novembro",
            "Dezembro",
        ]
        return nomes[self.mes]


@dataclass(frozen=True)
class DocumentoPortalISS:
    """Linha da planilha que será transportada para o formulário do portal."""

    arquivo_pdf: str
    cnpj_prestador: str
    numero_nf: str
    id_cnae_final: str
    data_emissao: str
    descricao_servico: str
    uf_local_prestacao: str
    cidade_local_prestacao: str
    natureza_operacao: str
    iss_retido: str
    valor_servico: str
    # Campos de identificação e endereço do prestador (opcionais — usados no preenchimento manual)
    prefeitura: str = ""
    nome_prestador: str = ""
    uf_prestador: str = ""
    cidade_prestador: str = ""
    cep_prestador: str = ""
    logradouro_prestador: str = ""
    numero_prestador: str = ""
    bairro_prestador: str = ""
    email_prestador: str = ""
    valor_deducoes: str = ""
    descontos_incondicionados: str = ""
    descontos_condicionados: str = ""
    outras_retencoes: str = ""
    ir: str = ""
    pis_nao_retido: str = ""
    cofins_nao_retido: str = ""
    csrf: str = ""
    inss: str = ""
    aliquota: str = ""
    regime_tributario: str = "OUTROS"
    tipo_cliente_prestador: str = ""
    id_cnae: str = ""
    tipo_tributacao: str = "Normal"
    # Empresa (tomador) dona da nota no modo Multi-CNPJ — vem da coluna CNPJ_TOMADOR da planilha
    cnpj_tomador: str = ""


# ---------------------------------------------------------------------------
# Exceções customizadas
# ---------------------------------------------------------------------------


class PrestadorNaoEncontradoError(RuntimeError):
    """Sinaliza que o portal não encontrou uma razão social para o CNPJ informado."""

    def __init__(self, cnpj: str, detalhe_tela: str = "") -> None:
        self.cnpj = cnpj
        self.detalhe_tela = detalhe_tela.strip()
        mensagem = f"Nenhuma Razão Social Encontrada para o CNPJ {cnpj}."
        if self.detalhe_tela:
            mensagem = f"{mensagem} Detalhe da tela: {self.detalhe_tela}"
        super().__init__(mensagem)


class CompetenciaSemEscriturarDisponivelError(RuntimeError):
    """Sinaliza que a competência consultada não liberou o botão de escriturar."""


class AutomacaoCanceladaError(RuntimeError):
    """Sinaliza que o operador cancelou a automação em execução pela GUI."""


class GravacaoBloqueadaError(RuntimeError):
    """Tentativa de gravar uma nota com a gravação bloqueada (modo simulação, testes ou uso
    fora de `gravacao_habilitada(True)`). Nada foi gravado no portal."""


class SessaoExpiradaError(RuntimeError):
    """A sessão do portal expirou (ex.: pausa longa) durante a escrituração da empresa."""


# ---------------------------------------------------------------------------
# Saída, marcos e gravação (ganchos próprios da cópia Multi-CNPJ)
# ---------------------------------------------------------------------------

_CALLBACK_SAIDA: Callable[[str], None] | None = None
_CALLBACK_EVENTO: Callable[[str], None] | None = None


def configurar_saida(
    callback_saida: Callable[[str], None] | None = None,
    callback_evento: Callable[[str], None] | None = None,
) -> None:
    """Liga as mensagens do robô (`callback_saida`, tela + log) e os marcos internos
    (`callback_evento`, só arquivo) ao log único da execução. Sem callback, a mensagem é descartada
    (as mensagens do robô nunca vão para o terminal nem para o `tratamento_erros`)."""

    global _CALLBACK_SAIDA, _CALLBACK_EVENTO
    _CALLBACK_SAIDA = callback_saida
    _CALLBACK_EVENTO = callback_evento


def _emitir(*partes: Any) -> None:
    """Substitui o `print` do original: envia a mensagem ao log da execução."""

    callback = _CALLBACK_SAIDA
    if callback is not None:
        try:
            callback(" ".join(str(p) for p in partes))
        except Exception:
            pass


def registrar_evento_execucao(evento: str, contexto: str, caminhos_log: Any = None) -> None:
    """Mesma assinatura do `tratamento_erros.registrar_evento_execucao` do original, mas o marco
    vai para o gancho de eventos da execução (nunca para pastas de log de outra função)."""

    callback = _CALLBACK_EVENTO
    if callback is not None:
        try:
            callback(f"{contexto}: {evento}")
        except Exception:
            pass


_GRAVACAO_PERMITIDA = False  # bloqueada por padrão: ver `gravacao_habilitada`


@contextmanager
def gravacao_habilitada(permitida: bool) -> Iterator[None]:
    """Libera (ou mantém bloqueado) o clique em "Gravar" dentro do bloco. Só o modo
    "Escriturar de verdade" usa `True`; ao sair do bloco a gravação volta a ficar bloqueada."""

    global _GRAVACAO_PERMITIDA
    anterior = _GRAVACAO_PERMITIDA
    _GRAVACAO_PERMITIDA = bool(permitida)
    try:
        yield
    finally:
        _GRAVACAO_PERMITIDA = anterior


# ---------------------------------------------------------------------------
# Verificação de pausa global do robô
# ---------------------------------------------------------------------------

# Protegido por Lock para evitar race conditions entre threads de pausa e execução (FALHA-04)
_LOCK_CALLBACK_PAUSA = threading.Lock()
_CALLBACK_PAUSA = None
_TEMPO_PAUSADO_S = 0.0  # tempo total bloqueado em pausa desde a última consulta


def _verificar_pausa() -> None:
    """Bloqueia a execução do robô se o operador tiver solicitado a pausa na GUI."""
    global _TEMPO_PAUSADO_S
    with _LOCK_CALLBACK_PAUSA:
        cb = _CALLBACK_PAUSA
    if cb:
        inicio = time.monotonic()
        try:
            cb()
        except Exception:
            pass
        bloqueado = time.monotonic() - inicio
        if bloqueado >= 1.0:  # ignora a latência normal do callback: só conta pausa de verdade
            with _LOCK_CALLBACK_PAUSA:
                _TEMPO_PAUSADO_S += bloqueado


def consumir_tempo_pausado() -> float:
    """Devolve (e zera) quantos segundos o robô ficou bloqueado em pausa desde a última chamada."""
    global _TEMPO_PAUSADO_S
    with _LOCK_CALLBACK_PAUSA:
        total, _TEMPO_PAUSADO_S = _TEMPO_PAUSADO_S, 0.0
    return total


# ---------------------------------------------------------------------------
# Verificação de cancelamento global do robô
# ---------------------------------------------------------------------------

# Mesmo padrão de Lock+global usado para a pausa (FALHA-04), mas SEM engolir
# exceção: é este utilitário quem levanta AutomacaoCanceladaError, e ela precisa
# se propagar livremente até o finally de executar_fluxo_iss() e o except de app.py.
_LOCK_CALLBACK_CANCELAMENTO = threading.Lock()
_CALLBACK_CANCELAMENTO = None

def configurar_controle(
    callback_pausa: Callable[[], None] | None = None,
    callback_cancelamento: Callable[[], bool] | None = None,
) -> None:
    """Registra a pausa e o cancelamento consultados a cada clique/digitação do robô.
    Chame `configurar_controle()` sem argumentos ao terminar para zerá-los."""

    global _CALLBACK_PAUSA, _CALLBACK_CANCELAMENTO, _TEMPO_PAUSADO_S
    with _LOCK_CALLBACK_PAUSA:
        _CALLBACK_PAUSA = callback_pausa
        _TEMPO_PAUSADO_S = 0.0
    with _LOCK_CALLBACK_CANCELAMENTO:
        _CALLBACK_CANCELAMENTO = callback_cancelamento


def _verificar_cancelamento() -> None:
    """Levanta AutomacaoCanceladaError se o operador tiver solicitado o cancelamento
    da automação pela GUI. O callback registrado deve apenas retornar um bool
    (nunca lançar por si só) — é aqui, fora de qualquer try/except supressor, que
    a exceção de cancelamento é de fato levantada."""
    with _LOCK_CALLBACK_CANCELAMENTO:
        cb = _CALLBACK_CANCELAMENTO
    if cb and cb():
        raise AutomacaoCanceladaError("Automação cancelada pelo operador.")


# ---------------------------------------------------------------------------
# Utilitários de texto
# ---------------------------------------------------------------------------


def normalizar_texto(texto: str) -> str:
    """Remove acentos e padroniza o texto para facilitar o parse."""

    normalizado = unicodedata.normalize("NFD", texto)
    return "".join(
        caractere for caractere in normalizado if unicodedata.category(caractere) != "Mn"
    ).lower()


def _interpretar_iss_retido(valor: str) -> tuple[bool, bool]:
    """Interpreta o valor da coluna ISS_RETIDO.

    Retorna (deve_marcar, reconhecido). Quando `reconhecido` é False, o valor não
    corresponde a um formato claro de "Sim"/"Não" — é tratado como "Não retido"
    por segurança (mesmo comportamento de sempre), mas o chamador deve registrar
    um aviso em vez de deixar isso passar silenciosamente.
    """

    normalizado = normalizar_texto(valor or "").strip()
    if normalizado.startswith("sim"):
        return True, True
    if normalizado.startswith("nao") or normalizado == "n":
        return False, True
    return False, False


def _aplicar_override_iss_retido_por_id_cnae(deve_marcar: bool, id_cnae: str) -> bool:
    """Aplica a regra de negócio pela qual o ID_CNAE (código bruto de serviço/
    tributação extraído do XML, antes da classificação final de CNAE) tem
    prioridade sobre o valor de ISS_RETIDO já lido da planilha: 1207 sempre
    marca o checkbox, 1213 sempre desmarca. Qualquer outro ID_CNAE não altera
    `deve_marcar`.
    """

    id_cnae_norm = re.sub(r"\D+", "", id_cnae or "")
    if id_cnae_norm == "1207":
        return True
    if id_cnae_norm == "1213":
        return False
    return deve_marcar


def _aplicar_override_iss_retido_por_tipo_tributacao(deve_marcar: bool, tipo_tributacao: str) -> bool:
    """Simples Nacional MEI sempre desmarca o checkbox ISS Retido — tem
    prioridade sobre o override de ID_CNAE (1207/1213), então deve ser
    aplicado depois dele.
    """

    if normalizar_texto(tipo_tributacao or "").strip() == "simples nacional mei":
        return False
    return deve_marcar


def _eh_prestador_fortaleza_ce(documento: "DocumentoPortalISS") -> bool:
    """True se o prestador estiver estabelecido em Fortaleza/CE."""

    cidade_norm = normalizar_texto(documento.cidade_prestador).upper()
    uf_norm = normalizar_texto(documento.uf_prestador).upper()
    return uf_norm == "CE" and "FORTALEZA" in cidade_norm


def limpar_cnpj_para_digitacao(cnpj: str) -> str:
    """Retorna o CNPJ pronto para digitação no portal.

    Remove apenas a máscara de pontuação (".", "-", "/", espaços) e qualquer
    outro caractere que não seja letra ou dígito, preservando as letras do
    novo formato alfanumérico de CNPJ (Receita Federal). O resultado é
    normalizado para maiúsculas, conforme o layout oficial (12 caracteres
    alfanuméricos + 2 dígitos verificadores numéricos).

    Não usar re.sub(r"\\D+", ...) aqui: isso destruiria as letras do CNPJ
    alfanumérico.
    """

    texto = (cnpj or "").strip().upper()
    return re.sub(r"[^0-9A-Z]", "", texto)


def _eh_cpf_prestador(documento: DocumentoPortalISS) -> bool:
    """True se o prestador for pessoa física.

    Fonte primária: a coluna TIPO_CLIENTE da planilha (explícita, auditável,
    gravada por processamento_xml.py para os dois layouts de XML suportados).
    Fallback: comprimento do CPF/CNPJ (11 dígitos = CPF, 14 = CNPJ, mesmo
    alfanumérico), usado apenas quando a coluna estiver ausente ou vazia —
    planilhas geradas antes dessa coluna existir, por exemplo.
    """

    tipo_norm = normalizar_texto(documento.tipo_cliente_prestador or "").strip()
    if tipo_norm.startswith("pessoa fisica"):
        return True
    if tipo_norm.startswith("pessoa juridica"):
        return False
    return len(limpar_cnpj_para_digitacao(documento.cnpj_prestador or "")) == 11


def validar_campos_obrigatorios(documento: DocumentoPortalISS) -> list[str]:
    """Valida se todos os campos obrigatórios necessários para a escrituração e cadastro do prestador
    estão presentes na planilha."""
    ausentes = []
    
    campos_validar = {
        "cnpj_prestador": "CNPJ do Prestador",
        "numero_nf": "Número da Nota Fiscal",
        "data_emissao": "Data de Emissão",
        "id_cnae_final": "Código CNAE (Atividade)",
        "descricao_servico": "Descrição do Serviço",
        "uf_local_prestacao": "UF do Local de Prestação",
        "cidade_local_prestacao": "Cidade do Local de Prestação",
        "natureza_operacao": "Natureza da Operação",
        "iss_retido": "ISS Retido (Sim/Não)",
        "nome_prestador": "Razão Social / Nome do Prestador",
        "uf_prestador": "UF do Prestador",
        "cidade_prestador": "Cidade do Prestador",
        "cep_prestador": "CEP do Prestador",
        "logradouro_prestador": "Logradouro (Endereço) do Prestador",
        "bairro_prestador": "Bairro do Prestador"
    }
    
    for attr, label in campos_validar.items():
        valor = getattr(documento, attr, None)
        if valor is None:
            ausentes.append(label)
        elif isinstance(valor, str) and not valor.strip():
            ausentes.append(label)
            
    # Validação adicional: garante que ISS_RETIDO tenha formato reconhecível (Sim/Não).
    # Sem isso, um valor ambíguo (ex.: "talvez", "S", célula com espaço/typo) seria
    # tratado como "Não retido" no portal sem qualquer aviso ao operador.
    if documento.iss_retido and documento.iss_retido.strip():
        _, iss_retido_reconhecido = _interpretar_iss_retido(documento.iss_retido)
        if not iss_retido_reconhecido:
            ausentes.append(
                f"ISS Retido (valor '{documento.iss_retido}' não reconhecido; use apenas 'Sim' ou 'Não')"
            )

    # Validação do valor do serviço
    val_serv = (documento.valor_servico or "").strip().replace(" ", "").replace(".", "").replace(",", ".")
    try:
        if not val_serv or float(val_serv) <= 0.0:
            ausentes.append("Valor do Serviço (deve ser maior que zero)")
    except ValueError:
        ausentes.append("Valor do Serviço (formato numérico inválido)")
        
    return ausentes


def _formatar_celula_para_string_de_valor(valor: Any) -> str:
    """Normaliza valores numéricos lidos do Excel para formato de string brasileiro."""
    if valor is None:
        return "0,00"
    if isinstance(valor, (int, float)):
        return f"{valor:.2f}".replace(".", ",")
    
    texto = str(valor).strip()
    if not texto:
        return "0,00"
    # Se for string formatada em padrão americano com ponto decimal, ex: "1500.50"
    if "." in texto and "," not in texto:
        try:
            val_f = float(texto)
            return f"{val_f:.2f}".replace(".", ",")
        except ValueError:
            pass
    return texto


def _normalizar_celula_cnpj(valor: Any) -> str:
    """Normaliza a célula CNPJ_PRESTADOR lida do Excel para uma string consistente.

    openpyxl entrega o tipo que o Excel detectou para a célula: int/float quando
    a célula "parece" um número (típico de CNPJ numérico legado, o que pode ter
    descartado um zero à esquerda), ou str em qualquer outro caso — incluindo
    sempre o novo CNPJ alfanumérico, já que o Excel não converte letras em número.

    Valores impossíveis para um CNPJ (célula vazia, booleano, zero ou negativo)
    retornam string vazia de propósito: assim `validar_campos_obrigatorios()`
    continua acusando "CNPJ do Prestador" ausente e a nota é registrada em
    `log_notas_incompletas.txt` em vez de ser escriturada com um CNPJ inventado.
    """

    if valor is None or isinstance(valor, bool):
        return ""
    if isinstance(valor, (int, float)):
        if valor <= 0:
            return ""
        # zfill repõe zeros à esquerda que o Excel descartou ao tratar a célula
        # como número (ex.: 191 -> "00000000000191", CNPJ legado válido).
        return str(int(round(valor))).zfill(14)
    return str(valor).strip()


def limpar_valor_para_digitacao(valor: str) -> str:
    """Retorna o valor em formato digitável para o campo mascarado do portal."""

    texto = (valor or "").strip()
    texto = texto.replace("R$", "").strip()
    texto = texto.replace(".", "")
    return texto


def limpar_numero_para_digitacao(numero: str) -> str:
    """Mantém apenas dígitos para o campo de número do documento."""

    return re.sub(r"\D+", "", numero or "")


def _somente_digitos(texto: str) -> str:
    """Normaliza qualquer valor para comparação por dígitos."""

    return re.sub(r"\D+", "", texto or "")


def _texto_indica_prestador_nao_encontrado(texto: str) -> bool:
    """Reconhece mensagens visíveis que indicam ausência de razão social no portal."""

    texto_normalizado = normalizar_texto(texto or "")
    marcadores = (
        "nenhuma razao social encontrada",
        "nenhuma razao social",
        "nenhum resultado encontrado",
        "nao foram encontrados",
    )
    return any(marcador in texto_normalizado for marcador in marcadores)


# ---------------------------------------------------------------------------
# Logging da função 2
# ---------------------------------------------------------------------------


def registrar_log_funcao2(caminho_log: Path, mensagem: str) -> None:
    """Registra eventos da função 2 com data e hora para rastrear as linhas ignoradas."""

    caminho_log.parent.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    with caminho_log.open("a", encoding="utf-8") as arquivo:
        arquivo.write(f"[{timestamp}] {mensagem}\n")


# ---------------------------------------------------------------------------
# Competência
# ---------------------------------------------------------------------------


def construir_competencia(mes: int, ano: int) -> CompetenciaTrabalho:
    """Valida mês e ano e devolve a competência estruturada."""

    if not 1 <= mes <= 12:
        raise ValueError("O mês deve estar entre 1 e 12.")
    if ano < 1000 or ano > 9999:
        raise ValueError("O ano deve estar no formato AAAA.")
    return CompetenciaTrabalho(mes=mes, ano=ano, rotulo=f"{mes:02d}/{ano}")


def interpretar_competencia(texto: str) -> CompetenciaTrabalho:
    """Converte uma entrada do usuário em mês e ano.

    Aceita entradas como `maio 2026`, `maio/2026` ou `05/2026`.
    """

    texto_limpo = texto.strip()
    if not texto_limpo:
        raise ValueError("Competência vazia.")

    padrao_numerico = re.fullmatch(r"(\d{1,2})\s*(?:/|\s)\s*(\d{4})", texto_limpo)
    if padrao_numerico:
        mes = int(padrao_numerico.group(1))
        ano = int(padrao_numerico.group(2))
        return construir_competencia(mes, ano)

    texto_norm = normalizar_texto(texto_limpo)
    padrao_nome = re.fullmatch(r"([a-z]+)\s+de\s+(\d{4})", texto_norm) or re.fullmatch(
        r"([a-z]+)\s+(\d{4})", texto_norm
    )
    if not padrao_nome:
        raise ValueError(
            "Informe a competência no formato 'maio 2026' ou '05/2026'."
        )

    nome_mes = padrao_nome.group(1)
    ano = int(padrao_nome.group(2))
    mes = MESES_POR_NOME.get(nome_mes)
    if mes is None:
        raise ValueError(
            "Não foi possível identificar o mês informado. Use, por exemplo, 'maio 2026'."
        )
    return construir_competencia(mes, ano)








# ---------------------------------------------------------------------------
# Carregamento da planilha
# ---------------------------------------------------------------------------


def carregar_documentos_xlsx(caminho_xlsx: Path) -> list[DocumentoPortalISS]:
    """Lê a planilha de notas e devolve as linhas prontas para o portal."""

    workbook = load_workbook(caminho_xlsx, data_only=True)
    planilha = workbook.active

    cabecalhos = [celula.value for celula in next(planilha.iter_rows(min_row=1, max_row=1))]
    colunas = {nome: indice for indice, nome in enumerate(cabecalhos)}

    # Suporte flexível para arquivo_xml ou arquivo_pdf
    col_arquivo = "ARQUIVO_XML" if "ARQUIVO_XML" in colunas else ("ARQUIVO_PDF" if "ARQUIVO_PDF" in colunas else None)
    if not col_arquivo:
        raise ValueError("A planilha informada deve conter a coluna 'ARQUIVO_XML' ou 'ARQUIVO_PDF'.")

    campos_obrigatorios = [
        "CNPJ_PRESTADOR",
        "NUMERO_NF",
        "ID_CNAE_FINAL",
        "DATA_EMISSAO",
        "DESCRICAO_SERVICO",
        "UF_LOCAL_PRESTACAO",
        "CIDADE_LOCAL_PRESTACAO",
        "NATUREZA_OPERACAO",
        "ISS_RETIDO",
        "VALOR_SERVICO",
    ]
    faltantes = [campo for campo in campos_obrigatorios if campo not in colunas]
    if faltantes:
        raise ValueError(
            f"A planilha informada não contém as colunas obrigatórias: {', '.join(faltantes)}"
        )

    documentos: list[DocumentoPortalISS] = []
    for linha in planilha.iter_rows(min_row=2, values_only=True):
        if not any(valor is not None and str(valor).strip() for valor in linha):
            continue

        valor_servico = _formatar_celula_para_string_de_valor(linha[colunas["VALOR_SERVICO"]]) if "VALOR_SERVICO" in colunas else "0,00"
        valor_deducoes = _formatar_celula_para_string_de_valor(linha[colunas["VALOR_DEDUCOES"]]) if "VALOR_DEDUCOES" in colunas else "0,00"
        descontos_incondicionados = _formatar_celula_para_string_de_valor(linha[colunas["DESCONTOS_INCONDICIONADOS"]]) if "DESCONTOS_INCONDICIONADOS" in colunas else "0,00"
        descontos_condicionados = _formatar_celula_para_string_de_valor(linha[colunas["DESCONTOS_CONDICIONADOS"]]) if "DESCONTOS_CONDICIONADOS" in colunas else "0,00"
        outras_retencoes = _formatar_celula_para_string_de_valor(linha[colunas["OUTRAS_RETENCOES"]]) if "OUTRAS_RETENCOES" in colunas else "0,00"
        ir = _formatar_celula_para_string_de_valor(linha[colunas["IR"]]) if "IR" in colunas else "0,00"
        pis_nao_retido = _formatar_celula_para_string_de_valor(linha[colunas["PIS_NAO_RETIDO"]]) if "PIS_NAO_RETIDO" in colunas else "0,00"
        cofins_nao_retido = _formatar_celula_para_string_de_valor(linha[colunas["COFINS_NAO_RETIDO"]]) if "COFINS_NAO_RETIDO" in colunas else "0,00"
        csrf = _formatar_celula_para_string_de_valor(linha[colunas["CSRF (CSLL + PIS + Cofins Retidos)"]]) if "CSRF (CSLL + PIS + Cofins Retidos)" in colunas else "0,00"
        inss = _formatar_celula_para_string_de_valor(linha[colunas["INSS"]]) if "INSS" in colunas else "0,00"
        aliquota = _formatar_celula_para_string_de_valor(linha[colunas["ALIQUOTA"]]) if "ALIQUOTA" in colunas else ""

        # Leitura flexível da prefeitura
        prefeitura_val = str(linha[colunas["PREFEITURA"]] or "") if "PREFEITURA" in colunas else ""

        documentos.append(
            DocumentoPortalISS(
                arquivo_pdf=str(linha[colunas[col_arquivo]] or ""),
                prefeitura=prefeitura_val,
                cnpj_prestador=_normalizar_celula_cnpj(linha[colunas["CNPJ_PRESTADOR"]]),
                numero_nf=str(linha[colunas["NUMERO_NF"]] or ""),
                id_cnae_final=str(linha[colunas["ID_CNAE_FINAL"]] or ""),
                data_emissao=str(linha[colunas["DATA_EMISSAO"]] or ""),
                descricao_servico=str(linha[colunas["DESCRICAO_SERVICO"]] or ""),
                uf_local_prestacao=str(linha[colunas["UF_LOCAL_PRESTACAO"]] or ""),
                cidade_local_prestacao=str(linha[colunas["CIDADE_LOCAL_PRESTACAO"]] or ""),
                natureza_operacao=str(linha[colunas["NATUREZA_OPERACAO"]] or ""),
                iss_retido=str(linha[colunas["ISS_RETIDO"]] or ""),
                valor_servico=valor_servico,
                # Campos de endereço do prestador (lidos quando presentes na planilha)
                nome_prestador=str(linha[colunas["NOME_PRESTADOR"]] or "") if "NOME_PRESTADOR" in colunas else "",
                uf_prestador=str(linha[colunas["UF_PRESTADOR"]] or "") if "UF_PRESTADOR" in colunas else "",
                cidade_prestador=str(linha[colunas["CIDADE_PRESTADOR"]] or "") if "CIDADE_PRESTADOR" in colunas else "",
                cep_prestador=str(linha[colunas["CEP_PRESTADOR"]] or "") if "CEP_PRESTADOR" in colunas else "",
                logradouro_prestador=str(linha[colunas["LOGRADOURO_PRESTADOR"]] or "") if "LOGRADOURO_PRESTADOR" in colunas else "",
                numero_prestador=str(linha[colunas["NUMERO_PRESTADOR"]] or "") if "NUMERO_PRESTADOR" in colunas else "",
                bairro_prestador=str(linha[colunas["BAIRRO_PRESTADOR"]] or "") if "BAIRRO_PRESTADOR" in colunas else "",
                email_prestador=str(linha[colunas["EMAIL_PRESTADOR"]] or "") if "EMAIL_PRESTADOR" in colunas else "",
                valor_deducoes=valor_deducoes,
                descontos_incondicionados=descontos_incondicionados,
                descontos_condicionados=descontos_condicionados,
                outras_retencoes=outras_retencoes,
                ir=ir,
                pis_nao_retido=pis_nao_retido,
                cofins_nao_retido=cofins_nao_retido,
                csrf=csrf,
                inss=inss,
                aliquota=aliquota,
                regime_tributario=str(linha[colunas["REGIME_TRIBUTARIO"]] or "OUTROS") if "REGIME_TRIBUTARIO" in colunas else "OUTROS",
                tipo_cliente_prestador=str(linha[colunas["TIPO_CLIENTE"]] or "") if "TIPO_CLIENTE" in colunas else "",
                id_cnae=str(linha[colunas["ID_CNAE"]] or "") if "ID_CNAE" in colunas else "",
                tipo_tributacao=str(linha[colunas["TIPO_TRIBUTACAO"]] or "Normal") if "TIPO_TRIBUTACAO" in colunas else "Normal",
                cnpj_tomador=_normalizar_celula_cnpj(linha[colunas["CNPJ_TOMADOR"]]) if "CNPJ_TOMADOR" in colunas else "",
            )
        )
    return documentos












# ---------------------------------------------------------------------------
# Esperas e cliques com Selenium
# ---------------------------------------------------------------------------


def _aguardar_elemento(
    driver: WebDriver,
    by: str,
    seletor: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda um elemento ficar visível e retorna o WebElement."""

    return WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((by, seletor))
    )


def _aguardar_elemento_clicavel(
    driver: WebDriver,
    by: str,
    seletor: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda um elemento ficar clicável (visível + habilitado) e o retorna."""

    return WebDriverWait(driver, timeout).until(
        EC.element_to_be_clickable((by, seletor))
    )


def _clicar_com_espera(
    driver: WebDriver,
    by: str,
    seletor: str,
    descricao: str,
    timeout: int = 10,
) -> WebElement:
    """Aguarda o elemento e efetua o clique, com fallback via JavaScript e retry para elementos obsoletos."""

    _verificar_pausa()
    _verificar_cancelamento()
    for tentativa in range(3):
        try:
            elemento = _aguardar_elemento_clicavel(driver, by, seletor, timeout=timeout)
            try:
                elemento.click()
            except (ElementNotInteractableException, WebDriverException) as exc:
                if isinstance(exc, StaleElementReferenceException):
                    raise
                # Fallback via JavaScript para elementos que bloqueiam o clique nativo
                driver.execute_script("arguments[0].click();", elemento)
            return elemento
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)
    # Sentinela: garante que falhas de timeout ou inesperadas sejam exibidas ao chamador
    raise RuntimeError(f"Não foi possível clicar no elemento '{descricao}' após 3 tentativas.")


def _clicar_por_id(driver: WebDriver, elemento_id: str, timeout: int = 10) -> None:
    """Clica em um elemento pelo ID com fallback via JavaScript."""

    _clicar_com_espera(driver, By.ID, elemento_id, f"elemento ID={elemento_id}", timeout=timeout)


# O portal usa o MESMO componente de modal (RichFaces) para vários avisos de
# confirmação ao clicar em Gravar — só o texto interno muda. Exemplos já
# observados: "Já existe documento fiscal escriturado com o CNPJ do
# prestador, com o mesmo número de nota, na competência ..." (possível
# duplicata) e "Prestador não inscrito no CPOM. ..." (aviso informativo).
_XPATH_MODAL_CONFIRMACAO_CONTAINER = (
    '//*[@id="digitarDocumentoForm:confirmacao_customizadaContainer"]'
)
_XPATH_MODAL_CONFIRMACAO_H3 = _XPATH_MODAL_CONFIRMACAO_CONTAINER + '//h3'
_XPATH_HEADING_SUCESSO_GRAVACAO = (
    '//*[@id="content"]/legend/h2[contains(., "digitado com") '
    'and (contains(., "Sucesso") or contains(., "sucesso"))]'
)

# Marcador de texto (normalizado, sem acento/caixa) que identifica especificamente
# o aviso de possível nota duplicada — o único caso em que a decisão é "Não".
# Qualquer outro texto nesse mesmo componente é tratado como aviso informativo
# (decisão "Sim").
_MARCADOR_TEXTO_NOTA_DUPLICADA = "documento fiscal escriturado"

# Mensagem de validação inline (NÃO é o modal de confirmação acima — é o
# componente <rich:messages> do JSF, sem botão, que bloqueia a gravação até o
# formulário ser corrigido): "CNAE não incide Imposto Sobre Serviço. Solução:
# Selecione a Natureza de Operação 'Não Incidência'." A correção é trocar a
# Natureza da Operação e tentar gravar de novo — não há botão para clicar aqui.
_XPATH_MENSAGEM_CNAE_NAO_INCIDE = (
    '//span[contains(@class, "rich-messages-label") and contains(., "CNAE não incide")]'
)


def _clicar_botao_modal_confirmacao(driver: WebDriver, valor_botao: str) -> None:
    """Clica no botão ("Sim" ou "Não") do modal de confirmação genérico do
    portal e aguarda o modal desaparecer da tela antes de prosseguir."""

    xpath_botao = f'{_XPATH_MODAL_CONFIRMACAO_CONTAINER}//input[@value="{valor_botao}"]'
    try:
        driver.find_element(By.XPATH, xpath_botao).click()
    except Exception:
        # Fallback pelos IDs auto-gerados pelo JSF observados no momento da implementação
        # (frágil a mudanças de versão do portal, por isso o XPath acima é a via principal).
        id_fallback = "digitarDocumentoForm:j_id489" if valor_botao == "Sim" else "digitarDocumentoForm:j_id493"
        driver.find_element(By.ID, id_fallback).click()

    WebDriverWait(driver, 10).until(
        EC.invisibility_of_element_located((By.XPATH, _XPATH_MODAL_CONFIRMACAO_H3))
    )


def _aguardar_resultado_da_gravacao(driver: WebDriver, timeout: int = 15) -> str:
    """Aguarda o desfecho do clique em "Gravar": sucesso normal, ou algum dos
    modais de confirmação que o portal pode exibir antes de gravar de fato.

    - Se o texto do modal indicar possível nota duplicada (mesmo CNPJ do
      prestador + mesmo número de nota já escriturados em outra competência),
      clica em "Não" (decisão do operador: nunca confirmar a geração de uma
      possível duplicata automaticamente) e retorna "duplicata".
    - Qualquer outro texto de confirmação (ex.: "Prestador não inscrito no
      CPOM...") é tratado como aviso informativo: clica em "Sim" e retorna
      "confirmacao_generica" — cabe ao chamador clicar em Gravar de novo.
    - Se a mensagem de validação "CNAE não incide Imposto Sobre Serviço..."
      aparecer (sem botão, bloqueia a gravação), retorna "cnae_nao_incide" —
      cabe ao chamador trocar a Natureza da Operação e tentar gravar de novo.
    - Quando o cabeçalho de sucesso aparece, retorna "sucesso".
    - Se nada aparecer dentro do prazo, levanta TimeoutException.
    """

    fim = time.monotonic() + timeout
    while time.monotonic() < fim:
        _verificar_cancelamento()

        if driver.find_elements(By.XPATH, _XPATH_HEADING_SUCESSO_GRAVACAO):
            return "sucesso"

        elementos_cnae = driver.find_elements(By.XPATH, _XPATH_MENSAGEM_CNAE_NAO_INCIDE)
        if any(el.is_displayed() for el in elementos_cnae):
            return "cnae_nao_incide"

        elementos_modal = driver.find_elements(By.XPATH, _XPATH_MODAL_CONFIRMACAO_H3)
        elementos_modal = [el for el in elementos_modal if el.is_displayed()]
        if elementos_modal:
            texto_modal = normalizar_texto(elementos_modal[0].text or "")
            if _MARCADOR_TEXTO_NOTA_DUPLICADA in texto_modal:
                registrar_evento_execucao(
                    "Portal indicou que já existe documento fiscal escriturado com o mesmo "
                    "CNPJ do prestador e o mesmo número de nota (em outra competência). "
                    "Clicando em 'Não' para não gravar uma possível duplicata.",
                    "ISS Fortaleza",
                )
                _clicar_botao_modal_confirmacao(driver, "Não")
                return "duplicata"

            registrar_evento_execucao(
                f"Portal exibiu confirmação: '{elementos_modal[0].text.strip()}'. "
                "Clicando em 'Sim'.",
                "ISS Fortaleza",
            )
            _clicar_botao_modal_confirmacao(driver, "Sim")
            return "confirmacao_generica"

        time.sleep(0.3)

    raise TimeoutException("Timeout aguardando confirmação de gravação do documento.")


def _clicar_gravar_documento(driver: WebDriver) -> None:
    """Clica no botão "Gravar Documento", com fallback via XPath por valor.

    ÚNICO ponto do módulo que grava uma nota. Bloqueado por padrão: só funciona dentro de
    `gravacao_habilitada(True)`."""

    if not _GRAVACAO_PERMITIDA:
        raise GravacaoBloqueadaError(
            "Gravação bloqueada: o robô está em modo simulação (ou fora de gravacao_habilitada(True)). "
            "Nada foi gravado no portal."
        )
    try:
        _clicar_por_id(driver, "digitarDocumentoForm:j_id477", timeout=15)
    except Exception:
        el = driver.find_element(
            By.XPATH,
            "//*[@id='digitarDocumentoForm:j_id477'] | //input[@value='Gravar' or @value='Gravar Documento']",
        )
        driver.execute_script("arguments[0].click();", el)


def _gravar_documento_com_confirmacoes(
    driver: WebDriver, numero_nf: str, timeout: int = 15, max_confirmacoes: int = 4
) -> str:
    """Clica em Gravar Documento e resolve os bloqueios que o portal pode
    encadear antes de gravar de fato: nota duplicada (clica "Não", nunca
    grava), avisos informativos como "Prestador não inscrito no CPOM" (clica
    "Sim" e tenta Gravar de novo), ou o erro de validação "CNAE não incide
    Imposto Sobre Serviço" (troca a Natureza da Operação para "Não
    Incidência" e tenta Gravar de novo) — repete até `max_confirmacoes`
    vezes, já que o portal pode encadear mais de um bloqueio para a mesma
    nota. Retorna "sucesso" ou "duplicata".
    """

    _clicar_gravar_documento(driver)
    for _ in range(max_confirmacoes):
        resultado = _aguardar_resultado_da_gravacao(driver, timeout=timeout)
        if resultado in ("sucesso", "duplicata"):
            return resultado

        if resultado == "cnae_nao_incide":
            # O portal não deixa gravar esse CNAE com a Natureza atual; a correção
            # é forçar "Não Incidência" (sobrescrevendo o que a planilha continha
            # para esta nota) e tentar gravar de novo. O checkbox de ISS Retido
            # não precisa de tratamento aqui: o próprio portal o remove da tela ao
            # selecionar "Não Incidência".
            registrar_evento_execucao(
                f"Portal recusou a gravação da NF {numero_nf}: CNAE não incide ISS. "
                "Selecionando Natureza da Operação = 'Não Incidência' (sobrescrevendo "
                "o valor da planilha para esta nota) e tentando gravar novamente.",
                "ISS Fortaleza",
            )
            _selecionar_opcao_por_texto(
                driver,
                By.ID,
                "digitarDocumentoForm:comboEscolherLocalPrestacao",
                ["Não Incidência", "Nao Incidencia"],
            )
            # Mesma pausa de estabilização de AJAX já usada para esse campo em
            # preencher_documento_servico().
            time.sleep(1.5)
        else:
            # resultado == "confirmacao_generica": _aguardar_resultado_da_gravacao já
            # clicou "Sim" no modal; tenta gravar de novo para ver o próximo desfecho.
            registrar_evento_execucao(
                f"Confirmação genérica do portal resolvida (Sim) para a NF {numero_nf}; "
                "tentando Gravar Documento novamente.",
                "ISS Fortaleza",
            )

        _clicar_gravar_documento(driver)

    raise TimeoutException(
        f"Excedeu o limite de {max_confirmacoes} confirmações encadeadas ao gravar a NF {numero_nf}."
    )


def _pausar_para_depuracao(
    driver: WebDriver,
    habilitado: bool,
    mensagem: str,
) -> None:
    """Interrompe o fluxo para inspeção manual quando o modo de depuração estiver ativo."""

    if not habilitado:
        return
    _emitir(mensagem)
    if sys.stdin.isatty():
        input(f"{mensagem}\nPressione Enter para continuar com o próximo clique...")


# ---------------------------------------------------------------------------
# Digitação e seleção de campos
# ---------------------------------------------------------------------------


def _digitar_campo(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
    delay_ms: int = 80,
    normalizador_fallback: Callable[[str], str] | None = None,
) -> None:
    """Localiza o campo, limpa e digita o texto tecla por tecla com pequena pausa entre elas.

    Possui lógica de retry para se recuperar caso o elemento mude sob AJAX durante a digitação.

    O parâmetro opcional `normalizador_fallback` substitui a normalização padrão
    (somente dígitos) usada na verificação de estabilidade do campo. É necessário
    para campos que podem conter letras (ex.: CNPJ alfanumérico): usar apenas
    dígitos como fallback nesse caso mascararia silenciosamente uma eventual perda
    de letras durante a digitação. Quando omitido, o comportamento é idêntico ao
    anterior para todos os outros campos.
    """

    _verificar_pausa()
    _verificar_cancelamento()
    normalizar = normalizador_fallback or _somente_digitos
    for tentativa in range(3):
        try:
            campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
            campo.click()
            campo.send_keys(Keys.CONTROL + "a")
            campo.send_keys(Keys.BACKSPACE)
            time.sleep(0.12)

            # Digita caractere por caractere para disparar os eventos key* do portal
            for caractere in texto:
                campo.send_keys(caractere)
                time.sleep(delay_ms / 1000)

            # Verifica se o valor ficou estável antes de seguir
            esperado_normalizado = normalizar(texto)
            for _ in range(20):
                try:
                    valor_atual = campo.get_attribute("value") or ""
                    if valor_atual.strip() == texto.strip():
                        return
                    if esperado_normalizado and normalizar(valor_atual) == esperado_normalizado:
                        return
                except StaleElementReferenceException:
                    raise
                time.sleep(0.15)

            raise RuntimeError(f"O campo {seletor} não permaneceu com o valor esperado: {texto}")
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)


def _definir_valor_instantaneo(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
) -> None:
    """Define o valor de um campo de texto instantaneamente executando JavaScript no navegador.

    Isso é utilizado para campos de textos muito longos, como a descrição dos serviços,
    evitando o gargalo de tempo da digitação caractere por caractere via Selenium.
    Dispara os eventos 'input' e 'change' para garantir que os ouvintes de eventos da página
    e o estado do JSF reconheçam a alteração do valor.
    """
    _verificar_pausa()
    _verificar_cancelamento()
    for tentativa in range(3):
        try:
            campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
            driver.execute_script(
                "arguments[0].value = arguments[1];"
                "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
                "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
                campo,
                texto,
            )
            return
        except StaleElementReferenceException:
            if tentativa == 2:
                raise
            time.sleep(0.5)


def _definir_campo_rapido(
    driver: WebDriver,
    by: str,
    seletor: str,
    texto: str,
    delay_ms: int = 80,
    normalizador_fallback: Callable[[str], str] | None = None,
) -> None:
    """Tenta preencher o campo instantaneamente via JavaScript (bem mais rápido
    que digitar caractere por caractere) e confere se o valor realmente ficou
    correto. Se não ficou — vazio, ou diferente do esperado (ex.: alguma
    validação/máscara do site descartando o valor) — cai para `_digitar_campo`
    (digitação caractere por caractere, comportamento de sempre) como fallback
    garantido.
    """

    _verificar_pausa()
    _verificar_cancelamento()
    normalizar = normalizador_fallback or _somente_digitos

    try:
        campo = _aguardar_elemento_clicavel(driver, by, seletor, timeout=10)
        driver.execute_script(
            "arguments[0].value = arguments[1];"
            "arguments[0].dispatchEvent(new Event('input', { bubbles: true }));"
            "arguments[0].dispatchEvent(new Event('change', { bubbles: true }));",
            campo,
            texto,
        )
        time.sleep(0.3)  # pequena espera para eventual validação JS do portal reagir

        valor_atual = campo.get_attribute("value") or ""
        if valor_atual.strip() == texto.strip() or normalizar(valor_atual) == normalizar(texto):
            return
    except (StaleElementReferenceException, NoSuchElementException):
        pass  # cai para o fallback abaixo

    _digitar_campo(driver, by, seletor, texto, delay_ms=delay_ms, normalizador_fallback=normalizador_fallback)


def _campo_aliquota_esta_editavel(driver: WebDriver) -> WebElement | None:
    """Retorna o elemento do campo Alíquota se ele estiver editável (um <input>
    sem disabled/readonly), ou None se estiver bloqueado pelo portal.

    O campo é calculado automaticamente para alguns CNAEs — nesse caso o portal
    renderiza um <span> (não um <input>), sem nenhum campo para digitar. Só
    tentamos preencher quando existe de fato um <input> habilitado.
    """

    try:
        elemento = driver.find_element(By.ID, "digitarDocumentoForm:idAliquota")
    except NoSuchElementException:
        return None
    if elemento.tag_name.lower() != "input":
        return None
    if elemento.get_attribute("readonly") or elemento.get_attribute("disabled"):
        return None
    if "disabled" in (elemento.get_attribute("class") or ""):
        return None
    return elemento


def _ler_estado_checkbox_apos_espera(
    driver: WebDriver,
    by: str,
    seletor: str,
    espera: float,
):
    """Aguarda `espera` segundos e retorna checkbox.is_selected(), ou None se o
    elemento não for encontrado.

    Usado para dar tempo a um onchange AJAX assíncrono do portal (ex.: a
    mudança de Natureza da Operação, que pode marcar/desmarcar sozinho o
    checkbox de ISS Retido) terminar de processar antes de ler o estado real.
    Uma leitura por "debounce" (esperar até duas leituras seguidas baterem)
    foi cogitada e descartada: como o polling costuma acontecer mais rápido
    que a resposta do AJAX, duas leituras consecutivas podem coincidir
    ANTES da mudança tardia chegar, declarando "estável" cedo demais — uma
    espera de duração fixa (mesmo padrão já usado nesta função para Status
    NFSE e UF do Prestador) é mais confiável aqui.
    """

    time.sleep(espera)
    try:
        return driver.find_element(by, seletor).is_selected()
    except NoSuchElementException:
        return None


def _selecionar_opcao_por_texto(
    driver: WebDriver,
    by: str,
    seletor: str,
    textos: list[str],
    timeout: int = 10,
) -> None:
    """Seleciona a primeira opção disponível em um <select> a partir de uma lista de rótulos."""

    _verificar_pausa()
    _verificar_cancelamento()
    # Aumentado para 8 tentativas para maior tolerância ao AJAX do portal
    for tentativa in range(8):
        try:
            elemento = _aguardar_elemento_clicavel(driver, by, seletor, timeout=timeout)
            select = Select(elemento)

            # Se a lista de opções estiver vazia ou contiver apenas o placeholder inicial, forçamos o retry
            # pois o portal (RichFaces) provavelmente ainda está popularizando o <select> via AJAX assíncrono.
            if len(select.options) <= 1:
                raise NoSuchElementException("Select ainda não populado com opções (AJAX pendente).")

            # 1. Tenta correspondência exata normalizada (ignora caixa e acentos)
            for texto in textos:
                texto_norm = normalizar_texto(texto)
                for opcao in select.options:
                    opcao_texto_norm = normalizar_texto(opcao.text)
                    opcao_valor_norm = normalizar_texto(opcao.get_attribute("value") or "")
                    if (texto_norm == opcao_texto_norm) or (texto_norm == opcao_valor_norm):
                        select.select_by_value(opcao.get_attribute("value"))
                        return

            # 2. Tenta correspondência parcial normalizada (ex: "Fortaleza" contido em "Fortaleza - CE")
            for texto in textos:
                texto_norm = normalizar_texto(texto)
                for opcao in select.options:
                    opcao_texto_norm = normalizar_texto(opcao.text)
                    if texto_norm in opcao_texto_norm:
                        # Ignora placeholders de instruções como "selecione" ou "escolha"
                        if "selecione" in opcao_texto_norm or "escolha" in opcao_texto_norm:
                            continue
                        select.select_by_value(opcao.get_attribute("value"))
                        return

            raise NoSuchElementException(f"Nenhuma das opções {textos} foi localizada no select.")
            
        except (StaleElementReferenceException, NoSuchElementException) as exc:
            if tentativa == 7:
                raise RuntimeError(
                    f"Não foi possível selecionar nenhuma opção entre {textos} após várias tentativas."
                ) from exc
            time.sleep(0.8)


# ---------------------------------------------------------------------------
# Aguardando login manual
# ---------------------------------------------------------------------------






# ---------------------------------------------------------------------------
# Seleção de competência via calendário RichFaces
# ---------------------------------------------------------------------------


def _valor_campo_competencia(driver: WebDriver, base_id: str) -> str:
    """Lê o valor do campo de data do calendário RichFaces."""

    try:
        campo = driver.find_element(By.ID, f"{base_id}InputDate")
        return campo.get_attribute("value") or ""
    except NoSuchElementException:
        return ""


def _abrir_editor_calendario(driver: WebDriver, base_id: str, rotulo: str) -> None:
    """Clica no botão de edição do calendário RichFaces para abrir o editor de mês/ano.

    Garante que o calendário popup seja aberto antes de tentar acessar seu cabeçalho.
    """

    # Clica no botão popup do calendário (ícone do calendário) para exibir o calendário na tela.
    # Sem isso, o cabeçalho '#{base_id}Header' não é renderizado ou fica invisível no DOM.
    popup_btn_id = f"{base_id}PopupButton"
    try:
        popup_btn = WebDriverWait(driver, 8).until(
            EC.element_to_be_clickable((By.ID, popup_btn_id))
        )
        popup_btn.click()
        # Pequena pausa para garantir a renderização visual do calendário popup
        time.sleep(0.5)
    except Exception as e:
        _emitir(f"Tive um problema ao abrir o calendário de {rotulo}, mas vou tentar seguir em frente: {e}")

    # O botão de edição fica no cabeçalho do calendário.
    # Atenção: o base_id contém ':' (ex: 'manterEscrituracaoForm:dataInicial'),
    # portanto não pode ser usado diretamente como seletor CSS '#id'. Usamos o seletor
    # de atributo [id='...'] para contornar isso sem necessidade de escape manual.
    botao = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable(
            (By.CSS_SELECTOR, f"[id='{base_id}Header'] .rich-calendar-tool-btn")
        )
    )
    botao.click()
    # Aguarda o botão OK do editor ficar visível como sinal de abertura do editor
    WebDriverWait(driver, 10).until(
        EC.visibility_of_element_located((By.ID, f"{base_id}DateEditorButtonOk"))
    )


def _selecionar_mes_no_editor(driver: WebDriver, base_id: str, competencia: CompetenciaTrabalho) -> None:
    """Clica no mês desejado na grade do editor do calendário RichFaces."""

    indice_mes = competencia.mes - 1
    id_mes = f"{base_id}DateEditorLayoutM{indice_mes}"
    _clicar_por_id(driver, id_mes)


def _selecionar_ano_no_editor(driver: WebDriver, base_id: str, competencia: CompetenciaTrabalho) -> None:
    """Navega pelos anos do editor do calendário RichFaces e clica no ano desejado."""

    for _ in range(20):
        anos_visiveis: list[tuple[int, WebElement]] = []
        for indice in range(10):
            id_ano = f"{base_id}DateEditorLayoutY{indice}"
            try:
                el = driver.find_element(By.ID, id_ano)
                if not el.is_displayed():
                    continue
                texto = el.text.strip()
                if texto.isdigit():
                    anos_visiveis.append((int(texto), el))
            except NoSuchElementException:
                continue

        for ano_visivel, el_ano in anos_visiveis:
            if ano_visivel == competencia.ano:
                el_ano.click()
                return

        if not anos_visiveis:
            raise RuntimeError(
                f"Não consegui ler os anos visíveis no calendário de competência ({base_id})."
            )

        menor_ano = min(a for a, _ in anos_visiveis)
        maior_ano = max(a for a, _ in anos_visiveis)

        # Navega para esquerda ou direita conforme o ano-alvo
        if competencia.ano < menor_ano:
            sinal = "<"
        else:
            sinal = ">"

        try:
            # Seletor por atributo para evitar o erro de CSS com ':' no ID
            botoes = driver.find_elements(By.CSS_SELECTOR, f"[id='{base_id}'] .rich-calendar-editor-btn")
            clicou = False
            for btn in botoes:
                if btn.text.strip() == sinal and btn.is_displayed():
                    btn.click()
                    clicou = True
                    break
            if not clicou:
                raise RuntimeError("Botão de navegação do calendário não encontrado.")
        except Exception as exc:
            raise RuntimeError(
                f"Não consegui navegar no calendário para chegar ao ano {competencia.ano}."
            ) from exc

        time.sleep(0.25)

    raise RuntimeError(
        f"Não consegui localizar o ano {competencia.ano} no calendário de competência ({base_id})."
    )


def _confirmar_editor_calendario(driver: WebDriver, base_id: str, esperado: str) -> None:
    """Clica em OK no editor do calendário e confirma que o campo assumiu o valor esperado."""

    _clicar_por_id(driver, f"{base_id}DateEditorButtonOk")
    # Aguarda o campo de data refletir o valor confirmado
    WebDriverWait(driver, 10).until(
        lambda d: (
            lambda el: el is not None and el.get_attribute("value") == esperado
        )(d.find_element(By.ID, f"{base_id}InputDate"))
    )


def _selecionar_campo_competencia(
    driver: WebDriver,
    base_id: str,
    rotulo: str,
    competencia: CompetenciaTrabalho,
) -> None:
    """Seleciona o mês e o ano no calendário RichFaces de um campo de competência."""

    esperado = f"{competencia.mes:02d}/{competencia.ano}"
    valor_atual = _valor_campo_competencia(driver, base_id)

    if valor_atual == esperado:
        registrar_evento_execucao(
            f"Campo {rotulo} da competência já estava em {esperado}.",
            "ISS Fortaleza",
        )
        return

    _abrir_editor_calendario(driver, base_id, rotulo)
    _selecionar_mes_no_editor(driver, base_id, competencia)
    _selecionar_ano_no_editor(driver, base_id, competencia)
    _confirmar_editor_calendario(driver, base_id, esperado)
    registrar_evento_execucao(
        f"Campo {rotulo} da competência ajustado para {esperado}.",
        "ISS Fortaleza",
    )


def selecionar_competencia_na_tela_richfaces(
    driver: WebDriver,
    competencia: CompetenciaTrabalho,
) -> None:
    """Seleciona a competência nos calendários `De` e `Até` usando o editor do RichFaces."""

    # Garante que a tela carregou antes de tentar acessar os calendários
    WebDriverWait(driver, 120).until(
        EC.element_to_be_clickable((By.ID, "manterEscrituracaoForm:btnConsultar"))
    )
    time.sleep(1.2)

    _selecionar_campo_competencia(driver, "manterEscrituracaoForm:dataInicial", "De", competencia)
    _selecionar_campo_competencia(driver, "manterEscrituracaoForm:dataFinal", "Até", competencia)
    registrar_evento_execucao(
        f"Competência selecionada na tela: {competencia.rotulo}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Botão Escriturar
# ---------------------------------------------------------------------------


def aguardar_botao_escriturar(driver: WebDriver, timeout: int = 120) -> bool:
    """Espera a consulta retornar e verifica se o botão de escriturar está habilitado."""

    id_ativo = "manterEscrituracaoForm:dataTable:0:linkEscriturar"
    id_desabilitado = "manterEscrituracaoForm:dataTable:0:linkEscriturarDesabilitado"

    prazo = time.monotonic() + timeout
    while time.monotonic() < prazo:
        try:
            el = driver.find_element(By.ID, id_ativo)
            if el.is_displayed():
                return True
        except NoSuchElementException:
            pass
        try:
            el = driver.find_element(By.ID, id_desabilitado)
            if el.is_displayed():
                return False
        except NoSuchElementException:
            pass
        time.sleep(0.5)

    raise CompetenciaSemEscriturarDisponivelError(
        "Não consegui identificar se a competência consultada liberou o botão de escriturar."
    )


# ---------------------------------------------------------------------------
# Tela de Escrituração Fiscal
# ---------------------------------------------------------------------------


def aguardar_tela_escrituracao_fiscal(driver: WebDriver, timeout: int = 120) -> None:
    """Espera a tela de Escrituração Fiscal ficar pronta após o clique em Escriturar."""

    prazo = time.monotonic() + timeout
    while time.monotonic() < prazo:
        try:
            el = driver.find_element(By.ID, "tabEncerramentoEscrituracao")
            if el.is_displayed():
                return
        except NoSuchElementException:
            pass
        try:
            el = driver.find_element(By.ID, "aba_tomados_lbl")
            if el.is_displayed():
                return
        except NoSuchElementException:
            pass
        time.sleep(0.25)

    raise RuntimeError(
        "A tela de Escrituração Fiscal não ficou visível após o clique em Escriturar."
    )


# ---------------------------------------------------------------------------
# Aba Serviços Tomados
# ---------------------------------------------------------------------------


def clicar_aba_servicos_tomados(driver: WebDriver) -> None:
    """Clica na aba 'Serviços Tomados' com estratégia específica para o portal RichFaces."""

    # Aguarda a tela de Escrituração Fiscal aparecer primeiro
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "tabEncerramentoEscrituracao"))
    )
    try:
        _clicar_por_id(driver, "aba_tomados_lbl", timeout=120)
        return
    except Exception:
        pass

    # Fallback: busca o elemento pelo onclick
    try:
        el = driver.find_element(By.CSS_SELECTOR, '[onclick*="atualizarDadosDaAbaTomados"]')
        el.click()
    except NoSuchElementException:
        raise RuntimeError("Não foi possível clicar na aba 'Serviços Tomados'.")


# ---------------------------------------------------------------------------
# Tela Digitar Documento
# ---------------------------------------------------------------------------


def aguardar_tela_digitar_documento(driver: WebDriver, timeout: int = 120) -> None:
    """Espera a tela de Digitar Documento carregar por completo."""

    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm"))
    )
    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:divPesquisaTomador"))
    )
    # Aguarda o rádio de CNPJ aparecer como confirmação final de carregamento
    WebDriverWait(driver, timeout).until(
        EC.visibility_of_element_located(
            (By.ID, "digitarDocumentoForm:tipoPesquisaTomadorRb:1")
        )
    )


# ---------------------------------------------------------------------------
# Preenchimento manual dos dados do prestador
# ---------------------------------------------------------------------------


def preencher_dados_prestador(
    driver: WebDriver,
    documento: "DocumentoPortalISS",
    depuracao: bool = False,
) -> None:
    """Preenche manualmente o formulário de identificação do prestador com os dados da planilha.

    Este fluxo substitui a busca via autocomplete pelo CNPJ, permitindo escriturar
    prestadores que ainda não estão cadastrados no portal da ISS Fortaleza.
    Os dados utilizados são os extraídos pela IA durante o processamento das notas fiscais.
    """

    registrar_evento_execucao(
        f"Preenchendo dados do prestador manualmente para CNPJ {documento.cnpj_prestador}",
        "ISS Fortaleza",
    )
    aguardar_tela_digitar_documento(driver)

    # --- Tipo de Cliente/Fornecedor: "Pessoa Física" ou "Pessoa Jurídica" conforme o prestador ---
    eh_cpf = _eh_cpf_prestador(documento)
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherTipoNaturezaJuridica",
        ["Pessoa Física", "Pessoa Fisica"] if eh_cpf else ["Pessoa Jurídica", "Pessoa Juridica", "PJ"],
        timeout=10,
    )
    registrar_evento_execucao(
        f"Tipo Natureza Jurídica definido como {'Pessoa Física' if eh_cpf else 'Pessoa Jurídica'} "
        f"(prestador identificado por {'CPF' if eh_cpf else 'CNPJ'})",
        "ISS Fortaleza",
    )
    time.sleep(0.5)

    # --- CPF/CNPJ do Prestador ---
    cnpj_limpo = limpar_cnpj_para_digitacao(documento.cnpj_prestador)
    if cnpj_limpo:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idCPFCNPJ",
            cnpj_limpo,
            delay_ms=60,
            normalizador_fallback=limpar_cnpj_para_digitacao,
        )
        registrar_evento_execucao(f"CPF/CNPJ preenchido: {cnpj_limpo}", "ISS Fortaleza")

    # --- Nome/Denominação ---
    nome = documento.nome_prestador.strip()
    if nome:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idNome",
            nome,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Nome/Denominação preenchido: {nome}", "ISS Fortaleza")

    # --- UF do Prestador (dispara AJAX de carregamento de cidades) ---
    uf = documento.uf_prestador.strip().upper()
    if uf:
        _selecionar_opcao_por_texto(
            driver,
            By.ID,
            "digitarDocumentoForm:comboEscolherEstado",
            [uf],
            timeout=10,
        )
        registrar_evento_execucao(f"UF do prestador selecionada: {uf}", "ISS Fortaleza")
        # Aguarda o AJAX recarregar o combo de cidades após a seleção de estado
        time.sleep(1.5)

    # --- Cidade do Prestador ---
    cidade = documento.cidade_prestador.strip()
    if cidade:
        _selecionar_opcao_por_texto(
            driver,
            By.ID,
            "digitarDocumentoForm:comboEscolherCidade",
            [cidade],
            timeout=15,
        )
        registrar_evento_execucao(f"Cidade do prestador selecionada: {cidade}", "ISS Fortaleza")
        # Aguarda o AJAX da Cidade processar antes de digitar o CEP — o combo de
        # cidade também dispara um postback do próprio portal (mesmo padrão de
        # estabilização já usado acima para a UF), e sem essa espera o CEP pode
        # ser digitado com sucesso e depois limpo por uma resposta tardia desse
        # AJAX, sem que nada perceba.
        time.sleep(1.5)

    # --- Tipo de Tributação (só existe na tela quando o prestador não é de
    # Fortaleza/CE — para prestador de Fortaleza/CE o campo nem aparece no DOM) ---
    if not _eh_prestador_fortaleza_ce(documento):
        _selecionar_opcao_por_texto(
            driver,
            By.ID,
            "digitarDocumentoForm:tipoTributacaoPrestadorExternoId",
            [documento.tipo_tributacao],
            timeout=10,
        )
        registrar_evento_execucao(
            f"Tipo de Tributação definido como '{documento.tipo_tributacao}'.",
            "ISS Fortaleza",
        )

    # --- CEP ---
    cep = re.sub(r"\D+", "", documento.cep_prestador or "")
    if cep:
        _digitar_campo(
            driver,
            By.ID,
            "digitarDocumentoForm:idCEP",
            cep,
            delay_ms=40,
        )

        # Confere de novo após uma pequena espera extra: se uma resposta AJAX
        # ainda mais tardia (da seleção de Cidade) limpar o campo, redigita e
        # registra a correção no log de auditoria.
        time.sleep(0.8)
        try:
            valor_pos_espera = re.sub(r"\D+", "", driver.find_element(By.ID, "digitarDocumentoForm:idCEP").get_attribute("value") or "")
        except NoSuchElementException:
            valor_pos_espera = None
        if valor_pos_espera is not None and valor_pos_espera != cep:
            registrar_evento_execucao(
                f"CEP do prestador foi sobrescrito/limpo pelo portal após a digitação "
                f"(ficou '{valor_pos_espera}'); redigitando '{cep}'.",
                "ISS Fortaleza",
            )
            _digitar_campo(
                driver,
                By.ID,
                "digitarDocumentoForm:idCEP",
                cep,
                delay_ms=40,
            )

        registrar_evento_execucao(f"CEP preenchido: {cep}", "ISS Fortaleza")

    # --- Logradouro ---
    logradouro = documento.logradouro_prestador.strip()
    if logradouro:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idEndereco",
            logradouro,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Logradouro preenchido: {logradouro}", "ISS Fortaleza")

    # --- Número (usa 'S/N' se o campo estiver vazio na planilha) ---
    numero = documento.numero_prestador.strip() or "S/N"
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:idNumero",
        numero,
        delay_ms=40,
    )
    registrar_evento_execucao(f"Número preenchido: {numero}", "ISS Fortaleza")

    # --- Bairro ---
    bairro = documento.bairro_prestador.strip()
    if bairro:
        _definir_campo_rapido(
            driver,
            By.ID,
            "digitarDocumentoForm:idBairro",
            bairro,
            delay_ms=20,
        )
        registrar_evento_execucao(f"Bairro preenchido: {bairro}", "ISS Fortaleza")

    # --- Email (campo não obrigatório — preenche somente se disponível) ---
    email = documento.email_prestador.strip()
    if email:
        try:
            _definir_campo_rapido(
                driver,
                By.ID,
                "digitarDocumentoForm:inputEmail3",
                email,
                delay_ms=20,
            )
            registrar_evento_execucao(f"E-mail preenchido: {email}", "ISS Fortaleza")
        except Exception as exc_email:
            # E-mail é opcional: registra o aviso mas não interrompe o fluxo
            registrar_evento_execucao(
                f"Aviso: não foi possível preencher o e-mail '{email}': {exc_email}",
                "ISS Fortaleza",
            )

    _pausar_para_depuracao(
        driver,
        depuracao,
        f"Dados do prestador {documento.cnpj_prestador} preenchidos. Confira antes de continuar.",
    )
    registrar_evento_execucao(
        f"Dados do prestador preenchidos com sucesso para CNPJ {documento.cnpj_prestador}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Modal de pesquisa de CNAE
# ---------------------------------------------------------------------------


def _forcar_abertura_modal_cnae(driver: WebDriver) -> None:
    """Chama a API nativa do RichFaces via JavaScript para forçar a abertura do modal."""

    try:
        driver.execute_script("Richfaces.showModalPanel('pesquisarCnaeModal')")
    except JavascriptException:
        # Fallback: chama a função com prefixo alternativo que alguns portais usam
        driver.execute_script("RichFaces.showModalPanel('pesquisarCnaeModal')")


def preencher_modal_pesquisar_cnae(
    driver: WebDriver,
    id_cnae_final: str,
    depuracao: bool = False,
) -> None:
    """Preenche o modal de pesquisa de CNAE com o código final da planilha e dispara a busca.

    Caso o modal não apareça dentro do tempo padrão após o clique, utiliza a API
    nativa do RichFaces via JavaScript para forçá-lo a abrir — estratégia de
    fallback que corrige o timeout observado quando o JSF ainda está reprocessando.
    """

    valor_cnae = limpar_numero_para_digitacao(id_cnae_final)
    registrar_evento_execucao(
        f"Abrindo modal de pesquisa de CNAE para o código {valor_cnae}",
        "ISS Fortaleza",
    )

    # Tenta aguardar o modal ficar visível normalmente
    try:
        WebDriverWait(driver, 8).until(
            EC.visibility_of_element_located(
                (By.ID, "digitarDocumentoForm:pesquisarCnaeModalContainer")
            )
        )
    except TimeoutException:
        # Fallback via JavaScript: força o RichFaces a exibir o modal
        registrar_evento_execucao(
            "Modal de CNAE não abriu sozinho; forçando via JavaScript do RichFaces.",
            "ISS Fortaleza",
        )
        _forcar_abertura_modal_cnae(driver)
        # Aguarda mais tempo após a chamada forçada
        WebDriverWait(driver, 10).until(
            EC.visibility_of_element_located(
                (By.ID, "digitarDocumentoForm:pesquisarCnaeModalContainer")
            )
        )

    # Preenche o campo de pesquisa dentro do modal
    _definir_campo_rapido(
        driver,
        By.ID,
        "digitarDocumentoForm:idFormularioPesquisaCnae:idCnaePesquisa",
        valor_cnae,
        delay_ms=45,
    )
    _pausar_para_depuracao(
        driver,
        depuracao,
        f"O modal de pesquisa de CNAE está preenchido com {valor_cnae}.",
    )

    # Dispara a pesquisa AJAX do CNAE escolhido
    _clicar_por_id(driver, "digitarDocumentoForm:idFormularioPesquisaCnae:idPesquisar")
    time.sleep(1.2)
    registrar_evento_execucao(
        f"Pesquisa de CNAE enviada para {valor_cnae}",
        "ISS Fortaleza",
    )
    _pausar_para_depuracao(
        driver,
        depuracao,
        f"A pesquisa de CNAE para {valor_cnae} foi enviada. Confirme o resultado na lista.",
    )

    # Seleciona a primeira linha de resultado retornada pelo portal
    # O ID do link de seleção usa um índice dinâmico; buscamos com seletor parcial
    seletor_resultado = (
        "[id^='digitarDocumentoForm:idFormularioPesquisaCnae:idDatatableListaCnae:0:']"
    )
    el_resultado = WebDriverWait(driver, 10).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, seletor_resultado))
    )
    el_resultado.click()
    time.sleep(1.2)
    registrar_evento_execucao(
        f"CNAE selecionado na lista de resultados: {valor_cnae}",
        "ISS Fortaleza",
    )


def _digitar_campo_se_nao_zero(
    driver: WebDriver,
    seletor_id: str,
    valor: str,
) -> None:
    """Preenche o campo apenas se o valor for numérico e diferente de zero."""
    if not valor:
        return

    # Remove tudo que não for dígito para checar se é zero
    digitos = re.sub(r"\D+", "", valor)
    if not digitos or int(digitos) == 0:
        return

    _digitar_campo(
        driver,
        By.ID,
        seletor_id,
        limpar_valor_para_digitacao(valor),
    )


# ---------------------------------------------------------------------------
# Preenchimento da aba Serviço
# ---------------------------------------------------------------------------


def preencher_documento_servico(
    driver: WebDriver,
    documento: DocumentoPortalISS,
    depuracao: bool = False,
) -> None:
    """Preenche a aba de serviço com os dados de uma linha da planilha."""

    # Ativa a aba Serviço real do portal antes de preencher os campos
    _clicar_por_id(driver, "digitarDocumentoForm:abaServico_lbl")
    registrar_evento_execucao("Aba Serviço acionada", "ISS Fortaleza")
    time.sleep(0.5)

    # Aguarda os campos da aba Serviço ficarem disponíveis
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm"))
    )
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:tipoDocumentoDigitado"))
    )
    WebDriverWait(driver, 120).until(
        EC.visibility_of_element_located((By.ID, "digitarDocumentoForm:statusNfse"))
    )
    registrar_evento_execucao(
        "Formulário da aba Serviço ficou visível e pronto para preenchimento",
        "ISS Fortaleza",
    )

    registrar_evento_execucao(
        f"Iniciando preenchimento da aba Serviço para a NF {documento.numero_nf} ({documento.cnpj_prestador})",
        "ISS Fortaleza",
    )

    # Tipo de documento — MEI estabelecido em Fortaleza/CE (a exceção que faz a nota
    # ser escriturada apesar de ser da própria capital, ver regra em executar_fluxo_iss)
    # usa "NFS-e Nacional" em vez do "NFS-e de Outro Município" padrão.
    is_fortaleza_ce_doc = _eh_prestador_fortaleza_ce(documento)
    is_mei_doc = (documento.regime_tributario.upper().strip() == "MEI")

    if is_mei_doc and is_fortaleza_ce_doc:
        opcoes_tipo_documento = ["NFS-e Nacional", "NFS-e nacional"]
        registrar_evento_execucao(
            f"Prestador MEI estabelecido em Fortaleza/CE (NF {documento.numero_nf}): "
            "Tipo do Documento Digitado definido como 'NFS-e Nacional'.",
            "ISS Fortaleza",
        )
    elif _eh_cpf_prestador(documento):
        opcoes_tipo_documento = ["NFS Avulsa de outro município", "NFS Avulsa de outro municipio"]
        registrar_evento_execucao(
            f"Prestador pessoa física (NF {documento.numero_nf}): "
            "Tipo do Documento Digitado definido como 'NFS Avulsa de outro município'.",
            "ISS Fortaleza",
        )
    else:
        opcoes_tipo_documento = ["NFS-e de Outro Município", "NFS-e de outro município", "707"]

    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:tipoDocumentoDigitado",
        opcoes_tipo_documento,
    )

    # Número da nota fiscal
    # Usa digitação real (não _definir_campo_rapido): este campo é seguido de perto
    # por um select (Status NFSE) que dispara reprocessamento AJAX do JSF — sem
    # foco/clique real no campo, o valor definido via JS pode nunca ser confirmado
    # pelo lado servidor e acaba sendo limpo quando esse AJAX reprocessa o painel.
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:numeroDocumentoDigitado",
        limpar_numero_para_digitacao(documento.numero_nf),
    )

    # Data de emissão
    _digitar_campo(
        driver,
        By.ID,
        "digitarDocumentoForm:dataEmissaoInputDate",
        documento.data_emissao,
    )

    # Status NFSE
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:statusNfse",
        ["NORMAL", "Normal", "472"],
    )
    registrar_evento_execucao(
        f"Campos iniciais da aba Serviço preenchidos para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Pausa de estabilização para o JSF reprocessar após selecionar o Status
    time.sleep(1.2)

    # Confere se o reprocessamento do JSF não limpou o Número da NF (mesma classe
    # de corrida já vista para Natureza/ISS Retido e Cidade/CEP nesta sessão)
    numero_nf_limpo = limpar_numero_para_digitacao(documento.numero_nf)
    try:
        valor_pos_status = driver.find_element(By.ID, "digitarDocumentoForm:numeroDocumentoDigitado").get_attribute("value") or ""
    except NoSuchElementException:
        valor_pos_status = None
    if valor_pos_status is not None and _somente_digitos(valor_pos_status) != numero_nf_limpo:
        registrar_evento_execucao(
            f"Número da NF foi limpo pelo portal após selecionar Status NFSE "
            f"(ficou '{valor_pos_status}'); redigitando '{numero_nf_limpo}'.",
            "ISS Fortaleza",
        )
        _digitar_campo(driver, By.ID, "digitarDocumentoForm:numeroDocumentoDigitado", numero_nf_limpo, delay_ms=40)

    # Abre o modal de pesquisa de CNAE
    _clicar_por_id(driver, "digitarDocumentoForm:idLinkPesquisarCnae")
    registrar_evento_execucao(
        f"Botão Pesquisar CNAE acionado para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Preenche o modal de CNAE (com fallback via JavaScript)
    preencher_modal_pesquisar_cnae(driver, documento.id_cnae_final, depuracao=depuracao)

    # --- Alíquota (bloqueada pelo portal para alguns CNAEs, calculada automaticamente) ---
    aliquota = documento.aliquota.strip()
    if aliquota:
        if _campo_aliquota_esta_editavel(driver) is not None:
            _digitar_campo(driver, By.ID, "digitarDocumentoForm:idAliquota", aliquota, delay_ms=40)
            registrar_evento_execucao(f"Alíquota preenchida: {aliquota}", "ISS Fortaleza")
        else:
            registrar_evento_execucao(
                f"Alíquota não preenchida: campo bloqueado pelo portal para este CNAE "
                f"(calculado automaticamente). Valor da planilha: {aliquota}.",
                "ISS Fortaleza",
            )

    # Descrição do serviço preenchida de forma instantânea via JS
    _definir_valor_instantaneo(
        driver,
        By.ID,
        "digitarDocumentoForm:idDescricaoServico",
        documento.descricao_servico,
    )

    # UF de prestação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherEstadoLocalPrestacao",
        [documento.uf_local_prestacao],
    )

    # Cidade de prestação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherCidadeLocalPrestacao",
        [documento.cidade_local_prestacao],
    )

    # Natureza da operação
    _selecionar_opcao_por_texto(
        driver,
        By.ID,
        "digitarDocumentoForm:comboEscolherLocalPrestacao",
        [documento.natureza_operacao],
    )
    # Aguarda o AJAX da Natureza da Operação processar antes de mexer no checkbox
    # ISS Retido — o portal marca/desmarca esse checkbox sozinho ao mudar para
    # certas naturezas (ex.: "Tributação Fora do Município"), e sem essa espera
    # o código abaixo lê o estado do checkbox antes dessa reação do portal
    # terminar (mesmo padrão de estabilização já usado para Status NFSE e UF do
    # Prestador nesta função).
    time.sleep(1.5)

    # Checkbox ISS retido
    deve_marcar, iss_retido_reconhecido = _interpretar_iss_retido(documento.iss_retido)
    if not iss_retido_reconhecido:
        # Defesa em profundidade: validar_campos_obrigatorios() já deveria ter
        # barrado esta nota antes de chegar aqui, mas registra o aviso mesmo assim.
        registrar_evento_execucao(
            f"AVISO: valor de ISS_RETIDO '{documento.iss_retido}' não reconhecido para "
            f"a NF {documento.numero_nf}. Tratado como 'Não retido' por padrão.",
            "ISS Fortaleza",
        )

    # ID_CNAE (código bruto de serviço/tributação extraído do XML) tem
    # prioridade sobre o valor de ISS_RETIDO já decidido acima — cobre também
    # planilhas geradas antes dessa regra existir, ou editadas manualmente.
    deve_marcar_antes_do_override = deve_marcar
    deve_marcar = _aplicar_override_iss_retido_por_id_cnae(deve_marcar, documento.id_cnae)
    # Simples Nacional MEI tem prioridade máxima — aplicado por último, pode
    # sobrescrever até o override de ID_CNAE acima.
    deve_marcar = _aplicar_override_iss_retido_por_tipo_tributacao(deve_marcar, documento.tipo_tributacao)
    if deve_marcar != deve_marcar_antes_do_override:
        registrar_evento_execucao(
            f"ID_CNAE {re.sub(r'\\D+', '', documento.id_cnae or '')} / Tipo de Tributação "
            f"'{documento.tipo_tributacao}' na NF {documento.numero_nf}: ISS Retido forçado "
            f"para '{'Sim' if deve_marcar else 'Não'}' (planilha tinha '{documento.iss_retido}').",
            "ISS Fortaleza",
        )
    try:
        seletor_checkbox_iss = (By.NAME, "digitarDocumentoForm:j_id361")

        # A essa altura já esperamos 1.5s pelo AJAX da Natureza da Operação
        # (acima), então esta leitura reflete o estado já assentado do portal.
        marcado_agora = driver.find_element(*seletor_checkbox_iss).is_selected()

        if deve_marcar != marcado_agora:
            driver.find_element(*seletor_checkbox_iss).click()

        # Confere de novo após uma pequena espera extra: se uma resposta AJAX
        # ainda mais tardia sobrescrever o clique, corrige e registra a
        # correção no log de auditoria.
        estado_pos_clique = _ler_estado_checkbox_apos_espera(driver, *seletor_checkbox_iss, espera=1.0)
        if estado_pos_clique is not None and estado_pos_clique != deve_marcar:
            registrar_evento_execucao(
                f"Checkbox ISS Retido da NF {documento.numero_nf} foi sobrescrito pelo "
                f"portal logo após o clique (ficou "
                f"{'marcado' if estado_pos_clique else 'desmarcado'}); corrigindo para "
                f"{'marcado' if deve_marcar else 'desmarcado'}.",
                "ISS Fortaleza",
            )
            driver.find_element(*seletor_checkbox_iss).click()

        estado_final = driver.find_element(*seletor_checkbox_iss).is_selected()
        if estado_final != deve_marcar:
            registrar_evento_execucao(
                f"ALERTA: checkbox ISS Retido da NF {documento.numero_nf} ficou "
                f"{'marcado' if estado_final else 'desmarcado'}, mas o esperado era "
                f"{'marcado' if deve_marcar else 'desmarcado'} "
                f"(valor na planilha: '{documento.iss_retido}').",
                "ISS Fortaleza",
            )
        else:
            registrar_evento_execucao(
                f"Checkbox ISS Retido da NF {documento.numero_nf} definido como "
                f"{'marcado' if estado_final else 'desmarcado'} "
                f"(valor na planilha: '{documento.iss_retido}').",
                "ISS Fortaleza",
            )
    except NoSuchElementException:
        registrar_evento_execucao(
            f"ERRO: elemento do checkbox 'ISS Retido' (digitarDocumentoForm:j_id361) não "
            f"foi encontrado na tela para a NF {documento.numero_nf}. O estado do ISS "
            "Retido NÃO pôde ser verificado/definido nesta nota — revise manualmente no portal.",
            "ISS Fortaleza",
        )
    registrar_evento_execucao(
        f"Local de prestação e ISS tratados para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )

    # Valor do serviço e retenções/deduções (somente preenche se não for zero)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idValorServicoPrestado", documento.valor_servico)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idValorDeducoes", documento.valor_deducoes)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idDescontosIncondicionados", documento.descontos_incondicionados)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idDescontosCondicionados", documento.descontos_condicionados)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idOutrasRetencoes", documento.outras_retencoes)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idIR", documento.ir)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idPis", documento.pis_nao_retido)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idConfins", documento.cofins_nao_retido)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idCSLL", documento.csrf)
    _digitar_campo_se_nao_zero(driver, "digitarDocumentoForm:idINSS", documento.inss)

    registrar_evento_execucao(
        f"Aba Serviço preenchida com sucesso para a NF {documento.numero_nf}",
        "ISS Fortaleza",
    )


# ---------------------------------------------------------------------------
# Fluxo principal da ISS
# ---------------------------------------------------------------------------


def resetar_tela_para_digitar_documento(driver: WebDriver, competencia: CompetenciaTrabalho) -> None:
    """Tenta navegar de volta para a tela inicial de Digitar Documento caso a gravação de uma nota tenha falhado
    e o formulário tenha ficado travado em estado intermediário (efeito cascata)."""
    registrar_evento_execucao(
        "Iniciando recuperação de estado do portal para limpar formulário travado...",
        "ISS Fortaleza",
    )
    # 1. Clica no menu topo Escrituração
    try:
        el_menu = WebDriverWait(driver, 10).until(
            EC.element_to_be_clickable(
                (By.XPATH, "//a[contains(@class, 'dropdown-toggle') and (contains(., 'Escritura') or contains(., 'Escrituracao'))]")
            )
        )
        el_menu.click()
        time.sleep(0.3)
        _clicar_por_id(driver, "formMenuTopo:menuEscrituracao:j_id80")
    except Exception:
        # Fallback de menu por índice
        menus = driver.find_elements(By.CSS_SELECTOR, "a.dropdown-toggle")
        if len(menus) > 4:
            menus[4].click()
            time.sleep(0.3)
            _clicar_por_id(driver, "formMenuTopo:menuEscrituracao:j_id80")
            
    # 2. Aguarda a tela de Manter Escrituração e consulta a competência novamente
    WebDriverWait(driver, 30).until(
        EC.element_to_be_clickable((By.ID, "manterEscrituracaoForm:btnConsultar"))
    )
    selecionar_competencia_na_tela_richfaces(driver, competencia)
    _clicar_por_id(driver, "manterEscrituracaoForm:btnConsultar")
    
    # 3. Aguarda o botão Escriturar e clica
    aguardar_botao_escriturar(driver)
    _clicar_por_id(driver, "manterEscrituracaoForm:dataTable:0:linkEscriturar")
    
    # 4. Vai para a aba Serviços Tomados e clica em Digitar Documento
    clicar_aba_servicos_tomados(driver)
    WebDriverWait(driver, 30).until(
        EC.visibility_of_element_located((By.ID, "servico_tomado_form:seamj_id849"))
    )
    _clicar_por_id(driver, "servico_tomado_form:seamj_id849")
    aguardar_tela_digitar_documento(driver)
    time.sleep(0.5)


# ---------------------------------------------------------------------------
# Classificação e análise das notas (regras idênticas às do laço da escrituração individual;
# funções puras, sem navegador — usadas pela pré-análise e pelo laço de cada empresa)
# ---------------------------------------------------------------------------

CATEGORIA_ESCRITURAR = "ESCRITURAR"
CATEGORIA_FORTALEZA = "FORTALEZA"  # prestador de Fortaleza/CE que não é MEI: não entra na escrituração
CATEGORIA_INCOMPLETA = "INCOMPLETA"  # falta algum campo obrigatório na planilha


def classificar_documento(documento: DocumentoPortalISS) -> tuple[str, list[str]]:
    """Decide o destino da nota. Devolve (categoria, campos_ausentes). Ordem das regras igual
    à do original: primeiro Fortaleza/CE não MEI, depois campos obrigatórios."""

    cidade_prest = normalizar_texto(documento.cidade_prestador).upper()
    uf_prest = normalizar_texto(documento.uf_prestador).upper()
    is_fortaleza_ce = uf_prest == "CE" and "FORTALEZA" in cidade_prest
    is_mei = documento.regime_tributario.upper().strip() == "MEI"
    if is_fortaleza_ce and not is_mei:
        return CATEGORIA_FORTALEZA, []
    campos_ausentes = validar_campos_obrigatorios(documento)
    if campos_ausentes:
        return CATEGORIA_INCOMPLETA, campos_ausentes
    return CATEGORIA_ESCRITURAR, []


def chave_da_nota(documento: DocumentoPortalISS) -> str:
    """Identificador estável da nota dentro de uma empresa (arquivo, CNPJ do prestador e número da
    NF). Usado na retomada para não escriturar duas vezes a mesma nota. Não contém dado sigiloso."""

    return "|".join(
        (
            (documento.arquivo_pdf or "").strip(),
            limpar_cnpj_para_digitacao(documento.cnpj_prestador or ""),
            (documento.numero_nf or "").strip(),
        )
    )


@dataclass
class AnaliseNotas:
    """Resultado da pré-análise offline das notas de uma empresa."""

    total: int = 0
    a_escriturar: int = 0
    fortaleza: int = 0
    incompletas: int = 0

    def descricao(self) -> str:
        return (
            f"{self.total} nota(s): {self.a_escriturar} a escriturar, "
            f"{self.fortaleza} ignorada(s) por serem de prestador de Fortaleza/CE não MEI, "
            f"{self.incompletas} incompleta(s) (campo obrigatório ausente)"
        )


def analisar_documentos(documentos: list[DocumentoPortalISS]) -> AnaliseNotas:
    """Conta quantas notas seriam escrituradas, ignoradas ou incompletas, sem tocar no portal."""

    analise = AnaliseNotas(total=len(documentos))
    for documento in documentos:
        categoria, _ = classificar_documento(documento)
        if categoria == CATEGORIA_FORTALEZA:
            analise.fortaleza += 1
        elif categoria == CATEGORIA_INCOMPLETA:
            analise.incompletas += 1
        else:
            analise.a_escriturar += 1
    return analise


def agrupar_por_tomador(documentos: list[DocumentoPortalISS]) -> dict[str, list[DocumentoPortalISS]]:
    """Agrupa as notas pelo CNPJ_TOMADOR (empresa dona da escrituração), na ordem da planilha."""

    grupos: dict[str, list[DocumentoPortalISS]] = {}
    for documento in documentos:
        grupos.setdefault(limpar_cnpj_para_digitacao(documento.cnpj_tomador), []).append(documento)
    return grupos
